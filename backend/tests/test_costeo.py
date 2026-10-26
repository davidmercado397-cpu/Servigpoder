"""Costeo y MFP's: motor de costeo (contra los valores del Excel de Palmaseca), MFP con topes, autorizaciones,
bolsa de horas por año, reembolso por terminación anticipada, importación del Excel y costeo → MFP."""

from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest

from app.apps.costeo.services import calculo
from app.apps.costeo.services.calculo import PARAMETROS, Rango
from app.apps.costeo.services.mfp import meses_entre, sumar_meses
from tests.conftest import ACCESOS, archivo, entrar

API = "/api/costeo"
P = {k: v[0] for k, v in PARAMETROS.items()}
RANGOS = [Rango(h, pct, calculo.CARGA_PRESTACIONAL) for h, pct in calculo.COMISIONES]

SERV24 = "Servicio de vigilancia 24 horas de lunes a domingo todo el mes"
PALMASECA = {
    "duracion_meses": 41, "pct_nota_credito": 0.04, "pct_cupo_inversion": 0.0856, "horas_bolsa_anual": 0, "donaciones": 0,
    "servicios": [
        {"descripcion": SERV24, "puesto": "PORTERIA", "cantidad": 1, "tarifa": 18490849, "hombres": 3},
        {"descripcion": SERV24, "puesto": "MOTORIZADO", "cantidad": 1, "tarifa": 18838226, "hombres": 3},
        {"descripcion": SERV24, "puesto": "PORTERIA 1", "cantidad": 1, "tarifa": 18838226, "hombres": 3},
        {"descripcion": SERV24, "puesto": "PORTERIA 2", "cantidad": 1, "tarifa": 18490849, "hombres": 3},
        {"descripcion": "Apoyo porterías 12 h", "puesto": "APOYO", "cantidad": 2, "tarifa": 6556331, "hombres": 1.5},
    ],
    "personal": [
        {"nombre": "Guarda", "hombres": 12, "horas_extras": 1019110.6666666666, "auxilio_transporte": True, "auxilios_no_salariales": 100000},
        {"nombre": "Motorizado", "hombres": 3, "horas_extras": 1087014.3333333333, "auxilio_transporte": True,
         "auxilios_no_salariales": 200000, "rodamiento": 672000},
    ],
    "otros": [
        {"nombre": "ACCES TOO", "valor_unitario": 3600000, "cantidad": 1, "tipo_item": "ACCES_TOO"},
        {"nombre": "MODEM", "valor_unitario": 5017297, "cantidad": 4},
        {"nombre": "CELULAR", "valor_unitario": 700000, "cantidad": 3},
        {"nombre": "VISITRACK", "valor_unitario": 69020, "cantidad": 1, "vida_util_meses": 1},
        {"nombre": "BONO TARIFA", "valor_unitario": 11238358, "cantidad": 1, "tipo_item": "NOTA_CREDITO"},
        {"nombre": "DISPENSADOR", "valor_unitario": 1106700, "cantidad": 3},
        {"nombre": "IN HOUSE", "valor_unitario": 2500000, "cantidad": 1, "vida_util_meses": 1},
    ],
    "dotacion_tarifa": [{"nombre": "PANTALON", "valor_unitario": 44030, "cantidad": 15, "entregas_anio": 3},
                        {"nombre": "RADIO KENWOOD", "valor_unitario": 987700, "cantidad": 6, "vida_util_meses": 36}],
    "dotacion_puesto": [{"nombre": "RADIO ADICIONAL", "valor_unitario": 987700, "cantidad": 15, "vida_util_meses": 36}],
}


def test_motor_reproduce_el_excel():
    r = calculo.calcular(PALMASECA, P, RANGOS)
    assert r["tarifa_decreto"] == pytest.approx(87770812) and r["tarifa_neta"] == pytest.approx(84259979.52)
    assert r["hombres_servicios"] == 15 and r["hombres_personal"] == 15
    # Costo de un guarda sin el auxilio no salarial = el "total gasto personal" del Excel
    guarda = r["personal"][0]
    assert guarda["costo_hombre"] - 100000 == pytest.approx(4295913.40, abs=0.01)
    assert guarda["neto_hombre"] == pytest.approx(2957509.41, abs=0.01)  # el devengado del Excel ya incluye el auxilio
    # MFP, cupo de inversión, otros costos directos y comisiones, iguales al Excel
    assert r["costos"]["otros_directos"] == pytest.approx(3552621.12, abs=0.01)
    assert r["mfp"]["pct_total"] == pytest.approx(0.16607611, abs=1e-8)
    assert r["mfp"]["cupo_inversion_vp"] == pytest.approx(251689245.56, abs=0.01)
    assert r["costos"]["cupo_inversion"] == pytest.approx(6138762.09, abs=0.01)
    assert r["comision"]["pct"] == 0.002
    assert r["comision"]["vendedor_total"] == pytest.approx(5249913.39, abs=0.01)
    assert r["comision"]["director_total"] == pytest.approx(3387040.90, abs=0.01)
    assert r["costos"]["publicidad"] == pytest.approx(5095.97, abs=0.01)
    assert r["costos"]["representacion"] == pytest.approx(21407.52, abs=0.01)
    assert r["costos"]["ica"] == pytest.approx(56674.00, abs=0.01)
    # Dotación: anual ÷ 12 por entregas o mensual por vida útil, sin dividir dos veces
    assert r["costos"]["dotacion_tarifa"] == pytest.approx(44030 * 15 * 3 / 12 + 987700 * 6 / 36)
    assert r["costos"]["dotacion_puesto"] == pytest.approx(987700 * 15 / 36)
    # Con los auxilios no salariales el negocio pierde: la renta no se vuelve negativa
    assert r["margen_mensual"] < 0 and r["costos"]["renta"] == 0 and "negativo" in " ".join(r["alertas"])
    assert r["margen_total"] == pytest.approx(r["margen_mensual"] * 41)


def test_parametros_cambian_el_resultado():
    base = calculo.calcular(PALMASECA, P, RANGOS)
    sobre_contrato = calculo.calcular(PALMASECA, {**P, "publicidad_sobre_contrato": 1}, RANGOS)
    assert sobre_contrato["costos"]["publicidad"] == pytest.approx(base["costos"]["publicidad"] * 41)
    otra_comision = calculo.calcular(PALMASECA, P, [Rango(None, 0.01, 1.5)])
    assert otra_comision["comision"]["vendedor_total"] > base["comision"]["vendedor_total"]
    horas = calculo.calcular({**PALMASECA, "horas_bolsa_anual": 60}, P, RANGOS)
    assert horas["costos"]["horas_adicionales"] == pytest.approx(horas["costo_por_hombre"] / 30 / 12 * 60 / 12)


def test_fechas_de_contrato():
    assert sumar_meses(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert meses_entre(date(2026, 3, 1), date(2027, 3, 1)) == 12
    assert meses_entre(date(2026, 3, 1), date(2026, 3, 16)) == pytest.approx(15 / 31)


# ---------------- MFP ----------------

def _tipos(admin) -> dict[str, int]:
    return {t["codigo"]: t["id"] for t in admin.get(f"{API}/tipos-item").json()["data"]}


def _mfp(admin, **extra) -> dict:
    cid = admin.post(f"{API}/clientes", json={"nit": "900123", "nombre": "Conjunto Prueba"}).json()["data"]["id"]
    t = _tipos(admin)
    hoy = date.today()
    inicio = sumar_meses(hoy, -14)  # 14 meses atrás: va en el 2.º año del contrato
    datos = {"numero": "500", "cliente_id": cid, "contrato": "CT-001", "fecha_inicio": inicio.isoformat(), "duracion_meses": 36,
             "valor_mensual": 10_000_000, "pct_comercial": 0.10,
             "items": [{"tipo_id": t["REINVERSION"], "pactado": 36_000_000}, {"tipo_id": t["BOLSA_HORAS"], "pactado": 24},
                       {"tipo_id": t["SERVICIO_MENSUAL"], "pactado": 35_000, "descripcion": "Plan celular"},
                       {"tipo_id": t["PUBLICIDAD"], "pactado": 5_000_000}], **extra}
    r = admin.post(f"{API}/mfp", json=datos)
    assert r.status_code == 201, r.text
    return r.json()["data"]


def _item(m: dict, codigo: str) -> dict:
    return next(i for i in m["items"] if i["tipo_codigo"] == codigo)


def test_mfp_topes_horas_y_reembolso(admin):
    m = _mfp(admin)
    assert m["facturacion_total"] == 360_000_000 and m["estado"] == "vigente"
    # La publicidad es una bolsa interna: no suma en el % de MFP
    assert m["pactado"] == 36_000_000 + 35_000 * 36 and m["pct_mfp"] == pytest.approx(m["pactado"] / 360_000_000)
    hoy = date.today()
    reinv = _item(m, "REINVERSION")
    mov = lambda it, v, f=hoy, **x: admin.post(f"{API}/items/{it['id']}/movimientos",  # noqa: E731
                                                json={"fecha": f.isoformat(), "valor": v, "descripcion": "Pago", **x})
    assert mov(reinv, 30_000_000).status_code == 201

    # Superar lo pactado: sin la autorización escrita no se registra; el administrador puede autorizarlo
    r = mov(reinv, 10_000_000)
    assert r.status_code == 409 and r.json()["error"]["code"] == "TOPE_EXCEDIDO" and "supera el tope en $4.000.000" in r.json()["error"]["message"]
    r = mov(reinv, 10_000_000, autorizacion="Gerencia, acta 12 del 20/10/2026")
    assert r.status_code == 201
    m = r.json()["data"]
    assert _item(m, "REINVERSION")["excede"] and any(a["tipo"] == "excede" for a in m["alertas"])

    # Bolsa de horas por año de contrato: 24 h en el año 2; lo del año 1 no se usa (se perdió)
    horas = _item(m, "BOLSA_HORAS")
    assert [a["anio"] for a in horas["anios"]] == [1, 2, 3] and horas["anios"][1]["actual"]
    r = mov(horas, 2, descripcion="Acompañamiento 2 h por cumpleaños")
    h = _item(r.json()["data"], "BOLSA_HORAS")
    assert h["pendiente"] == 22 and h["anios"][1]["usado"] == 2 and h["anios"][0]["usado"] == 0
    r = mov(horas, 23)
    assert r.status_code == 409 and "22 h disponibles" in r.json()["error"]["message"]

    # Reembolso: lo ejecutado (40 M) × meses que faltan ÷ 36. A la fecha de inicio + 12 meses: 40 M × 24/36
    inicio = date.fromisoformat(m["fecha_inicio"])
    r = admin.get(f"{API}/mfp/{m['id']}?fecha_reembolso={sumar_meses(inicio, 12).isoformat()}").json()["data"]
    assert _item(r, "REINVERSION")["reembolso"] == pytest.approx(40_000_000 * 24 / 36)
    assert _item(r, "SERVICIO_MENSUAL")["reembolso"] == 0 and _item(r, "BOLSA_HORAS")["reembolso"] == 0
    # Después del fin del contrato no hay reembolso negativo
    r = admin.get(f"{API}/mfp/{m['id']}?fecha_reembolso={sumar_meses(inicio, 40).isoformat()}").json()["data"]
    assert r["reembolso"] == 0

    # Terminación anticipada
    r = admin.post(f"{API}/mfp/{m['id']}/terminar", json={"fecha": sumar_meses(inicio, 18).isoformat()}).json()["data"]
    assert r["estado"] == "terminada" and r["reembolso"] == pytest.approx(40_000_000 * 18 / 36)
    assert mov(horas, 1).status_code == 422  # sin bolsa después de terminar

    # Vistas generales
    res = admin.get(f"{API}/resumen").json()["data"]
    assert res["mfp"] == 1 and any(t["tipo"].startswith("Reinversión") and t["excedidos"] == 1 for t in res["por_tipo"])
    lista = admin.get(f"{API}/mfp?q=prueba").json()
    assert len(lista["data"]) == 1 and lista["meta"]["extra"]["totales"]["ejecutado"] == 40_000_000
    assert admin.get(f"{API}/clientes").json()["data"][0]["ejecutado"] == 40_000_000
    x = admin.get(f"{API}/mfp-excel")
    assert x.status_code == 200 and openpyxl.load_workbook(BytesIO(x.content)).active.max_row == 2
    assert admin.delete(f"{API}/mfp/{m['id']}").status_code == 409  # tiene movimientos


def test_programador_no_autoriza(admin):
    m = _mfp(admin)
    roles = {r["nombre"]: r["id"] for r in admin.get("/api/roles").json()["data"]}
    admin.post("/api/usuarios", json={"accesos": ACCESOS, "username": "prog", "nombre": "Programador",
                                      "password": "clave-segura-1", "roles": [roles["Programador MFP"]]})
    admin.post("/api/auth/logout")
    entrar(admin, "prog", "clave-segura-1")
    horas = _item(m, "BOLSA_HORAS")
    url = f"{API}/items/{horas['id']}/movimientos"
    hoy = date.today().isoformat()
    assert admin.post(url, json={"fecha": hoy, "valor": 4, "descripcion": "Evento"}).status_code == 201
    r = admin.post(url, json={"fecha": hoy, "valor": 30, "descripcion": "Evento", "autorizacion": "yo mismo lo autorizo"})
    assert r.status_code == 409 and r.json()["error"]["details"] == [{"puede_autorizar": False}]
    assert admin.post(f"{API}/mfp/{m['id']}/terminar", json={"fecha": hoy}).status_code == 403
    assert admin.put(f"{API}/parametros", json={"pct_renta": 30}).status_code == 403


# ---------------- Importación y costeo → MFP ----------------

def _excel_mfp() -> bytes:
    enc = ["NRO DE MFP FISICA", "NIT DEL CLIENTE", "NOMBRE DEL CLIENTE", "NRO DE CONTRATO", "FECHA DE INICIO ", "FECHA DE FINALIZACION",
           "DURACION EN MESES", "VALOR FACTURACION", "FACTURACION X CANTIDAD DE MESES", "FACTURACION REAL SISTEMA", "REINVERSION",
           "VALOR PAGADO DE REINVERSION", "FECHA DE PAGO DE REINVERSION", "% DE EJECUCION REINVERSION", "VALOR PENDIENTE POR EJECUTAR",
           "% PENDIENTE POR EJECUTAR", "NOTA CREDITO", "VALOR APLICADO DE NOTA CREDITO", "% DE NOTA CREDITO", "ACCES TOO",
           "VALOR APLICADO DE ACCES TOO", "VALOR PENDIENTE POR EJECUTAR", "CAPACITACION", "VALOR PAGADO CAPACITACION", "FECHA DE PAGO",
           "VALOR PENDIENTE POR EJECUTAR CAPACITACION", "OTROS", "MFP TOTAL", "% TOTAL DE MFP MILEYDA", "COMERCIAL FISICA"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "servig"
    ws.append(enc)
    ws.append([])
    ws.append([239, 800089377, "CONJUNTO A", "CT-002", datetime(2026, 9, 1), None, 12, 39372648, None, None, 0, None, None, None, None,
               None, 51971895, None, None, None, None, None, 0, 0, None, None, 0, None, None, 0.11])
    ws.append([238, 805025025, "TRANSPORTES B", "CT-001", datetime(2025, 11, 18), None, 36, 14939576, None, None, 50000000, 52000000,
               "2/10/2026 y 7/10/2026", None, None, None, 0, None, None, 3000000, 1000000, None, 12000000, 12000000, datetime(2026, 4, 9),
               None, 700000, None, None, 0.093])
    ws.append([238, 805025025, "TRANSPORTES B", "CT-002", None, None, None, 1000, None])
    otra = wb.create_sheet("Integrar")
    otra.append(enc)
    otra.append([1, 1, "OTRA EMPRESA", "X", datetime(2026, 1, 1), None, 12, 1000])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_importar_excel(admin):
    r = admin.post(f"{API}/mfp/importar", files=archivo(_excel_mfp(), "MFP GENERAL.xlsx"))
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["creadas"] == 3 and d["omitidas"] == 0
    textos = " ".join(d["advertencias"])
    assert "238 se repite" in textos and "sin fecha de inicio" in textos and "ejecutado mayor" in textos
    lista = admin.get(f"{API}/mfp").json()["data"]
    assert "OTRA EMPRESA" not in {x["cliente"]["nombre"] for x in lista}  # la hoja Integrar queda fuera
    b = next(x for x in lista if x["contrato"] == "CT-001")
    det = admin.get(f"{API}/mfp/{b['id']}").json()["data"]
    reinv = _item(det, "REINVERSION")
    assert reinv["pactado"] == 50_000_000 and reinv["ejecutado"] == 52_000_000 and reinv["excede"]
    assert reinv["movimientos"][0]["fecha"] == "2026-10-02" and _item(det, "CAPACITACION")["pendiente"] == 0
    assert _item(det, "OTROS")["ejecutado"] == 0
    # Volver a importar no duplica
    assert admin.post(f"{API}/mfp/importar", files=archivo(_excel_mfp())).json()["data"]["omitidas"] == 3


def test_costeo_aprobado_crea_la_mfp(admin):
    cid = admin.post(f"{API}/clientes", json={"nit": "1", "nombre": "Palmaseca"}).json()["data"]["id"]
    assert len(admin.get(f"{API}/catalogo?grupo=dotacion_tarifa").json()["data"]) > 10
    c = admin.post(f"{API}/costeos", json={"cliente_id": cid, "nombre": "Palmaseca", "datos": {**PALMASECA, "horas_bolsa_anual": 60}}).json()["data"]
    assert c["estado"] == "borrador" and c["resultado"]["tarifa_decreto"] == pytest.approx(87770812)
    assert admin.post(f"{API}/costeos/{c['id']}/crear-mfp", json={"numero": "900", "fecha_inicio": "2026-04-01"}).status_code == 409

    a = admin.post(f"{API}/costeos/{c['id']}/estado", json={"estado": "aprobado"}).json()["data"]
    assert a["estado"] == "aprobado" and a["parametros_aprobados"]
    # Cambiar un parámetro después no altera el costeo aprobado
    assert admin.put(f"{API}/parametros", json={"pct_renta": 30}).status_code == 200
    assert admin.get(f"{API}/costeos/{c['id']}").json()["data"]["resultado"]["margen_mensual"] == a["resultado"]["margen_mensual"]
    assert admin.put(f"{API}/costeos/{c['id']}", json={"cliente_id": cid, "nombre": "Otro nombre", "datos": PALMASECA}).status_code == 409

    r = admin.post(f"{API}/costeos/{c['id']}/crear-mfp", json={"numero": "900", "contrato": "CT-001", "fecha_inicio": "2026-04-01"})
    assert r.status_code == 201, r.text
    m = admin.get(f"{API}/mfp/{r.json()['data']['mfp_id']}").json()["data"]
    codigos = {i["tipo_codigo"] for i in m["items"]}
    assert {"NOTA_CREDITO", "REINVERSION", "ACCES_TOO", "OTROS", "BOLSA_HORAS", "PUBLICIDAD", "REPRESENTACION"} <= codigos
    assert _item(m, "BOLSA_HORAS")["pactado"] == 60 and m["valor_mensual"] == pytest.approx(87770812)
    assert m["pct_comercial"] == pytest.approx(0.16608, abs=1e-5)
    # Nueva versión: vuelve a borrador con los parámetros vigentes
    v2 = admin.post(f"{API}/costeos/{c['id']}/version").json()["data"]
    assert v2["version"] == 2 and v2["estado"] == "borrador"
    assert admin.post(f"{API}/costeos/calcular", json=PALMASECA).json()["data"]["costos"]["renta"] == 0
