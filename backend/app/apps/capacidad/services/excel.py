from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook


def leer_filas(contenido: bytes) -> list[list[object]]:
    """Filas de la primera hoja del libro, con los valores tal como vienen."""
    libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    try:
        hoja = libro.worksheets[0]
        # Algunos archivos exportados (SIESA) declaran mal su rango; se recalcula leyendo todo
        hoja.reset_dimensions()
        return [list(fila) for fila in hoja.iter_rows(values_only=True)]
    finally:
        libro.close()


def texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def a_fecha(valor: object) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return date.fromisoformat(texto(valor)[:10])


def a_hora(valor: object) -> time:
    if isinstance(valor, time):
        return valor
    if isinstance(valor, datetime):
        return valor.time()
    h, m = texto(valor).split(":")[:2]
    return time(int(h) % 24, int(m))
