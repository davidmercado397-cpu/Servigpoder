"""Ficha del personal: importación de la base de personal de gestión humana y del libro de vacaciones.

Base de personal (hojas OPERATIVOS, ADMINISTRATIVOS, RETIROS OP y RETIROS ADMIN): crea o actualiza cada empleado por
documento con sus datos personales y su historial de contratos. Si alguien está en OPERATIVOS y en ADMINISTRATIVOS
manda ADMINISTRATIVOS. Quien solo aparece en retiros queda retirado; quien está activo y en retiros es un reintegro.

Libro de vacaciones (hojas PERSONAL y FECHAS VACACIONES): el saldo de PERSONAL a su fecha de corte queda como punto
de partida (días disfrutados hasta el corte; la fecha de ingreso de la base manda sobre la del libro); las salidas de FECHAS VACACIONES quedan como historial (las posteriores al
corte, como programadas). Volver a importar el libro reemplaza ese historial.

Los archivos se procesan en memoria y no se guardan. Lo que no se puede leer se reporta con hoja y fila.
"""

import io
import re
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.apps.liquidador.models import LiqContrato, LiqEmpleado, LiqVacacionSalida
from app.apps.liquidador.services import festivos, vacaciones as vac

ACTIVO, RETIRADO = "activo", "retirado"
OPERATIVO, ADMINISTRATIVO = "operativo", "administrativo"


class ErrorArchivo(ValueError):
    pass


def _norm(texto: object) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


# Encabezado (normalizado, por prefijo) → campo
_COLUMNAS_BASE = (
    ("CEDULA", "documento"), ("NOMBRE Y APELLIDO", "nombre"), ("CARGO", "cargo"), ("PUESTO", "puesto"), ("CIUDAD", "ciudad"),
    ("SEXO", "sexo"), ("FECHA DE NACIMIENTO", "fecha_nacimiento"), ("RH", "rh"), ("CELULAR", "celular"),
    ("ID PAIS NACIMIENTO", "pais_nacimiento"), ("ID DEPTO NACIMIENTO", "depto_nacimiento"), ("ID CIUDAD NACIMIENTO", "ciudad_nacimiento"),
    ("FECHA EXPEDICION", "fecha_expedicion"), ("ID PAIS EXPEDICION", "pais_expedicion"), ("ID DEPTO EXPEDICION", "depto_expedicion"),
    ("ID CIUDAD EXPEDICION", "ciudad_expedicion"), ("INDICADOR ESTADO CIVIL", "estado_civil"), ("INDICADOR EXTRANJERO", "extranjero"),
    ("CORREO", "correo"), ("DIRECCION", "direccion"), ("BARRIO", "barrio"), ("CAJA DE COMPENSACION", "caja_compensacion"),
    ("FONDO DE PENSIONES", "fondo_pensiones"), ("FONDO DE CESANTIAS", "fondo_cesantias"), ("EPS", "eps"),
    ("PLAN EXEQUIAL", "plan_exequial"), ("NUMERO DE CUENTA", "numero_cuenta"), ("FECHA DE INICIO DEL CONTRATO", "fecha_inicio"),
    ("CONTRATO", "tipo_contrato"), ("NIVEL EDUCATIVO", "nivel_educativo"), ("CURSOS", "cursos"), ("FECHA DE RETIRO", "fecha_retiro"),
    ("OBSERVACION", "motivo_retiro"),
)
_LARGOS = {"puesto": 200, "tipo_contrato": 40, "rh": 5, "pais_nacimiento": 80, "depto_nacimiento": 80, "ciudad_nacimiento": 80,
           "pais_expedicion": 80, "depto_expedicion": 80, "ciudad_expedicion": 80, "celular": 30, "correo": 150, "direccion": 200,
           "barrio": 120, "ciudad": 80, "eps": 80, "fondo_pensiones": 80, "fondo_cesantias": 80, "caja_compensacion": 80,
           "plan_exequial": 80, "numero_cuenta": 40, "nivel_educativo": 60, "nombre": 200, "cargo": 80}
_NOMBRES_CAMPO = {"fecha_nacimiento": "fecha de nacimiento", "fecha_expedicion": "fecha de expedición", "fecha_inicio": "fecha de inicio del contrato",
                  "fecha_retiro": "fecha de retiro", "barrio": "barrio", "ciudad": "ciudad", "sexo": "sexo"}
_PUESTO = re.compile(r"^\d{3}\s*-\s*\d+")


class _Informe:
    def __init__(self) -> None:
        self.hallazgos: list[dict] = []

    def __call__(self, hoja: str, fila: int, documento: str | None, detalle: str) -> None:
        self.hallazgos.append({"hoja": hoja, "fila": fila, "documento": documento, "detalle": detalle})


class Indice(dict):
    """Empleados por documento; si no está exacto, busca por los números del documento (PPT 6623930 = 6623930)."""

    def __init__(self, empleados) -> None:
        super().__init__((e.documento, e) for e in empleados)
        self._numeros = {re.sub(r"\D", "", d): e for d, e in self.items() if not d.isdigit() and re.sub(r"\D", "", d)}

    def buscar(self, doc: str):
        if doc in self:
            return self[doc]
        if doc.isdigit():
            return self._numeros.get(doc)
        return self.get(re.sub(r"\D", "", doc))

    def __setitem__(self, doc, e) -> None:
        super().__setitem__(doc, e)
        if hasattr(self, "_numeros") and not doc.isdigit() and re.sub(r"\D", "", doc):
            self._numeros[re.sub(r"\D", "", doc)] = e


def documento(valor: object) -> tuple[str | None, bool]:
    """(documento, es_solo_numeros). Las cédulas que Excel guarda como número se leen sin decimales; los documentos
    con letras (p. ej. PPT 1234567) se dejan en mayúsculas con un solo espacio."""
    if valor is None or str(valor).strip() == "":
        return None, True
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor)), True
    if isinstance(valor, int):
        return str(valor), True
    texto = " ".join(str(valor).strip().upper().split())
    return texto[:32], texto.isdigit()


def _texto(valor: object, largo: int = 300) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = " ".join(str(valor).split())
    return texto[:largo] or None


def _fecha(valor: object) -> tuple[date | None, bool]:
    """(fecha, válida). Vacío no es error."""
    if valor is None or str(valor).strip() == "":
        return None, True
    if isinstance(valor, datetime):
        return valor.date(), True
    if isinstance(valor, date):
        return valor, True
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%Y%m%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date(), True
        except ValueError:
            continue
    return None, False


def _encabezados(fila: tuple) -> dict[str, int]:
    columnas: dict[str, int] = {}
    for i, celda in enumerate(fila):
        clave = _norm(celda)
        for prefijo, campo in _COLUMNAS_BASE:
            if clave.startswith(prefijo) and campo not in columnas:
                columnas[campo] = i
                break
    return columnas


def _hojas(libro, informe: _Informe):
    """Hojas de la base en orden: operativos, administrativos (mandan sobre operativos) y luego los retiros."""
    activos, retiros = [], []
    for hoja in libro.worksheets:
        nombre = _norm(hoja.title)
        tipo = ADMINISTRATIVO if "ADMIN" in nombre else OPERATIVO
        (retiros if "RETIRO" in nombre else activos).append((hoja, tipo))
    if not activos and not retiros:
        raise ErrorArchivo("El archivo no tiene las hojas de la base de personal (OPERATIVOS, ADMINISTRATIVOS, RETIROS…)")
    activos.sort(key=lambda x: x[1] == ADMINISTRATIVO)
    return activos, retiros


def _filas(hoja):
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None) or ()
    columnas = _encabezados(encabezado)
    if "documento" not in columnas:
        raise ErrorArchivo(f"La hoja {hoja.title} no tiene la columna CEDULA")
    # La columna sin título de texto después de OBSERVACION (retiros) trae el detalle del retiro
    extra = None
    if "motivo_retiro" in columnas:
        siguiente = columnas["motivo_retiro"] + 1
        if siguiente < len(encabezado) and not isinstance(encabezado[siguiente], str):
            extra = siguiente
    for num, fila in enumerate(filas, start=2):
        if not fila:
            continue
        valores = {campo: (fila[i] if i < len(fila) else None) for campo, i in columnas.items()}
        if extra is not None:
            valores["detalle_retiro"] = fila[extra] if extra < len(fila) else None
        yield num, valores


def _aplicar_datos(e: LiqEmpleado, v: dict, hoja: str, num: int, doc: str, informe: _Informe, solo_vacios: bool = False) -> None:
    """Copia a la ficha los campos que trae la fila (en retiros solo completa lo que esté vacío)."""
    def poner(campo: str, valor) -> None:
        if valor is None or (solo_vacios and getattr(e, campo) not in (None, "")):
            return
        setattr(e, campo, valor)

    for campo in ("puesto", "rh", "pais_nacimiento", "depto_nacimiento", "ciudad_nacimiento", "pais_expedicion", "depto_expedicion",
                  "ciudad_expedicion", "direccion", "eps", "fondo_cesantias", "caja_compensacion", "plan_exequial", "nivel_educativo",
                  "celular", "numero_cuenta", "tipo_contrato"):
        if campo in v:
            poner(campo, _texto(v[campo], _LARGOS[campo]))
    if "correo" in v:
        poner("correo", (_texto(v["correo"], 150) or "").lower() or None)
    if "cursos" in v:
        poner("cursos", _texto(v["cursos"], 2000))
    if "ciudad" in v and v["ciudad"] is not None:
        ciudad = _texto(v["ciudad"], 200) or ""
        if _PUESTO.match(ciudad):
            informe(hoja, num, doc, f"La columna CIUDAD trae un puesto («{ciudad}»): se guardó como puesto")
            poner("puesto", ciudad[:200])
        elif len(ciudad) < 3:
            informe(hoja, num, doc, f"Ciudad no válida («{ciudad}»): se dejó vacía")
        else:
            poner("ciudad", ciudad[:80])
    if "barrio" in v and v["barrio"] is not None:
        if isinstance(v["barrio"], (datetime, date)):
            informe(hoja, num, doc, "El barrio quedó convertido en fecha en Excel: se dejó vacío para corregirlo en la ficha")
        else:
            poner("barrio", _texto(v["barrio"], 120))
    if "sexo" in v and v["sexo"] is not None:
        sexo = _norm(v["sexo"])[:1]
        if sexo in ("M", "F"):
            poner("sexo", sexo)
        else:
            informe(hoja, num, doc, f"Sexo no válido («{v['sexo']}»)")
    for campo in ("fecha_nacimiento", "fecha_expedicion"):
        if campo in v:
            f, valida = _fecha(v[campo])
            if not valida:
                informe(hoja, num, doc, f"{_NOMBRES_CAMPO[campo].capitalize()} no válida («{v[campo]}»): se dejó vacía")
            poner(campo, f)
    if v.get("estado_civil") not in (None, ""):
        try:
            ec = int(float(v["estado_civil"]))
            if 0 <= ec <= 5:
                poner("estado_civil", ec)
            else:
                raise ValueError
        except (TypeError, ValueError):
            informe(hoja, num, doc, f"Estado civil no válido («{v['estado_civil']}»)")
    if v.get("extranjero") not in (None, "") and not solo_vacios:
        e.extranjero = str(v["extranjero"]).strip() in ("1", "1.0", "SI", "Si", "si")
    if v.get("fondo_pensiones") not in (None, ""):
        fondo = _texto(v["fondo_pensiones"], 120) or ""
        if "PENSIONAD" in _norm(fondo):  # «PENSIONADA VEJEZ - COLPENSIONES»: pensionado y su fondo (COLPENSIONES solo es afiliado)
            e.pensionado = True
            partes = [p.strip() for p in fondo.split("-") if p.strip()]
            poner("fondo_pensiones", partes[-1][:80] if len(partes) > 1 else None)
        else:
            if not solo_vacios:
                e.pensionado = False
            poner("fondo_pensiones", fondo[:80])


def sincronizar_estado(db: Session, e: LiqEmpleado) -> None:
    """Estado, fecha de ingreso y de retiro según los contratos: activo si alguno sigue abierto."""
    contratos = list(db.scalars(select(LiqContrato).where(LiqContrato.empleado_id == e.id).order_by(LiqContrato.fecha_inicio)))
    if not contratos:
        return
    abiertos = [c for c in contratos if c.fecha_retiro is None]
    if abiertos:
        actual = abiertos[-1]
        e.estado, e.fecha_retiro, e.fecha_ingreso = ACTIVO, None, actual.fecha_inicio
        e.tipo_contrato = actual.tipo_contrato or e.tipo_contrato
    else:
        ultimo = max(contratos, key=lambda c: c.fecha_retiro)
        e.estado, e.fecha_retiro = RETIRADO, ultimo.fecha_retiro
        e.fecha_ingreso = e.fecha_ingreso or ultimo.fecha_inicio


def importar_base(db: Session, contenido: bytes) -> dict:
    libro = _abrir(contenido)
    informe = _Informe()
    activos, retiros = _hojas(libro, informe)
    empleados = Indice(db.scalars(select(LiqEmpleado)))
    contratos = {(c.empleado_id, c.fecha_inicio): c for c in db.scalars(select(LiqContrato))}
    creados = actualizados = 0
    activos_en: dict[str, str] = {}
    tocados: dict[str, LiqEmpleado] = {}

    def empleado(doc: str, nombre: str | None, cargo: str | None) -> LiqEmpleado:
        nonlocal creados, actualizados
        e = empleados.buscar(doc)
        if e is None:
            e = LiqEmpleado(documento=doc, nombre=(nombre or doc)[:200], cargo=(cargo or "Operario")[:80], extranjero=False, pensionado=False)
            db.add(e)
            db.flush()
            empleados[doc] = e
            creados += 1
        elif doc not in tocados:
            actualizados += 1
        tocados[doc] = e
        return e

    def leer_doc(hoja: str, num: int, valor) -> str | None:
        doc, solo_numeros = documento(valor)
        if doc and not solo_numeros:
            informe(hoja, num, doc, "El documento tiene letras (p. ej. PPT): se cargó tal cual; verifique que coincida con el de la nómina")
        return doc

    for hoja, tipo in activos:
        for num, v in _filas(hoja):
            doc = leer_doc(hoja.title, num, v.get("documento"))
            if not doc:
                continue
            if doc in activos_en and activos_en[doc] != hoja.title:
                informe(hoja.title, num, doc, f"También está en {activos_en[doc]}: se toma la información de {hoja.title}")
            activos_en[doc] = hoja.title
            nombre, cargo = _texto(v.get("nombre"), 200), _texto(v.get("cargo"), 80)
            e = empleado(doc, nombre, cargo)
            e.nombre, e.cargo, e.tipo_personal = nombre or e.nombre, cargo or e.cargo, tipo
            _aplicar_datos(e, v, hoja.title, num, doc, informe)
            inicio, valida = _fecha(v.get("fecha_inicio"))
            if inicio is None:
                informe(hoja.title, num, doc, "Sin fecha de inicio del contrato válida: no se registró el contrato" if not valida
                        else "Sin fecha de inicio del contrato: no se registró el contrato")
                continue
            c = contratos.get((e.id, inicio)) or LiqContrato(empleado_id=e.id, fecha_inicio=inicio)
            if c.id is None:
                db.add(c)
                contratos[(e.id, inicio)] = c
            c.fecha_retiro, c.motivo_retiro = None, None
            c.tipo_contrato, c.cargo, c.puesto = e.tipo_contrato, e.cargo, e.puesto

    vistos_retiro: set[tuple[str, date]] = set()
    retirados = 0
    for hoja, tipo in retiros:
        for num, v in _filas(hoja):
            doc = leer_doc(hoja.title, num, v.get("documento"))
            if not doc:
                continue
            inicio, _ = _fecha(v.get("fecha_inicio"))
            retiro, valida_retiro = _fecha(v.get("fecha_retiro"))
            if inicio is None or retiro is None:
                informe(hoja.title, num, doc, "Sin fecha de inicio o de retiro válidas: no se registró el retiro")
                continue
            if (doc, inicio) in vistos_retiro:
                informe(hoja.title, num, doc, "Retiro repetido (misma cédula y fecha de inicio): se tomó una sola vez")
                continue
            vistos_retiro.add((doc, inicio))
            nombre, cargo = _texto(v.get("nombre"), 200), _texto(v.get("cargo"), 80)
            e = empleado(doc, nombre, cargo)
            if doc not in activos_en:
                e.nombre, e.cargo = nombre or e.nombre, cargo or e.cargo
                e.tipo_personal = e.tipo_personal or tipo
            _aplicar_datos(e, v, hoja.title, num, doc, informe, solo_vacios=doc in activos_en)
            c = contratos.get((e.id, inicio))
            if c is not None and c.fecha_retiro is None and doc in activos_en:
                informe(hoja.title, num, doc, f"Está activo en {activos_en[doc]} con el mismo contrato (inicio {inicio:%d/%m/%Y}) que aparece "
                        "retirado: se dejó activo")
                continue
            if c is None:
                c = LiqContrato(empleado_id=e.id, fecha_inicio=inicio)
                db.add(c)
                contratos[(e.id, inicio)] = c
            c.fecha_retiro = retiro
            c.motivo_retiro = _texto(v.get("motivo_retiro"), 200)
            c.observacion = _texto(v.get("detalle_retiro"), 300)
            c.cargo, c.tipo_contrato = cargo or c.cargo, _texto(v.get("tipo_contrato"), 40) or c.tipo_contrato
            puesto = _texto(v.get("puesto"), 200) or (e.puesto if doc not in activos_en else None)
            c.puesto = puesto or c.puesto
            retirados += 1
    db.flush()
    for e in tocados.values():
        sincronizar_estado(db, e)
    sin_base = sorted(d for d in empleados if d not in tocados)
    db.flush()
    resumen = {
        "empleados": len(tocados), "creados": creados, "actualizados": actualizados,
        "activos": sum(1 for e in tocados.values() if e.estado == ACTIVO),
        "retirados": sum(1 for e in tocados.values() if e.estado == RETIRADO), "contratos_retirados": retirados,
        "sin_base": len(sin_base), "hallazgos": informe.hallazgos,
    }
    if sin_base:
        resumen["sin_base_documentos"] = sin_base[:50]
    return resumen


# --- Libro de vacaciones ---

_MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9,
          "SETIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12}
_RANGO_DOS_MESES = re.compile(r"(\d{1,2}) DE ([A-Z]+) AL (\d{1,2}) DE ([A-Z]+) DEL? (\d{4,5})")
_RANGO_UN_MES = re.compile(r"(\d{1,2}) AL (\d{1,2}) DE ([A-Z]+) DEL? (\d{4,5})")
_DIA = re.compile(r"(\d{1,2}) DE ([A-Z]+) DEL? (\d{4,5})")


def _anio(texto: str) -> tuple[int | None, bool]:
    """(año, corregido): «20226» se lee como 2026 (se quita un dígito repetido por error de digitación)."""
    if len(texto) == 4:
        return int(texto), False
    candidatos = {int(texto[:i] + texto[i + 1:]) for i in range(len(texto) - 1) if texto[i] == texto[i + 1]}
    validos = sorted(a for a in candidatos if 2000 <= a <= 2100)
    return (validos[0], True) if len(validos) == 1 else (None, False)


def rango_vacaciones(texto: str) -> tuple[date | None, date | None, bool]:
    """«03 al 22 de julio del 2026» o «16 de julio al 03 de agosto del 2026» → (desde, hasta, año corregido)."""
    t = _norm(texto)
    m = _RANGO_DOS_MESES.search(t)
    if m:
        d1, m1, d2, m2, a = m.groups()
    else:
        m = _RANGO_UN_MES.search(t)
        if not m:
            return None, None, False
        d1, d2, m2, a = m.groups()
        m1 = m2
    anio, corregido = _anio(a)
    if anio is None or m1 not in _MESES or m2 not in _MESES:
        return None, None, False
    mes1, mes2 = _MESES[m1], _MESES[m2]
    try:
        hasta = date(anio, mes2, int(d2))
        desde = date(anio - 1 if mes1 > mes2 else anio, mes1, int(d1))
    except ValueError:
        return None, None, False
    return desde, hasta, corregido


def dia_texto(texto: str) -> date | None:
    m = _DIA.search(_norm(texto))
    if not m:
        return None
    anio, _ = _anio(m.group(3))
    if anio is None or m.group(2) not in _MESES:
        return None
    try:
        return date(anio, _MESES[m.group(2)], int(m.group(1)))
    except ValueError:
        return None


def _abrir(contenido: bytes):
    try:
        return load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:  # archivo dañado: no se expone el detalle técnico
        raise ErrorArchivo("No se pudo leer el archivo Excel. Verifica que sea un .xlsx válido.") from e


def _hoja(libro, nombre: str):
    for hoja in libro.worksheets:
        if _norm(hoja.title) == nombre:
            return hoja
    return None


def importar_vacaciones(db: Session, contenido: bytes) -> dict:
    libro = _abrir(contenido)
    informe = _Informe()
    personal, fechas = _hoja(libro, "PERSONAL"), _hoja(libro, "FECHAS VACACIONES")
    if personal is None:
        raise ErrorArchivo("El libro no tiene la hoja PERSONAL con el saldo de vacaciones")
    empleados = Indice(db.scalars(select(LiqEmpleado)))

    filas = personal.iter_rows(values_only=True)
    encabezado = [_norm(c) for c in (next(filas, None) or ())]

    def col(*nombres: str) -> int | None:
        return next((i for i, c in enumerate(encabezado) if any(c.startswith(n) for n in nombres)), None)

    c_doc, c_nombre, c_puesto = col("CEDULA"), col("NOMBRE"), col("PUESTO")
    c_ingreso, c_corte, c_disfrutados = col("FECHA INGRESO"), col("FECHA CORTE"), col("DIAS DISFRUTADOS")
    if None in (c_doc, c_ingreso, c_corte, c_disfrutados):
        raise ErrorArchivo("La hoja PERSONAL debe tener Cédula, Fecha Ingreso, Fecha Corte y Días Disfrutados")
    saldos = creados = 0
    cortes: set[date] = set()
    for num, fila in enumerate(filas, start=2):
        if not fila or c_doc >= len(fila):
            continue
        doc, _ = documento(fila[c_doc])
        if not doc:
            continue
        ingreso, _ = _fecha(fila[c_ingreso])
        corte, _ = _fecha(fila[c_corte])
        try:
            disfrutados = float(fila[c_disfrutados] or 0)
        except (TypeError, ValueError):
            informe("PERSONAL", num, doc, f"Días disfrutados no válidos («{fila[c_disfrutados]}»): no se cargó el saldo")
            continue
        if ingreso is None or corte is None:
            informe("PERSONAL", num, doc, "Sin fecha de ingreso o de corte válidas: no se cargó el saldo")
            continue
        e = empleados.buscar(doc)
        if e is not None and e.documento != doc:
            informe("PERSONAL", num, doc, f"Se tomó como el documento {e.documento} de la base de personal")
        if e is None:
            nombre = _texto(fila[c_nombre], 200) if c_nombre is not None else None
            e = LiqEmpleado(documento=doc, nombre=nombre or doc, cargo="Operario", extranjero=False, pensionado=False,
                            puesto=_texto(fila[c_puesto], 200) if c_puesto is not None else None)
            db.add(e)
            db.flush()
            empleados[doc] = e
            creados += 1
            informe("PERSONAL", num, doc, "No está en la base de personal: se creó con nombre y puesto; complete su ficha")
        if e.fecha_ingreso and e.fecha_ingreso != ingreso:
            informe("PERSONAL", num, doc, f"La fecha de ingreso del libro ({ingreso:%d/%m/%Y}) es distinta a la de la base "
                    f"({e.fecha_ingreso:%d/%m/%Y}): se deja la de la base")
        e.fecha_ingreso = e.fecha_ingreso or ingreso
        e.vacaciones_corte = corte
        e.vacaciones_previas = round(disfrutados, 2)
        cortes.add(corte)
        saldos += 1

    salidas = programadas = compensadas = 0
    if fechas is not None:
        db.execute(delete(LiqVacacionSalida).where(LiqVacacionSalida.origen == "libro"))
        corte_libro = max(cortes) if cortes else date.today()
        filas_f = list(fechas.iter_rows(values_only=True))
        rango_fest = festivos.conjunto(db, date(corte_libro.year - 1, 1, 1), date(corte_libro.year + 1, 12, 31))
        for num, fila in enumerate(filas_f, start=1):
            f = list(fila or ()) + [None] * 7
            doc, solo_numeros = documento(f[0])
            if not doc or not (solo_numeros or re.match(r"^[A-Z]{2,4} \d+$", doc)):
                continue  # títulos de mes y encabezados
            e = empleados.buscar(doc)
            if e is None:
                informe("FECHAS VACACIONES", num, doc, "No está en la base de personal: no se registró la salida")
                continue
            texto, retorno_txt, nota = str(f[4] or ""), str(f[5] or ""), _texto(f[6], 300)
            x = LiqVacacionSalida(empleado_id=e.id, origen="libro", observacion=nota)
            if "COMPENS" in _norm(texto):
                x.tipo, x.estado = "compensadas", "historico"
                compensadas += 1
            else:
                desde, hasta, corregido = rango_vacaciones(texto)
                if desde is None:
                    informe("FECHAS VACACIONES", num, doc, f"No se entendió el periodo «{texto}»: no se registró la salida")
                    continue
                if corregido:
                    informe("FECHAS VACACIONES", num, doc, f"El año de «{texto}» tiene un error de digitación: se leyó como {hasta.year}")
                x.tipo, x.desde, x.hasta, x.retorno = "disfrutadas", desde, hasta, dia_texto(retorno_txt)
                x.dias_habiles = sum(1 for i in range((hasta - desde).days + 1)
                                     if vac.es_habil(desde + timedelta(days=i), e.dias_habiles, rango_fest))
                corte_e = e.vacaciones_corte or corte_libro
                x.estado = "programada" if desde > corte_e else "historico"
            if nota and "CANCEL" in _norm(nota):
                x.estado = "cancelada"
            programadas += x.estado == "programada"
            db.add(x)
            salidas += 1
    db.flush()
    return {"saldos": saldos, "creados": creados, "salidas": salidas, "programadas": programadas, "compensadas": compensadas,
            "corte": max(cortes).isoformat() if cortes else None, "hallazgos": informe.hallazgos}
