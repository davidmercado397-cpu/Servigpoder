import json
from io import BytesIO

from openpyxl import load_workbook

from app.models import Usuario
from app.apps.liquidador.services import asistente
from tests.conftest import archivo, entrar, xlsx

API = "/api/liquidador"


def _quincena(cliente, anio=2026, mes=9, quincena=2) -> int:
    r = cliente.post(f"{API}/periodos", json={"anio": anio, "mes": mes, "quincena": quincena})
    assert r.status_code == 201, r.text
    return r.json()["data"]["id"]


def _excel_sep_q2() -> bytes:
    # Septiembre 2026: 20 y 27 son domingo; no hay festivos en la segunda quincena
    dias = list(range(16, 31))
    filas = [["documento", "SEPTIEMBRE", *dias]]
    filas.append([1001, 1001, *(["A"] * 4 + ["Z"] + ["A"] * 6 + ["Z"] + ["A"] * 3)])  # 16-19 A, 20 Z, 21-26 A, 27 Z, 28-30 A
    filas.append(["2002", "2002", *(["Z", "Z", "N", "N", "Z", "D", "D", "N", "N", "Z", "Z", "D", "D", "N", "N"])])
    filas.append(["3003", "3003", *(["V"] * 5 + ["INC"] * 2 + ["LR"] + ["AUS"] + ["L"] + ["A"] * 4 + ["XX"])])
    return xlsx(filas)


def _cargar(cliente, pid: int, contenido: bytes | None = None):
    return cliente.post(f"{API}/periodos/{pid}/cargar", files=archivo(contenido or _excel_sep_q2(), "turnos.xlsx"))


def _resultados(cliente, pid: int) -> dict[str, dict]:
    datos = cliente.get(f"{API}/periodos/{pid}/resultados?tamano=150").json()["data"]
    return {r["documento"]: r for r in datos}


def test_configuracion_inicial_trae_los_turnos_de_produccion(admin):
    r = admin.get(f"{API}/turnos?tamano=150").json()
    assert r["meta"]["extra"]["total"] == 129
    por_codigo = {t["codigo"]: t for t in r["data"]}
    assert por_codigo["D"]["horas_ordinarias"] == 8.4 and por_codigo["D"]["horas_extras"] == 2.8
    assert por_codigo["INC"]["incapacidad"] and por_codigo["INC"]["clase"] == "incapacidad"
    assert por_codigo["V"]["clase"] == "novedad" and por_codigo["Z"]["clase"] == "descanso_pago"
    assert por_codigo["A"]["clase"] == "trabajado"
    assert admin.get(f"{API}/parametros").json()["data"]["hora_inicio_nocturna"] == 19


def test_carga_cuenta_horas_y_dias(admin):
    pid = _quincena(admin)
    r = _cargar(admin, pid)
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["empleados"] == 3 and d["periodo"]["estado"] == "calculada" and d["periodo"]["empleados"] == 3
    assert any("'XX' no existe" in a for a in d["advertencias"])

    res = _resultados(admin, pid)
    # A = 06:00, 7 ordinarias + 1 extra; 13 días hábiles (ningún domingo trabajado)
    a = res["1001"]
    assert a["horas"]["ordinary_day"] == 91 and a["horas"]["overtime_day"] == 13
    assert a["dias"] == {"trabajado": 13, "descanso_pago": 2, "libre": 0, "ausencia": 0, "incapacidad": 0, "novedad": 0}
    assert a["total_horas"] == 104

    # V y LR (sin horas) cuentan como días con novedad, no como días trabajados
    c = res["3003"]
    assert c["dias"] == {"trabajado": 4, "descanso_pago": 0, "libre": 1, "ausencia": 1, "incapacidad": 2, "novedad": 6}

    # N del sábado 19 → recargo nocturno dominical a partir de la medianoche
    b = admin.get(f"{API}/periodos/{pid}/empleados/{res['2002']['empleado_id']}").json()["data"]
    sabado = next(x for x in b["dias"] if x["fecha"] == "2026-09-19")
    assert sabado["tipo_dia"] == "saturday" and sabado["clase"] == "trabajado"
    assert sabado["horas"]["sunday_night_surcharge"] == 2.4 and sabado["horas"]["ordinary_night"] == 5


def test_archivo_de_liquidacion_conserva_columnas_y_agrega_novedad_al_final(admin):
    pid = _quincena(admin)
    _cargar(admin, pid)
    r = admin.get(f"{API}/periodos/{pid}/liquidacion")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="liquidacion_2026_09_Q2.xlsx"'
    hoja = load_workbook(BytesIO(r.content)).active
    assert hoja.title == "2026-09-Q2"
    filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    assert filas[0][:19] == [
        "DOCUMENTO", "EMPLEADO", "DIURNAS ORDINARIA", "RECARGO NOCTURNO", "RECARGO FESTIVO DIURNO",
        "RECARGO FESTIVO NOCTURNO", "RECARGO DOMINICAL DIURNO", "RECARGO DOMINICAL NOCTURNO", "EXTRAS ORDINARIA DIURNAS",
        "EXTRAS ORDINARIA NOCTURNAS", "EXTRAS FESTIVAS DIURNAS", "EXTRAS FESTIVAS NOCTURNAS", "EXTRAS DOMINICALES DIURNAS",
        "EXTRAS DOMINICALES NOCTURNAS", "DIAS TRABAJADOS", "DESCANSOS PAGOS", "LIBRES", "AUSENCIAS", "INCAPACIDADES"]
    assert filas[0][19] == "DIAS CON NOVEDAD" and len(filas[0]) == 20
    assert [f[0] for f in filas[1:]] == ["1001", "2002", "3003"]
    assert filas[1][2:9] == [91, 0, 0, 0, 0, 0, 13] and filas[1][14:] == [13, 2, 0, 0, 0, 0]
    assert filas[3][14:] == [4, 0, 1, 1, 2, 6]


def test_plantilla_de_la_quincena(admin):
    pid = _quincena(admin)
    r = admin.get(f"{API}/periodos/{pid}/plantilla")
    hoja = load_workbook(BytesIO(r.content)).active
    encabezado = next(hoja.iter_rows(values_only=True))
    assert encabezado[:3] == ("documento", "SEPTIEMBRE 2026", 16) and encabezado[-1] == 30


def test_validaciones_del_excel(admin):
    pid = _quincena(admin)
    r = _cargar(admin, pid, xlsx([["cedula", "x", 16], [1, 1, "A"]]))
    assert r.status_code == 422 and "documento" in r.json()["error"]["message"]
    r = _cargar(admin, pid, xlsx([["documento", "SEPTIEMBRE", 1, 2], [1, 1, "A", "A"]]))
    assert r.status_code == 422 and "1 al 2 de septiembre" in r.json()["error"]["message"]
    r = _cargar(admin, pid, xlsx([["documento", "SEPTIEMBRE", 16], [1, 1, "NOEXISTE"]]))
    assert r.status_code == 422 and "ningún día" in r.json()["error"]["message"]
    assert admin.get(f"{API}/periodos/{pid}").json()["data"]["estado"] == "borrador"


def test_cierre_bloquea_cambios_y_reabrir_los_permite(admin):
    pid = _quincena(admin)
    assert admin.post(f"{API}/periodos/{pid}/cerrar").status_code == 409  # sin Excel
    _cargar(admin, pid)
    assert admin.post(f"{API}/periodos/{pid}/cerrar").json()["data"]["estado"] == "cerrada"
    r = _cargar(admin, pid)
    assert r.status_code == 409 and r.json()["error"]["code"] == "PERIODO_CERRADO"
    assert admin.post(f"{API}/periodos/{pid}/recalcular").status_code == 409
    assert admin.delete(f"{API}/periodos/{pid}").status_code == 409
    assert admin.post(f"{API}/periodos/{pid}/reabrir").json()["data"]["estado"] == "calculada"
    assert _cargar(admin, pid).status_code == 200


def test_festivo_manual_y_recalculo(admin):
    pid = _quincena(admin)
    _cargar(admin, pid)
    antes = _resultados(admin, pid)["1001"]["horas"]
    assert antes["holiday_day_surcharge"] == 0
    # El jueves 17 pasa a ser festivo: el turno A de ese día genera recargo festivo y extras festivas
    assert admin.post(f"{API}/festivos", json={"fecha": "2026-09-17", "descripcion": "Prueba"}).status_code == 201
    admin.post(f"{API}/periodos/{pid}/recalcular")
    despues = _resultados(admin, pid)["1001"]["horas"]
    assert despues["holiday_day_surcharge"] == 7 and despues["holiday_overtime_day"] == 1 and despues["overtime_day"] == 12

    festivos = admin.get(f"{API}/festivos?anio=2026").json()["data"]
    manual = next(f for f in festivos if f["fecha"] == "2026-09-17")
    assert manual["origen"] == "manual"
    # Quitar un festivo nacional y restablecerlo
    r = admin.post(f"{API}/festivos/quitar", json={"fecha": "2026-12-25"}).json()["data"]
    assert r["origen"] == "quitado" and not r["vigente"]
    assert admin.delete(f"{API}/festivos/ajustes/{r['ajuste_id']}").status_code == 200
    assert next(f for f in admin.get(f"{API}/festivos?anio=2026").json()["data"] if f["fecha"] == "2026-12-25")["vigente"]


def test_turnos_crear_editar_y_hora_nocturna(admin):
    datos = {"codigo": "t1", "nombre": "Prueba", "hora_inicio": "14:00", "hora_fin": "22:00",
             "horas_ordinarias": 8, "horas_extras": 0, "remunerado": True, "incapacidad": False}
    r = admin.post(f"{API}/turnos", json=datos)
    assert r.status_code == 201
    t = r.json()["data"]
    assert t["codigo"] == "T1" and t["matriz"]["ordinary_night"]["weekday"] == 3  # 19:00 a 22:00
    assert admin.post(f"{API}/turnos", json=datos).status_code == 409

    # Incapacidad: nunca remunerada ni con horas
    r = admin.put(f"{API}/turnos/{t['id']}", json=datos | {"incapacidad": True})
    assert r.json()["data"]["horas_ordinarias"] == 0 and not r.json()["data"]["remunerado"]
    admin.put(f"{API}/turnos/{t['id']}", json=datos)

    # La hora nocturna regenera todas las matrices
    r = admin.put(f"{API}/parametros", json={"hora_inicio_nocturna": 21})
    assert r.json()["meta"]["extra"]["turnos_regenerados"] == 130
    assert admin.get(f"{API}/turnos/{t['id']}").json()["data"]["matriz"]["ordinary_night"]["weekday"] == 1

    previa = admin.post(f"{API}/turnos/vista-previa", json={"hora_inicio": "18:00", "horas_ordinarias": 8.4, "horas_extras": 2.8}).json()["data"]
    assert previa["ordinary_night"]["weekday"] == 5.4  # 21:00 a 02:24
    assert admin.post(f"{API}/turnos/{t['id']}/activo").json()["data"]["activo"] is False
    assert admin.post(f"{API}/turnos", json=datos | {"codigo": "X2", "horas_ordinarias": 20, "horas_extras": 5}).status_code == 422


def test_empleados_y_busqueda(admin):
    pid = _quincena(admin)
    _cargar(admin, pid)
    e = admin.get(f"{API}/empleados?q=300").json()
    assert e["meta"]["extra"]["total"] == 1 and e["data"][0]["ultima"] == "2026-09 Q2" and e["data"][0]["quincenas"] == 1
    r = admin.get(f"{API}/periodos/{pid}/resultados?q=2002").json()
    assert [x["documento"] for x in r["data"]] == ["2002"]
    assert r["meta"]["extra"]["totales"]["dias"]["novedad"] == 0


def test_permisos_del_liquidador(admin):
    admin.post("/api/roles", json={"nombre": "Consulta liq", "permisos": ["liquidador.periodos.ver"]})
    rid = next(r["id"] for r in admin.get("/api/roles").json()["data"] if r["nombre"] == "Consulta liq")
    admin.post("/api/usuarios", json={"username": "consulta", "nombre": "Consulta", "password": "clave-segura-1", "roles": [rid]})
    pid = _quincena(admin)
    admin.post("/api/auth/logout")
    entrar(admin, "consulta", "clave-segura-1")
    assert admin.get(f"{API}/periodos").status_code == 200
    assert admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 10, "quincena": 1}).status_code == 403
    assert _cargar(admin, pid).status_code == 403
    assert admin.get(f"{API}/turnos").status_code == 403
    assert admin.get("/api/capacidad/maestros/puestos").status_code == 403
    apps = [a["codigo"] for a in admin.get("/api/plataforma/apps").json()["data"]]
    assert apps == ["liquidador"]


def test_herramientas_del_asistente(admin, db):
    pid = _quincena(admin)
    _cargar(admin, pid)
    with db() as s:
        usuario = s.query(Usuario).filter_by(username="admin").one()
        h = {t.nombre: t for t in asistente.crear_herramientas(asistente.Contexto(s, usuario, True))}
        estado = json.loads(h["estado_datos"].call())
        assert estado["quincenas"][0]["quincena"] == "2026-09 Q2" and estado["quincenas"][0]["personas"] == 3
        resumen = json.loads(h["resumen_quincena"].call({"anio": 2026, "mes": 9, "quincena": 2}))
        assert resumen["personas"] == 3 and resumen["dias"]["días con novedad (vacaciones, licencias…)"] == 6
        detalle = json.loads(h["detalle_persona"].call({"documento": "1001"}))
        assert detalle["horas_trabajadas"] == 104 and len(detalle["dia_a_dia"]) == 15
        turno = json.loads(h["explicar_turno"].call({"codigo": "n"}))
        assert turno["matriz_por_tipo_de_dia"]["Sábado"]["Recargo nocturno dominical"] == 2.4
        assert "error" in json.loads(h["detalle_persona"].call({"documento": "999"}))
        sin_datos = asistente.crear_herramientas(asistente.Contexto(s, usuario, False))
        assert "1001" not in next(t for t in sin_datos if t.nombre == "resumen_quincena").call()
