"""Carga inicial de las MFP desde el Excel con el que se llevaban ("MFP GENERAL SERVIGPODER").

Se lee la hoja de Servigpoder (la de otras empresas, como "Integrar", queda por fuera). Por cada fila: cliente
(por NIT), MFP y sus ítems pactados; lo ya ejecutado entra como un movimiento "saldo inicial" con la fecha que
traiga el Excel (o la de inicio del contrato). El Excel no se guarda. Una MFP que ya existe (mismo número, NIT,
contrato y fecha de inicio) no se vuelve a crear.
"""

import re
from datetime import date, datetime
from io import BytesIO

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.apps.costeo.models import CstCliente, CstMfp, CstMfpItem, CstMovimiento, CstTipoItem

HOJA_EXCLUIDA = {"INTEGRAR", "HOJA2"}

# Encabezado normalizado → campo
COLUMNAS = {
    "NRO DE MFP FISICA": "numero", "NIT DEL CLIENTE": "nit", "NOMBRE DEL CLIENTE": "cliente", "NRO DE CONTRATO": "contrato",
    "FECHA DE INICIO": "inicio", "DURACION EN MESES": "duracion", "VALOR FACTURACION": "valor_mensual",
    "REINVERSION": "reinversion", "VALOR PAGADO DE REINVERSION": "reinversion_pagado", "FECHA DE PAGO DE REINVERSION": "reinversion_fecha",
    "NOTA CREDITO": "nota", "VALOR APLICADO DE NOTA CREDITO": "nota_aplicado", "ACCES TOO": "acces", "VALOR APLICADO DE ACCES TOO": "acces_aplicado",
    "CAPACITACION": "capacitacion", "VALOR PAGADO CAPACITACION": "capacitacion_pagado", "FECHA DE PAGO": "capacitacion_fecha",
    "OTROS": "otros", "COMERCIAL FISICA": "pct_comercial", "OTRO SI": "otrosi",
}
# ítem: (tipo, campo pactado, campo ejecutado, campo fecha)
ITEMS = [("REINVERSION", "reinversion", "reinversion_pagado", "reinversion_fecha"), ("NOTA_CREDITO", "nota", "nota_aplicado", None),
         ("ACCES_TOO", "acces", "acces_aplicado", None), ("CAPACITACION", "capacitacion", "capacitacion_pagado", "capacitacion_fecha"),
         ("OTROS", "otros", None, None)]


class ErrorImportacion(Exception):
    pass


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().upper())


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return 0.0


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or ""))
    if m:  # en el Excel conviven dd/mm/aaaa y mm/dd/aaaa: se toma la primera fecha válida
        a, b, anio = int(m.group(1)), int(m.group(2)), int(m.group(3))
        for d, mes in ((a, b), (b, a)):
            try:
                return date(anio, mes, d)
            except ValueError:
                continue
    return None


def _hoja(wb):
    for ws in wb.worksheets:
        if _norm(ws.title) in HOJA_EXCLUIDA:
            continue
        if _norm(ws.cell(1, 1).value).startswith("NRO DE MFP"):
            return ws
    raise ErrorImportacion("No se encontró la hoja de MFP de Servigpoder (la columna A debe ser «NRO DE MFP FISICA»)")


def importar(db: Session, contenido: bytes, usuario_id: int | None) -> dict:
    wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True, read_only=True)
    ws = _hoja(wb)
    filas = list(ws.iter_rows(values_only=True))
    cols: dict[str, int] = {}
    for i, v in enumerate(filas[0]):
        if _norm(v) in COLUMNAS:
            cols.setdefault(COLUMNAS[_norm(v)], i)
    for campo in ("numero", "cliente", "inicio", "duracion", "valor_mensual"):
        if campo not in cols:
            raise ErrorImportacion(f"Falta la columna del Excel para «{campo}»")
    tipos = {t.codigo: t for t in db.scalars(select(CstTipoItem))}
    clientes = {(c.nit or "").strip(): c for c in db.scalars(select(CstCliente))}
    existentes = {(m.numero, m.cliente.nit, m.contrato, m.fecha_inicio) for m in db.scalars(select(CstMfp))}

    creadas, omitidas, advertencias = 0, 0, []
    vistos: dict[str, list[str]] = {}
    for n, fila in enumerate(filas[1:], start=2):
        val = lambda campo: fila[cols[campo]] if campo in cols and cols[campo] < len(fila) else None  # noqa: E731
        if not isinstance(val("numero"), (int, float)) or not str(val("cliente") or "").strip():
            continue
        numero = str(int(val("numero")))
        nit = str(val("nit") or "").strip()
        if isinstance(val("nit"), float):
            nit = str(int(val("nit")))
        nombre = str(val("cliente")).strip()
        inicio = _fecha(val("inicio"))
        duracion = int(_num(val("duracion"))) or None
        contrato = str(val("contrato") or "").strip()
        vistos.setdefault(numero, []).append(nombre)
        if not inicio or not duracion:
            advertencias.append(f"Fila {n} (MFP {numero}): sin fecha de inicio o duración")
        cliente = clientes.get(nit) if nit else None
        if cliente is None:
            cliente = CstCliente(nit=nit, nombre=nombre[:200])
            db.add(cliente)
            db.flush()
            if nit:
                clientes[nit] = cliente
        if (numero, cliente.nit, contrato, inicio) in existentes:
            omitidas += 1
            continue
        pct = val("pct_comercial")
        notas = " · ".join(str(x) for x in (val("otrosi"),) if x)
        m = CstMfp(numero=numero, cliente_id=cliente.id, contrato=contrato[:80], fecha_inicio=inicio, duracion_meses=duracion,
                   valor_mensual=_num(val("valor_mensual")), pct_comercial=pct if isinstance(pct, (int, float)) else None,
                   notas=notas[:2000], creado_por=usuario_id)
        db.add(m)
        db.flush()
        for codigo, c_pactado, c_ejecutado, c_fecha in ITEMS:
            pactado = _num(val(c_pactado))
            ejecutado = _num(val(c_ejecutado)) if c_ejecutado else 0.0
            if not pactado and not ejecutado:
                continue
            it = CstMfpItem(mfp_id=m.id, tipo_id=tipos[codigo].id, pactado=pactado, descripcion="Importado del Excel de MFP")
            db.add(it)
            db.flush()
            if ejecutado:
                fecha_txt = val(c_fecha) if c_fecha else None
                excede = ejecutado > pactado + 1
                db.add(CstMovimiento(item_id=it.id, fecha=_fecha(fecha_txt) or inicio or date.today(), valor=ejecutado,
                                     descripcion="Saldo ejecutado según el Excel de MFP",
                                     soporte=str(fecha_txt)[:300] if isinstance(fecha_txt, str) else "",
                                     excede=excede, autorizacion="Ya ejecutado antes de la plataforma (importado)" if excede else "",
                                     creado_por=usuario_id))
                if excede:
                    advertencias.append(f"Fila {n} (MFP {numero}): {tipos[codigo].nombre} ejecutado mayor que lo pactado")
        existentes.add((numero, cliente.nit, contrato, inicio))
        creadas += 1
    for numero, nombres in vistos.items():
        if len(nombres) > 1:
            advertencias.append(f"El número de MFP {numero} se repite en {len(nombres)} filas ({', '.join(sorted(set(nombres)))[:150]})")
    db.flush()
    return {"creadas": creadas, "omitidas": omitidas, "clientes": len(clientes), "advertencias": advertencias}
