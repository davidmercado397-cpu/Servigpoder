"""Carga del Excel de turnos de una quincena (mismo formato que la app original).

Fila 1: `documento` | nombre (o el mes, p. ej. MAYO) | [salario] [cargo] | 1 | 2 | … (o fechas completas)
Filas siguientes: documento, nombre y el código del turno de cada día (vacío = sin turno).
El archivo se procesa en memoria y no se guarda en el servidor.
"""

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.apps.liquidador.models import CARGO_DEFECTO, LiqDia, LiqEmpleado, LiqPeriodo
from app.apps.liquidador.services.turnos import por_codigo

_ALIAS = {
    "documento": "documento", "document_id": "documento",
    "nombre_completo": "nombre", "nombres_apellidos": "nombre", "nombre": "nombre", "full_name": "nombre",
    # El salario ya no se usa (solo se cuentan horas): la columna se reconoce para no leerla como día
    "salario": "salario", "salario_basico": "salario", "salario_base": "salario", "base_salary": "salario",
    "cargo": "cargo", "position": "cargo",
    # Solo en el costeo: clase de riesgo ARL de cada persona (en la nómina la ARL se configura en Empleados)
    "arl": "arl", "riesgo": "arl", "clase_riesgo": "arl", "riesgo_arl": "arl", "clase_arl": "arl", "clase_de_riesgo": "arl",
}
ROMANOS = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}
TARIFAS_ARL = {1: 0.522, 2: 1.044, 3: 2.436, 4: 4.35, 5: 6.96}  # % tarifa inicial por clase (Decreto 1772 de 1994)


def clase_arl(valor: object) -> int | None:
    """1-5, I-V, "clase 4" o la tarifa (4,35) → clase de riesgo; None si no se reconoce."""
    if valor is None or str(valor).strip() == "":
        return None
    texto = str(valor).strip().upper().replace("CLASE", "").replace("RIESGO", "").strip()
    if texto in ROMANOS:
        return ROMANOS[texto]
    try:
        numero = float(texto.replace("%", "").replace(",", "."))
    except ValueError:
        return None
    if numero.is_integer() and 1 <= numero <= 5:
        return int(numero)
    return next((c for c, t in TARIFAS_ARL.items() if abs(t - numero) < 0.001), None)
_MESES = ("", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
          "noviembre", "diciembre")


class ErrorImportacion(ValueError):
    pass


@dataclass
class FilaResumen:
    documento: str
    nombre: str
    dias: int
    turnos: str  # p. ej. "D×8, N×2"


@dataclass
class Resumen:
    empleados: int = 0
    dias: int = 0
    fechas: int = 0
    advertencias: list[str] = field(default_factory=list)
    detalle: list[FilaResumen] = field(default_factory=list)
    # Nómina por horas: turnos de otras líneas de la misma persona el mismo día (se suman sus horas)
    adicionales: dict[tuple[str, date], list] = field(default_factory=dict)


def _normalizar(valor: object) -> str:
    clave = str(valor).strip().lower().replace(" ", "_")
    return _ALIAS.get(clave, clave)


def _rango(desde: date, hasta: date) -> str:
    if (desde.year, desde.month) == (hasta.year, hasta.month):
        return f"{desde.day} al {hasta.day} de {_MESES[desde.month]} de {desde.year}"
    if desde.year == hasta.year:
        return f"{desde.day} de {_MESES[desde.month]} al {hasta.day} de {_MESES[hasta.month]} de {desde.year}"
    return f"{desde.day} de {_MESES[desde.month]} de {desde.year} al {hasta.day} de {_MESES[hasta.month]} de {hasta.year}"


def _fecha_encabezado(valor: object, periodo) -> date | None:
    """Número de día (1-31) del mes de la quincena o fecha completa."""
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    if isinstance(valor, (int, float)) or texto.isdigit():
        try:
            dia = int(float(valor)) if isinstance(valor, (int, float)) else int(texto)
            return date(periodo.anio, periodo.mes, dia) if 1 <= dia <= 31 else None
        except ValueError:
            return None  # día inexistente en el mes (p. ej. 30 de febrero)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


@dataclass
class FilaLeida:
    num: int
    documento: str
    nombre: str
    cargo: str
    dias: list[tuple[date, object]]  # (fecha, turno)
    arl: int | None = None  # solo en el costeo
    con_cargo: bool = False  # el Excel trae la columna cargo


def leer(db: Session, periodo, contenido: bytes, con_arl: bool = False) -> tuple[list[FilaLeida], Resumen]:
    """Lee el Excel sin escribir nada: filas con su turno por fecha, advertencias y conteos. Con `con_arl` (costeo)
    lee también la columna "arl"; en la nómina esa columna se ignora."""
    try:
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:  # archivo dañado: no se expone el detalle técnico
        raise ErrorImportacion("No se pudo leer el archivo Excel. Verifica que sea un .xlsx válido.") from e
    hoja = libro.active
    hoja.reset_dimensions()
    filas = list(hoja.iter_rows(values_only=True))
    libro.close()
    if len(filas) < 2:
        raise ErrorImportacion("El archivo no contiene filas de datos. Asegúrate de incluir al menos un empleado.")

    encabezado = list(filas[0])
    if not encabezado or encabezado[0] is None or _normalizar(encabezado[0]) != "documento":
        raise ErrorImportacion("La primera columna debe llamarse 'documento'. Revisa el encabezado del Excel.")

    columnas_fecha: list[tuple[int, date]] = []
    col_cargo: int | None = None
    col_arl: int | None = None
    for i, celda in enumerate(encabezado[2:], start=2):
        if celda is None:
            continue
        clave = _normalizar(celda)
        if clave == "salario":
            continue
        if clave == "cargo":
            col_cargo = i
            continue
        if clave == "arl":
            col_arl = i if con_arl else None
            continue
        fecha = _fecha_encabezado(celda, periodo)
        if fecha is not None:
            columnas_fecha.append((i, fecha))
    if not columnas_fecha:
        raise ErrorImportacion("No se encontraron columnas de fecha. A partir de la tercera columna, cada encabezado "
                               "debe ser un número de día (1 al 31) o una fecha completa.")

    if any(f < periodo.desde or f > periodo.hasta for _, f in columnas_fecha):
        fechas = [f for _, f in columnas_fecha]
        raise ErrorImportacion(
            f"El archivo cargado contiene fechas del {_rango(min(fechas), max(fechas))}, pero el periodo seleccionado es "
            f"{_rango(periodo.desde, periodo.hasta)}. Verifica que estás trabajando sobre el periodo correcto, o ajusta "
            f"los encabezados del Excel para que coincidan.")

    turnos = por_codigo(db)
    res = Resumen(fechas=len(columnas_fecha))
    leidas: list[FilaLeida] = []
    vistos: dict[tuple[str, date], object] = {}  # (documento, fecha) → turno de la primera línea
    # En la nómina por horas una persona puede tener varias líneas el mismo día (varios puestos): se suman las horas
    por_horas = getattr(periodo, "tipo_nomina", "normal") == "horas"
    con_horas = lambda t: (t.horas_ordinarias or 0) + (t.horas_extras or 0) > 0  # noqa: E731
    for num, fila in enumerate(filas[1:], start=2):
        if not fila or fila[0] is None or str(fila[0]).strip() == "":
            continue
        documento = _texto_documento(fila[0])
        nombre = str((fila[1] if len(fila) > 1 else None) or "").strip() or documento
        cargo = CARGO_DEFECTO
        if col_cargo is not None and col_cargo < len(fila) and fila[col_cargo]:
            cargo = str(fila[col_cargo]).strip() or CARGO_DEFECTO

        leida = FilaLeida(num, documento, nombre[:200], cargo[:80], [], con_cargo=col_cargo is not None)
        if col_arl is not None and col_arl < len(fila) and fila[col_arl] not in (None, ""):
            leida.arl = clase_arl(fila[col_arl])
            if leida.arl is None:
                res.advertencias.append(f"Fila {num}: la clase de riesgo ARL '{fila[col_arl]}' no es válida (1 a 5) — se omitió")
        leidas.append(leida)
        res.empleados += 1

        por_turno: dict[str, int] = {}
        for col, fecha in columnas_fecha:
            if col >= len(fila) or fila[col] is None or str(fila[col]).strip() == "":
                continue
            codigo = str(fila[col]).strip().upper()
            turno = turnos.get(codigo)
            if turno is None:
                res.advertencias.append(f"Fila {num}, {fecha.isoformat()}: el turno '{codigo}' no existe — se omitió")
                continue
            if (documento, fecha) in vistos:
                if por_horas and con_horas(turno) and con_horas(vistos[(documento, fecha)]):
                    res.adicionales.setdefault((documento, fecha), []).append(turno)
                    por_turno[codigo] = por_turno.get(codigo, 0) + 1
                    continue
                motivo = (" (en la nómina por horas solo se suman turnos con horas)" if por_horas else
                          " (dos líneas el mismo día solo se permiten en la nómina por horas)")
                res.advertencias.append(f"Fila {num}, {fecha.isoformat()}: el documento {documento} ya tiene turno ese día — se omitió{motivo}")
                continue
            vistos[(documento, fecha)] = turno
            leida.dias.append((fecha, turno))
            res.dias += 1
            por_turno[codigo] = por_turno.get(codigo, 0) + 1
        res.detalle.append(FilaResumen(documento, nombre, sum(por_turno.values()),
                                       ", ".join(f"{c}×{n}" for c, n in sorted(por_turno.items())) or "—"))

    res.empleados = len({f.documento for f in leidas})  # quien tiene varias líneas (nómina por horas) cuenta una vez
    if res.adicionales:
        res.advertencias.append(f"Nómina por horas: {len(res.adicionales)} día(s) de personas con más de una línea (varios puestos); "
                                f"se suman las horas de cada línea y el auxilio de transporte va una vez por día")
    if res.dias == 0:
        raise ErrorImportacion("No se importó ningún día de trabajo. Verifica que las celdas de cada día contengan un "
                               "código de turno válido (D, N, Z, L, AUS, INC…).")
    return leidas, res


def mismo_nombre(a: str, b: str) -> bool:
    """Dos nombres de la misma persona aunque cambie el orden o falte un apellido: comparten al menos dos palabras
    (o todas, si alguno tiene una sola)."""
    from app.apps.liquidador.services.personal import _norm

    pa, pb = set(_norm(a).split()), set(_norm(b).split())
    return len(pa & pb) >= min(2, len(pa), len(pb))


def importar(db: Session, periodo: LiqPeriodo, contenido: bytes) -> Resumen:
    """Reemplaza los días de la nómina y crea a quien no exista. A los empleados que ya existen no les cambia nada (nombre,
    cargo…): la ficha manda. Si el nombre del Excel no coincide con el de la ficha, lo advierte (cédula mal digitada)."""
    leidas, res = leer(db, periodo, contenido)
    from app.apps.liquidador.services.personal import Indice

    empleados = Indice(db.scalars(select(LiqEmpleado)))
    db.execute(delete(LiqDia).where(LiqDia.periodo_id == periodo.id))
    for f in leidas:
        emp = empleados.buscar(f.documento)
        if emp is None:
            emp = LiqEmpleado(documento=f.documento, nombre=f.nombre, cargo=f.cargo)
            db.add(emp)
            empleados[f.documento] = emp
        elif f.nombre != f.documento and not f.nombre.isdigit() and not mismo_nombre(f.nombre, emp.nombre):
            res.advertencias.append(f"Fila {f.num}: la cédula {f.documento} es de {emp.nombre} en la ficha, pero el Excel dice "
                                    f"«{f.nombre}». Se cargó a {emp.nombre}; si la cédula está mal digitada, corríjala y vuelva a cargar")
        db.flush()
        for fecha, turno in f.dias:
            extra = res.adicionales.get((f.documento, fecha), [])
            db.add(LiqDia(periodo_id=periodo.id, empleado_id=emp.id, fecha=fecha, turno_id=turno.id, turnos_extra=[t.id for t in extra],
                          codigo="+".join([turno.codigo, *(t.codigo for t in extra)])[:40]))
    db.flush()
    return res


def _texto_documento(valor: object) -> str:
    """Cédulas que Excel guarda como número (1144127973.0) se leen sin decimales."""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()
