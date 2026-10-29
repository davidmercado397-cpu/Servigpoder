"""Comparación con la nómina de Siigo: el reporte de netos (una hoja por nómina con NOMBRE, CEDULA, INGRESOS, DEDUCCIONES y
NETO A PAGAR) contra lo liquidado aquí en todas las nóminas del mismo periodo. Se cruza persona por persona sumando
sus nóminas en cada lado, para que no importe cómo se repartieron. El archivo se procesa y se descarta.

Las posibles causas son pistas: se busca qué concepto liquidado aquí explica la diferencia (recargos, extras, auxilio de
transporte, auxilios no prestacionales, vacaciones, incapacidad, novedades, días de salario, descuentos).
"""

from collections import defaultdict
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.apps.liquidador.models import LiqPeriodo, LiqResultado
from app.apps.liquidador.services.exportar import XLSX, _CENTRO, _FUENTE, _RELLENO, _autoajustar, _guardar
from app.apps.liquidador.services.personal import ErrorArchivo, Indice, _abrir, _norm, _texto, documento

CAMPOS_SISTEMA = ("basico", "incapacidad", "vacaciones_valor", "recargos", "extras", "auxilios_no_prestacionales", "auxilio",
                  "novedades_devengo", "devengado", "salud", "pension", "embargos", "prestamos", "cuotas", "novedades_deduccion",
                  "deducciones", "neto")
# Conceptos que pueden explicar una diferencia de ingresos o de deducciones
PISTAS_INGRESOS = (("recargos", "los recargos"), ("extras", "las horas extras"), ("auxilio", "el auxilio de transporte"),
                   ("auxilios_no_prestacionales", "los auxilios no prestacionales"), ("vacaciones_valor", "las vacaciones"),
                   ("incapacidad", "la incapacidad"), ("novedades_devengo", "las novedades (devengo)"))
PISTAS_DEDUCCIONES = (("prestamos", "los préstamos"), ("cuotas", "las cuotas"), ("embargos", "los embargos"),
                      ("novedades_deduccion", "las novedades (deducción)"))


@dataclass
class PersonaSiigo:
    documento: str
    nombre: str
    hojas: list[str] = field(default_factory=list)
    ingresos: float = 0.0
    deducciones: float = 0.0
    neto: float = 0.0


def _num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("$", "").replace(".", "").replace(",", ".").strip()) if v not in (None, "") else None
    except ValueError:
        return None


def leer(contenido: bytes) -> tuple[dict[str, PersonaSiigo], list[dict], dict[str, float]]:
    """(personas por documento, hallazgos, neto total de cada hoja)."""
    libro = _abrir(contenido)
    personas: dict[str, PersonaSiigo] = {}
    hallazgos, por_hoja = [], {}
    for hoja in libro.worksheets:
        cols: dict[str, int] = {}
        for num, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            if not cols:
                claves = [_norm(c) for c in fila]
                if "CEDULA" in claves and any(c.startswith("NETO") for c in claves):
                    for i, c in enumerate(claves):
                        for prefijo, campo in (("NOMBRE", "nombre"), ("CEDULA", "documento"), ("INGRESO", "ingresos"), ("DEVENGADO", "ingresos"),
                                               ("DEDUCCION", "deducciones"), ("NETO", "neto")):
                            if c.startswith(prefijo) and campo not in cols:
                                cols[campo] = i
                continue
            v = lambda campo: fila[cols[campo]] if campo in cols and cols[campo] < len(fila) else None  # noqa: E731
            doc, _ = documento(v("documento"))
            if doc is None:
                continue  # filas de total o vacías
            neto = _num(v("neto"))
            if neto is None:
                hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": "Sin neto a pagar: no se comparó"})
                continue
            p = personas.setdefault(doc, PersonaSiigo(doc, _texto(v("nombre"), 200) or doc))
            p.hojas.append(hoja.title)
            p.ingresos += _num(v("ingresos")) or 0.0
            p.deducciones += _num(v("deducciones")) or 0.0
            p.neto += neto
            por_hoja[hoja.title] = por_hoja.get(hoja.title, 0.0) + neto
    if not personas:
        raise ErrorArchivo("No se encontraron filas con CEDULA y NETO A PAGAR. Use el reporte de netos de nómina de Siigo.")
    return personas, hallazgos, por_hoja


def _pistas(dif: float, componentes: dict[str, float], lista, tolerancia: float, verbo: str) -> list[str]:
    """Conceptos de aquí cuyo valor coincide con la diferencia (solo si aquí hay más: de Siigo no hay detalle)."""
    if dif <= 0:
        return []
    margen = max(tolerancia, 100.0)
    return [f"Aquí se {verbo} {texto} ({valor:,.0f}) y la diferencia es igual: revise si Siigo los tiene"
            for campo, texto in lista if (valor := componentes.get(campo, 0.0)) and abs(dif - valor) <= margen]


def comparar(db: Session, anio: int, mes: int, quincena: int, contenido: bytes, tolerancia: float = 1.0) -> dict:
    siigo, hallazgos, por_hoja = leer(contenido)
    nominas = list(db.scalars(select(LiqPeriodo).where(LiqPeriodo.anio == anio, LiqPeriodo.mes == mes, LiqPeriodo.quincena == quincena)
                              .order_by(LiqPeriodo.id)))
    if not nominas:
        raise ErrorArchivo("No hay nóminas liquidadas en ese periodo")
    nuestras: dict[str, dict] = {}
    empleados = {}
    por_nomina = {}
    for p in nominas:
        total = 0.0
        for r in db.scalars(select(LiqResultado).where(LiqResultado.periodo_id == p.id)):
            n = r.nomina or {}
            if not n:
                continue
            e = r.empleado
            empleados[e.documento] = e
            x = nuestras.setdefault(e.documento, {"nombre": e.nombre, "empleado_id": e.id, "nominas": [], **{c: 0.0 for c in CAMPOS_SISTEMA},
                                                  "salario": float(n.get("salario", 0))})
            x["nominas"].append({"periodo_id": p.id, "nombre": p.nombre, "estado": p.estado})
            for c in CAMPOS_SISTEMA:
                x[c] += float(n.get(c, 0) or 0)
            total += float(n.get("neto", 0) or 0)
        por_nomina[p.nombre] = total
    indice = Indice(empleados.values())

    filas, vistos = [], set()
    for doc, s in siigo.items():
        e = indice.buscar(doc)
        propio = nuestras.get(e.documento) if e else None
        if propio is None:
            filas.append({"documento": doc, "nombre": s.nombre, "estado": "solo_siigo", "hojas_siigo": s.hojas, "nominas": [],
                          "siigo": {"ingresos": s.ingresos, "deducciones": s.deducciones, "neto": s.neto}, "sistema": None,
                          "diferencia": {"ingresos": -s.ingresos, "deducciones": -s.deducciones, "neto": -s.neto},
                          "pistas": ["No está liquidado aquí en este periodo"]})
            continue
        vistos.add(e.documento)
        dif = {"ingresos": propio["devengado"] - s.ingresos, "deducciones": propio["deducciones"] - s.deducciones,
               "neto": propio["neto"] - s.neto}
        cuadra = all(abs(v) <= tolerancia for v in dif.values())
        pistas = []
        if not cuadra:
            if abs(dif["ingresos"]) > tolerancia:
                pistas += _pistas(dif["ingresos"], propio, PISTAS_INGRESOS, tolerancia, "pagan")
                valor_dia = propio["salario"] / 30 if propio["salario"] else 0
                if valor_dia and not pistas:
                    dias = dif["ingresos"] / valor_dia
                    if abs(dias - round(dias)) < 0.05 and round(dias):
                        pistas.append(f"La diferencia de ingresos equivale a {abs(round(dias))} día(s) de salario "
                                      f"({'más' if dias > 0 else 'menos'} aquí)")
            if abs(dif["deducciones"]) > tolerancia:
                dp = _pistas(dif["deducciones"], propio, PISTAS_DEDUCCIONES, tolerancia, "descuentan")
                if not dp and abs(dif["ingresos"]) > tolerancia and abs(dif["deducciones"] - 0.08 * dif["ingresos"]) <= max(tolerancia, 100):
                    dp.append("La diferencia de deducciones es el 8 % (salud y pensión) de la diferencia de ingresos")
                pistas += dp
        filas.append({"documento": e.documento, "nombre": propio["nombre"], "empleado_id": propio["empleado_id"],
                      "estado": "cuadra" if cuadra else "diferencia", "hojas_siigo": s.hojas, "nominas": propio["nominas"],
                      "siigo": {"ingresos": s.ingresos, "deducciones": s.deducciones, "neto": s.neto},
                      "sistema": {c: round(propio[c], 2) for c in CAMPOS_SISTEMA}, "diferencia": {k: round(v, 2) for k, v in dif.items()},
                      "pistas": pistas})
    for doc, propio in nuestras.items():
        if doc not in vistos:
            filas.append({"documento": doc, "nombre": propio["nombre"], "empleado_id": propio["empleado_id"], "estado": "solo_sistema",
                          "hojas_siigo": [], "nominas": propio["nominas"], "siigo": None,
                          "sistema": {c: round(propio[c], 2) for c in CAMPOS_SISTEMA},
                          "diferencia": {"ingresos": propio["devengado"], "deducciones": propio["deducciones"], "neto": propio["neto"]},
                          "pistas": ["No está en el archivo de Siigo"]})
    orden = {"diferencia": 0, "solo_siigo": 1, "solo_sistema": 2, "cuadra": 3}
    filas.sort(key=lambda f: (orden[f["estado"]], -abs(f["diferencia"]["neto"]), f["nombre"]))
    neto_siigo = sum(s.neto for s in siigo.values())
    neto_sistema = sum(x["neto"] for x in nuestras.values())
    conteo = defaultdict(int)
    for f in filas:
        conteo[f["estado"]] += 1
    return {
        "resumen": {"personas_siigo": len(siigo), "personas_sistema": len(nuestras), "cuadran": conteo["cuadra"],
                    "con_diferencia": conteo["diferencia"], "solo_siigo": conteo["solo_siigo"], "solo_sistema": conteo["solo_sistema"],
                    "neto_siigo": round(neto_siigo, 2), "neto_sistema": round(neto_sistema, 2),
                    "diferencia_neto": round(neto_sistema - neto_siigo, 2), "tolerancia": tolerancia},
        "hojas_siigo": {k: round(v, 2) for k, v in por_hoja.items()},
        "nominas_sistema": {k: round(v, 2) for k, v in por_nomina.items()},
        "filas": filas, "hallazgos": hallazgos,
    }


ESTADOS = {"cuadra": "Cuadra", "diferencia": "Con diferencia", "solo_siigo": "Solo en Siigo", "solo_sistema": "Solo en el sistema"}


def excel(res: dict, etiqueta: str) -> tuple[bytes, str, str]:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Comparacion"
    hoja.append(["DOCUMENTO", "EMPLEADO", "ESTADO", "HOJAS SIIGO", "NOMINAS SISTEMA", "INGRESOS SIIGO", "INGRESOS SISTEMA", "DIF INGRESOS",
                 "DEDUCCIONES SIIGO", "DEDUCCIONES SISTEMA", "DIF DEDUCCIONES", "NETO SIIGO", "NETO SISTEMA", "DIF NETO",
                 "BASICO", "INCAPACIDAD", "VACACIONES", "RECARGOS", "EXTRAS", "AUX. NO PRESTACIONALES", "AUX. TRANSPORTE",
                 "NOVEDADES DEVENGO", "SALUD", "PENSION", "EMBARGOS", "PRESTAMOS", "CUOTAS", "NOVEDADES DEDUCCION", "POSIBLES CAUSAS"])
    for f in res["filas"]:
        s, n, d = f["siigo"] or {}, f["sistema"] or {}, f["diferencia"]
        hoja.append([f["documento"], f["nombre"], ESTADOS[f["estado"]], " · ".join(f["hojas_siigo"]),
                     " · ".join(x["nombre"] for x in f["nominas"]),
                     s.get("ingresos"), n.get("devengado"), d["ingresos"], s.get("deducciones"), n.get("deducciones"), d["deducciones"],
                     s.get("neto"), n.get("neto"), d["neto"],
                     *(n.get(c) for c in ("basico", "incapacidad", "vacaciones_valor", "recargos", "extras", "auxilios_no_prestacionales",
                                          "auxilio", "novedades_devengo", "salud", "pension", "embargos", "prestamos", "cuotas",
                                          "novedades_deduccion")),
                     " | ".join(f["pistas"])])
    for celda in hoja[1]:
        celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
    for col in range(6, 29):
        for (celda,) in hoja.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = "$ #,##0.00"
    _autoajustar(hoja, minimo=10, maximo=40)
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.row_dimensions[1].height = 42

    r = res["resumen"]
    resumen = libro.create_sheet("Resumen")
    for fila in (["Personas en Siigo", r["personas_siigo"]], ["Personas en el sistema", r["personas_sistema"]], ["Cuadran", r["cuadran"]],
                 ["Con diferencia", r["con_diferencia"]], ["Solo en Siigo", r["solo_siigo"]], ["Solo en el sistema", r["solo_sistema"]],
                 ["Neto Siigo", r["neto_siigo"]], ["Neto sistema", r["neto_sistema"]], ["Diferencia (sistema − Siigo)", r["diferencia_neto"]],
                 ["Tolerancia por persona", r["tolerancia"]], [], ["NETO POR HOJA DE SIIGO", ""],
                 *([k, v] for k, v in res["hojas_siigo"].items()), [], ["NETO POR NOMINA DEL SISTEMA", ""],
                 *([k, v] for k, v in res["nominas_sistema"].items())):
        resumen.append(fila)
    for celda in resumen["A"]:
        celda.font = Font(bold=True)
    _autoajustar(resumen, minimo=12, maximo=45)
    return _guardar(libro), XLSX, f"comparacion_siigo_{etiqueta}.xlsx"
