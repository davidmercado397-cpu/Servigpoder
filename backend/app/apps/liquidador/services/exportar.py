"""Archivos que se descargan de una quincena.

- Liquidación ("Calendario + Liquidación" en la app original): el archivo que se sube al sistema de
  nómina. Conserva las 19 columnas y su orden; DIAS CON NOVEDAD se agrega al final.
- Plantilla: el calendario vacío que se llena y se vuelve a cargar.
"""

import io
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.apps.liquidador.dominio.tipos import HourConcept
from app.apps.liquidador.models import LiqPeriodo, LiqResultado
from app.apps.liquidador.services.calculo import AUSENCIA, DESCANSO_PAGO, INCAPACIDAD, LIBRE, NOVEDAD, TRABAJADO

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Mismo orden de la plantilla de referencia `Libro1.xlsx`
ENCABEZADOS: tuple[str, ...] = (
    "DOCUMENTO", "EMPLEADO",
    "DIURNAS ORDINARIA", "RECARGO NOCTURNO", "RECARGO FESTIVO DIURNO", "RECARGO FESTIVO NOCTURNO",
    "RECARGO DOMINICAL DIURNO", "RECARGO DOMINICAL NOCTURNO", "EXTRAS ORDINARIA DIURNAS", "EXTRAS ORDINARIA NOCTURNAS",
    "EXTRAS FESTIVAS DIURNAS", "EXTRAS FESTIVAS NOCTURNAS", "EXTRAS DOMINICALES DIURNAS", "EXTRAS DOMINICALES NOCTURNAS",
    "DIAS TRABAJADOS", "DESCANSOS PAGOS", "LIBRES", "AUSENCIAS", "INCAPACIDADES",
    "DIAS CON NOVEDAD",
)
ORDEN_CONCEPTOS: tuple[HourConcept, ...] = (
    HourConcept.ORDINARY_DAY, HourConcept.ORDINARY_NIGHT, HourConcept.HOLIDAY_DAY_SURCHARGE,
    HourConcept.HOLIDAY_NIGHT_SURCHARGE, HourConcept.SUNDAY_SURCHARGE, HourConcept.SUNDAY_NIGHT_SURCHARGE,
    HourConcept.OVERTIME_DAY, HourConcept.OVERTIME_NIGHT, HourConcept.HOLIDAY_OVERTIME_DAY,
    HourConcept.HOLIDAY_OVERTIME_NIGHT, HourConcept.SUNDAY_OVERTIME_DAY, HourConcept.SUNDAY_OVERTIME_NIGHT,
)
ORDEN_DIAS = (TRABAJADO, DESCANSO_PAGO, LIBRE, AUSENCIA, INCAPACIDAD, NOVEDAD)

_RELLENO = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
_FUENTE = Font(bold=True, color="FFFFFF", size=10)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DERECHA = Alignment(horizontal="right", vertical="center")
_MESES = ("", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
          "NOVIEMBRE", "DICIEMBRE")


def _titulo(p: LiqPeriodo) -> str:
    return f"{p.anio}-{p.mes:02d}-Q{p.quincena}"


def fila_liquidacion(r: LiqResultado) -> list:
    return ([r.empleado.documento, r.empleado.nombre]
            + [float(r.horas.get(c.value, 0)) for c in ORDEN_CONCEPTOS]
            + [int(r.dias.get(k, 0)) for k in ORDEN_DIAS])


def liquidacion(db: Session, periodo: LiqPeriodo) -> tuple[bytes, str, str]:
    # Orden por documento como texto (igual que la app original), sin depender de la intercalación de la base
    resultados = sorted(db.scalars(select(LiqResultado).where(LiqResultado.periodo_id == periodo.id)),
                        key=lambda r: r.empleado.documento)
    libro = Workbook()
    hoja = libro.active
    hoja.title = _titulo(periodo)
    hoja.append(list(ENCABEZADOS))
    for celda in hoja[1]:
        celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
    for r in resultados:
        hoja.append(fila_liquidacion(r))

    # Horas con un decimal (columnas 3 a 14); conteos como enteros
    for col in range(1, len(ENCABEZADOS) + 1):
        formato = "#,##0.0" if 3 <= col <= 14 else "#,##0"
        for (celda,) in hoja.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format, celda.alignment = formato, _DERECHA
    _autoajustar(hoja, minimo=10, maximo=22)
    hoja.freeze_panes = "C2"
    hoja.row_dimensions[1].height = 32
    return _guardar(libro), XLSX, f"liquidacion_{periodo.anio}_{periodo.mes:02d}_Q{periodo.quincena}.xlsx"


def plantilla(periodo: LiqPeriodo) -> tuple[bytes, str, str]:
    """documento | <MES AÑO> | día 1 | día 2 | … con una fila vacía para llenar."""
    dias, actual = [], periodo.desde
    while actual <= periodo.hasta:
        dias.append(actual.day)
        actual += timedelta(days=1)
    libro = Workbook()
    hoja = libro.active
    hoja.title = _titulo(periodo)
    encabezado = ["documento", f"{_MESES[periodo.mes]} {periodo.anio}", *dias]
    hoja.append(encabezado)
    for celda in hoja[1]:
        celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
    hoja.append([""] * len(encabezado))
    hoja.column_dimensions["A"].width = 16
    hoja.column_dimensions["B"].width = 28
    for i in range(len(dias)):
        hoja.column_dimensions[get_column_letter(3 + i)].width = 5
    hoja.freeze_panes = "C2"
    hoja.row_dimensions[1].height = 26
    return _guardar(libro), XLSX, f"plantilla_{periodo.anio}_{periodo.mes:02d}_Q{periodo.quincena}.xlsx"


def _autoajustar(hoja, *, minimo: int, maximo: int, margen: int = 2) -> None:
    for celdas in hoja.columns:
        largo = max((len(linea) for c in celdas if c.value is not None for linea in str(c.value).split("\n")), default=0)
        hoja.column_dimensions[celdas[0].column_letter].width = max(minimo, min(maximo, largo + margen))


def _guardar(libro: Workbook) -> bytes:
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()
