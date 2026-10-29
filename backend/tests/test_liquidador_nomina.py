"""Nómina del liquidador. Valores calculados a mano con el mínimo 2026 ($1.750.905), auxilio $249.095 y 210
horas al mes: valor día = 58.363,5 · valor hora = 8.337,642857…"""

from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import ACCESOS, entrar, xlsx
from tests.test_liquidador import API, _cargar, _excel_sep_q2, _quincena, _resultados

HORA = 1750905 / 210


def _periodo_con(cliente, anio, mes, quincena, filas, encabezado_dias) -> int:
    pid = _quincena(cliente, anio, mes, quincena)
    r = _cargar(cliente, pid, xlsx([["documento", "MES", *encabezado_dias], *filas]))
    assert r.status_code == 200, r.text
    return pid


def _nomina(cliente, pid, doc) -> dict:
    res = _resultados(cliente, pid)[doc]
    return cliente.get(f"{API}/periodos/{pid}/empleados/{res['empleado_id']}").json()["data"]["empleado"]["nomina"]


def test_tarifas_de_ley_por_vigencia(admin):
    r = admin.get(f"{API}/tarifas").json()
    t = {x["vigente_desde"]: x for x in r["data"]}
    assert set(t) == {"2026-01-01", "2026-07-01"}
    julio = t["2026-07-01"]
    assert (julio["smlmv"], julio["auxilio_transporte"], julio["horas_mes"], julio["salud_pct"], julio["pension_pct"]) == (1750905, 249095, 210, 4, 4)
    assert julio["porcentajes"] == {
        "ordinary_night": 35, "holiday_day_surcharge": 90, "sunday_surcharge": 90, "holiday_night_surcharge": 125,
        "sunday_night_surcharge": 125, "overtime_day": 125, "overtime_night": 175, "holiday_overtime_day": 215,
        "sunday_overtime_day": 215, "holiday_overtime_night": 265, "sunday_overtime_night": 265}
    assert t["2026-01-01"]["porcentajes"]["sunday_surcharge"] == 80 and t["2026-01-01"]["porcentajes"]["sunday_overtime_night"] == 255


def test_nomina_de_una_quincena(admin):
    pid = _quincena(admin)
    _cargar(admin, pid, _excel_sep_q2())
    # 1001: 13 turnos A (7 ordinarias + 1 extra diurna) y 2 descansos Z
    n = _nomina(admin, pid, "1001")
    assert n["dias_salario"] == 15 and n["dias_auxilio"] == 15 and n["salario_minimo"]
    assert n["basico"] == 875453  # 15 × 58.363,5
    assert n["valores"] == {"overtime_day": round(13 * HORA * 1.25)} == {"overtime_day": 135487}
    assert n["auxilio"] == 124548  # 15 × 249.095 / 30
    assert n["devengado"] == 875453 + 135487 + 124548 == 1135488
    assert n["ibc"] == 1010940 and n["salud"] == 40438 and n["pension"] == 40438
    assert n["neto"] == 1135488 - 80876 == 1054612

    # 3003: LR (otro proceso) y AUS no se pagan; V×5 como vacaciones; L sí; INC al 100 %; el auxilio solo por los 4 días A
    n = _nomina(admin, pid, "3003")
    assert n["dias_salario"] == 5 and n["dias_incapacidad"] == 2 and n["dias_sin_pago"] == 2 and n["dias_auxilio"] == 4
    assert n["vacaciones"]["dias_habiles"] == 4 and n["vacaciones"]["dias_no_habiles"] == 1 and n["vacaciones_valor"] == 291818
    assert n["basico"] == 291818 and n["incapacidad"] == 116727 and n["auxilio"] == 33213
    assert (n["incapacidad_empresa"], n["incapacidad_eps"]) == (2, 0)

    lista = admin.get(f"{API}/periodos/{pid}/resultados?orden=neto").json()
    netos = [x["neto"] for x in lista["data"]]
    assert netos == sorted(netos, reverse=True) and 1054612 in netos
    assert lista["meta"]["extra"]["totales"]["nomina"]["neto"] == sum(x["neto"] for x in lista["data"])


def test_incapacidad_larga_y_base_30(admin):
    # Octubre 2026 mensual (31 días): 31 turnos A suman 30 días de salario; INC del 1 al 5 → 2 empresa + 3 EPS
    dias = list(range(1, 32))
    pid = _periodo_con(admin, 2026, 10, 0, [[1, "1", *(["A"] * 31)], [2, "2", *(["INC"] * 5 + ["A"] * 26)]], dias)
    a, b = _nomina(admin, pid, "1"), _nomina(admin, pid, "2")
    assert a["dias_salario"] == 30 and a["basico"] == 1750905 and a["auxilio"] == 249095
    assert b["dias_incapacidad"] == 5 and (b["incapacidad_empresa"], b["incapacidad_eps"]) == (2, 3)
    assert b["incapacidad"] == round(5 * 1750905 / 30) and b["dias_salario"] == 25  # el 31 no suma
    # Febrero 2027 mensual: trabajar hasta el 28 completa los 30 días
    pid = _periodo_con(admin, 2027, 2, 0, [[1, "1", *(["A"] * 28)]], list(range(1, 29)))
    assert _nomina(admin, pid, "1")["dias_salario"] == 30 and _nomina(admin, pid, "1")["basico"] == 1750905


def test_cada_dia_usa_la_tarifa_vigente(admin):
    # Domingo 21-jun-2026 (dominical 80 %) y domingo 5-jul-2026 (90 %), turno D: 8,4 ordinarias + 2,8 extras
    jun = _periodo_con(admin, 2026, 6, 2, [[1, "1", "D"]], [21])
    jul = _periodo_con(admin, 2026, 7, 1, [[1, "1", "D"]], [5])
    assert _nomina(admin, jun, "1")["valores"] == {"sunday_surcharge": round(8.4 * HORA * 0.80), "sunday_overtime_day": round(2.8 * HORA * 2.05)}
    assert _nomina(admin, jul, "1")["valores"] == {"sunday_surcharge": round(8.4 * HORA * 0.90), "sunday_overtime_day": round(2.8 * HORA * 2.15)}

    # Cambiar un % aplica al recalcular
    t = next(x for x in admin.get(f"{API}/tarifas").json()["data"] if x["vigente_desde"] == "2026-07-01")
    datos = {k: t[k] for k in ("vigente_desde", "smlmv", "auxilio_transporte", "horas_mes", "salud_pct", "pension_pct", "porcentajes", "nota")}
    datos["porcentajes"] = {**t["porcentajes"], "sunday_surcharge": 100}
    assert admin.put(f"{API}/tarifas/{t['id']}", json=datos).status_code == 200
    admin.post(f"{API}/periodos/{jul}/recalcular")
    assert _nomina(admin, jul, "1")["valores"]["sunday_surcharge"] == round(8.4 * HORA)
    # Validaciones
    assert admin.put(f"{API}/tarifas/{t['id']}", json={**datos, "porcentajes": {"overtime_day": 125}}).status_code == 422
    assert admin.post(f"{API}/tarifas", json=datos).status_code == 409  # misma fecha


def test_salario_propio_y_auxilio(admin):
    pid = _quincena(admin)
    _cargar(admin, pid, _excel_sep_q2())
    eid = _resultados(admin, pid)["1001"]["empleado_id"]
    assert admin.put(f"{API}/empleados/{eid}/salario", json={"salario": 3000000}).status_code == 200
    admin.post(f"{API}/periodos/{pid}/recalcular")
    n = _nomina(admin, pid, "1001")
    assert n["basico"] == 1500000 and n["auxilio"] == 124548 and not n["salario_minimo"]  # 3 M ≤ 2 mínimos
    admin.put(f"{API}/empleados/{eid}/salario", json={"salario": 4000000})
    admin.post(f"{API}/periodos/{pid}/recalcular")
    assert _nomina(admin, pid, "1001")["auxilio"] == 0  # más de 2 mínimos
    admin.put(f"{API}/empleados/{eid}/salario", json={"salario": None})
    e = admin.get(f"{API}/empleados?q=1001").json()["data"][0]
    assert e["salario"] is None


def test_prestamo_con_saldo_y_embargo(admin):
    dias = list(range(1, 16))
    filas = [[1001, "1001", *(["A"] * 15)]]
    p1 = _periodo_con(admin, 2026, 11, 1, filas, dias)
    r = admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "prestamo", "descripcion": "Préstamo celular",
                                              "valor_mensual": 100000, "monto_total": 120000, "desde": "2026-11-01"})
    assert r.status_code == 201, r.text
    prestamo = r.json()["data"]["id"]
    assert admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "embargo", "descripcion": "Juzgado 3",
                                                 "porcentaje": 10, "desde": "2026-11-01"}).status_code == 201
    assert admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "embargo", "descripcion": "x",
                                                 "porcentaje": 10, "valor_mensual": 5, "desde": "2026-11-01"}).status_code == 422
    assert admin.post(f"{API}/descuentos", json={"documento": "999", "tipo": "embargo", "descripcion": "x",
                                                 "porcentaje": 10, "desde": "2026-11-01"}).status_code == 404

    admin.post(f"{API}/periodos/{p1}/recalcular")
    n = _nomina(admin, p1, "1001")
    embargo, cuota = n["descuentos"]
    assert embargo["tipo"] == "embargo" and embargo["valor"] == round(n["ibc"] * 0.10)
    assert cuota["valor"] == 50000  # la mitad de la cuota mensual en la quincena
    assert n["neto"] == n["devengado"] - n["salud"] - n["pension"] - embargo["valor"] - 50000

    p2 = _periodo_con(admin, 2026, 11, 2, filas, list(range(16, 31)))
    assert next(x for x in _nomina(admin, p2, "1001")["descuentos"] if x["id"] == prestamo)["valor"] == 50000
    p3 = _periodo_con(admin, 2026, 12, 1, filas, dias)
    x = next(x for x in _nomina(admin, p3, "1001")["descuentos"] if x["id"] == prestamo)
    assert x["valor"] == 20000 and "saldo" in x["observacion"]  # solo el saldo pendiente
    d = next(x for x in admin.get(f"{API}/descuentos?q=celular").json()["data"])
    assert d["descontado"] == 120000 and d["saldo"] == 0
    assert admin.delete(f"{API}/descuentos/{prestamo}").status_code == 409  # ya aplicado: se desactiva

    r = admin.get(f"{API}/periodos/{p1}/nomina")
    assert r.headers["content-disposition"] == 'attachment; filename="nomina_2026_11_Q1.xlsx"'
    libro = load_workbook(BytesIO(r.content))
    assert libro.sheetnames == ["Nomina 2026-11-Q1", "Prestamos embargos y cuotas", "Novedades", "Vacaciones", "Auxilios no prestacionales", "Tarifas usadas"]
    filas_x = list(libro.worksheets[0].iter_rows(values_only=True))
    enc = filas_x[0]
    fila = dict(zip(enc, filas_x[1]))
    assert fila["DOCUMENTO"] == "1001" and fila["NETO A PAGAR"] == n["neto"] and fila["TOTAL DEVENGADO"] == n["devengado"]
    assert "VR EXTRAS ORDINARIA DIURNAS" in enc and "VR DIURNAS ORDINARIA" not in enc
    assert filas_x[-1][0] == "TOTAL"


def test_permisos_de_nomina(admin):
    pid = _quincena(admin)
    _cargar(admin, pid, _excel_sep_q2())
    admin.post("/api/roles", json={"nombre": "Solo horas", "permisos": ["liquidador.periodos.ver"]})
    rid = next(r["id"] for r in admin.get("/api/roles").json()["data"] if r["nombre"] == "Solo horas")
    admin.post("/api/usuarios", json={"accesos": ACCESOS, "username": "horas", "nombre": "Horas", "password": "clave-segura-1", "roles": [rid]})
    admin.post("/api/auth/logout")
    entrar(admin, "horas", "clave-segura-1")
    fila = admin.get(f"{API}/periodos/{pid}/resultados").json()
    assert fila["data"][0]["neto"] is None and "nomina" not in fila["meta"]["extra"]["totales"]
    assert _nomina(admin, pid, "1001") is None
    assert admin.get(f"{API}/tarifas").status_code == 403
    assert admin.get(f"{API}/descuentos").status_code == 403
    assert admin.get(f"{API}/periodos/{pid}/nomina").status_code == 403


def test_auxilios_no_prestacionales_y_quincena(admin):
    dias = list(range(1, 16))
    # 1001: 15 días A (7 + 1 extra diurna); 2002: 13 días A y 2 ausencias
    filas = [[1001, "1001", *(["A"] * 15)], [2002, "2002", *(["A"] * 13), "AUS", "AUS"], [3003, "3003", *(["A"] * 15)]]
    p1 = _periodo_con(admin, 2026, 12, 1, filas, dias)
    base = {"valor_mensual": 300000, "desde": "2026-12-01"}
    for doc in ("1001", "2002"):
        r = admin.post(f"{API}/auxilios", json={**base, "documento": doc, "descripcion": "Auxilio de rodamiento"})
        assert r.status_code == 201, r.text
    assert admin.post(f"{API}/auxilios", json={**base, "documento": "3003", "descripcion": "Bonificación", "valor_mensual": 2000000}).status_code == 201
    # Préstamo solo en la 2.ª quincena (valor completo en esa) y auxilio solo en la 1.ª
    assert admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "prestamo", "descripcion": "Celular",
                                                 "valor_mensual": 80000, "desde": "2026-12-01", "quincena": "segunda"}).status_code == 201
    assert admin.post(f"{API}/auxilios", json={**base, "documento": "1001", "descripcion": "Alimentación", "valor_mensual": 100000,
                                               "quincena": "primera"}).status_code == 201
    assert admin.post(f"{API}/auxilios", json={**base, "documento": "1001", "descripcion": "x", "quincena": "tercera"}).status_code == 422
    admin.post(f"{API}/periodos/{p1}/recalcular")

    n = _nomina(admin, p1, "1001")
    # Rodamiento: mitad del mes en cada quincena (150.000); alimentación: completa en la 1.ª (100.000)
    assert sorted(x["valor"] for x in n["auxilios"]) == [100000, 150000] and n["auxilios_no_prestacionales"] == 250000
    horas = sum(n["valores"].values())  # extras y recargos (hay domingos y el festivo del 8)
    assert n["devengado"] == n["basico"] + horas + 250000 + n["auxilio"]
    # 250.000 no pasa del 40 % de la remuneración: la base de salud y pensión es solo lo salarial
    salarial = n["basico"] + horas
    assert n["exceso_no_salarial"] == 0 and n["ibc"] == salarial and not n["alertas"]
    # 1.000.000 (la mitad de 2.000.000) sí pasa del 40 %: el exceso entra en la base
    n3 = _nomina(admin, p1, "3003")
    sal3 = n3["basico"] + sum(n3["valores"].values())
    exceso = round(1000000 - 0.4 * (sal3 + 1000000))
    assert exceso > 0 and n3["exceso_no_salarial"] == exceso and n3["ibc"] == sal3 + exceso and n3["alertas"]
    assert n["prestamos"] == 0  # el préstamo es solo de la 2.ª quincena

    # Ausencias: 2 días × 300.000 / 30 = 20.000 menos
    n2 = _nomina(admin, p1, "2002")
    assert n2["dias_ausencia"] == 2 and n2["auxilios"][0]["valor"] == 150000 - 20000 and "ausencia" in n2["auxilios"][0]["calculo"]

    p2 = _periodo_con(admin, 2026, 12, 2, [[1001, "1001", *(["A"] * 16)]], list(range(16, 32)))
    n = _nomina(admin, p2, "1001")
    assert n["prestamos"] == 80000 and "solo 2.ª quincena" in n["descuentos"][0]["calculo"]
    assert [x["valor"] for x in n["auxilios"]] == [150000]  # la alimentación es solo de la 1.ª

    lista = admin.get(f"{API}/auxilios?q=1001").json()["data"]
    rodamiento = next(x for x in lista if x["descripcion"] == "Auxilio de rodamiento")
    assert rodamiento["pagado"] == 300000 and rodamiento["quincena"] == "ambas"
    assert admin.delete(f"{API}/auxilios/{rodamiento['id']}").status_code == 409  # ya se pagó: se desactiva
    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{p1}/nomina").content))
    assert "Auxilios no prestacionales" in libro.sheetnames
    enc = [c.value for c in libro.worksheets[0][1]]
    assert "AUXILIOS NO PRESTACIONALES" in enc


def test_vacaciones_habiles_promedio_y_saldo(admin):
    # Septiembre Q1 previa: 3003 hace 3 turnos N → recargo nocturno que alimenta el promedio sugerido (medio mes)
    p1 = _periodo_con(admin, 2026, 9, 1, [["3003", "3003", *(["N"] * 3)]], [1, 2, 3])
    nocturno = _nomina(admin, p1, "3003")["valores"]["ordinary_night"]
    assert nocturno > 0
    pid = _quincena(admin)
    assert admin.post(f"{API}/auxilios", json={"documento": "3003", "descripcion": "Rodamiento", "valor_mensual": 300000,
                                               "desde": "2026-09-01"}).status_code == 201
    _cargar(admin, pid, _excel_sep_q2())
    # V del miércoles 16 al domingo 20: lunes a sábado → 4 hábiles y el domingo no hábil
    n = _nomina(admin, pid, "3003")
    v = n["vacaciones"]
    sugerido = round(nocturno / 0.5)
    assert (v["dias_habiles"], v["dias_no_habiles"], v["meses_promedio"], v["promedio_sugerido"]) == (4, 1, 0.5, sugerido)
    assert not v["promedio_manual"] and v["saldo"] is None and v["base"] == 1750905 + sugerido
    assert v["valor_habiles"] == round(4 * (1750905 + sugerido) / 30) and n["ibc"] >= v["valor"]
    # El auxilio se descuenta por todo el tiempo no laborado: 5 vacaciones, 2 incapacidad, 1 LR y 1 AUS → 150.000 − 9 × 10.000
    assert n["auxilios"][0]["valor"] == 60000 and n["dias_sin_pago"] == 2 and n["dias_tnl"] == 9
    assert "1 de ausencia, 2 de incapacidad, 5 de vacaciones, 1 de otras novedades" in n["auxilios"][0]["calculo"]

    # Lunes a viernes, ingreso 1-jun-2026 (5 días causados al 30-sep) → 3 hábiles y 2 no hábiles, sin alerta
    eid = _resultados(admin, pid)["3003"]["empleado_id"]
    cambio = {"dias_habiles": [0, 1, 2, 3, 4], "fecha_ingreso": "2026-06-01", "vacaciones_previas": 0}
    assert admin.put(f"{API}/empleados/{eid}/vacaciones", json=cambio).status_code == 200
    assert admin.post(f"{API}/periodos/{pid}/recalcular").status_code == 200
    v = _nomina(admin, pid, "3003")["vacaciones"]
    assert (v["dias_habiles"], v["dias_no_habiles"]) == (3, 2) and v["saldo"]["causados"] == 5 and v["saldo"]["pendientes"] == 2
    assert not [x for x in _nomina(admin, pid, "3003")["alertas"] if "vacaciones" in x]
    # Con 4 días ya disfrutados antes del liquidador solo le queda 1: se avisa
    assert admin.put(f"{API}/empleados/{eid}/vacaciones", json={**cambio, "vacaciones_previas": 4}).status_code == 200
    admin.post(f"{API}/periodos/{pid}/recalcular")
    alertas = [x for x in _nomina(admin, pid, "3003")["alertas"] if "vacaciones" in x]
    assert len(alertas) == 1 and "solo tenía 1.00 pendientes" in alertas[0]

    # Nómina fija el promedio y se recalcula; vacío vuelve al sugerido
    r = admin.put(f"{API}/periodos/{pid}/vacaciones/{eid}", json={"promedio": 100000})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["promedio"] == 100000 and r.json()["data"]["promedio_manual"]
    assert r.json()["data"]["valor_habiles"] == 185091  # 3 × 1.850.905 / 30 redondeado hacia arriba
    lista = admin.get(f"{API}/periodos/{pid}/vacaciones").json()["data"]
    assert [x["documento"] for x in lista] == ["3003"] and lista[0]["alertas"]
    assert admin.put(f"{API}/periodos/{pid}/vacaciones/{eid}", json={"promedio": None}).json()["data"]["promedio"] == sugerido

    emp = admin.get(f"{API}/empleados?q=3003").json()["data"][0]
    assert emp["dias_habiles"] == [0, 1, 2, 3, 4] and emp["vacaciones_liquidador"] == 3 and emp["saldo_vacaciones"]["previos"] == 4
    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/nomina").content))
    assert libro["Vacaciones"].max_row == 2


def test_varias_nominas_en_un_periodo(admin):
    pid = _quincena(admin)  # nómina «General» de sep-2026 Q2
    _cargar(admin, pid)
    assert admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "prestamo", "descripcion": "Celular",
                                                 "valor_mensual": 80000, "desde": "2026-09-01"}).status_code == 201
    r = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "quincena": 2, "nombre": "  Cliente   2 "})
    assert r.status_code == 201 and r.json()["data"]["nombre"] == "Cliente 2"
    c2 = r.json()["data"]["id"]
    assert admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "quincena": 2, "nombre": "cliente 2"}).status_code == 409
    # 1001 también cubre el 16 y el 17 en el Cliente 2 (repetidos) y el 18 no; 4004 solo está aquí
    r = _cargar(admin, c2, xlsx([["documento", "SEPTIEMBRE", 16, 17, 18], [1001, "1001", "A", "A", ""], [4004, "4004", "A", "A", "A"]]))
    assert r.status_code == 200, r.text
    n2 = _nomina(admin, c2, "1001")
    assert n2["otras_nominas"] and n2["descuentos"] == [] and n2["prestamos"] == 0
    assert any("«General»" in a and "2 día(s) están en las dos (16/09, 17/09)" in a for a in n2["alertas"])
    assert any("se aplican en la nómina «General»" in a for a in n2["alertas"])
    # En la General (creada primero) sigue el préstamo y avisa que también está en Cliente 2
    n1 = _nomina(admin, pid, "1001")
    assert n1["prestamos"] == 40000 and any("«Cliente 2»" in a for a in n1["alertas"])
    assert not _nomina(admin, c2, "4004")["alertas"]
    assert "Cliente_2" in admin.get(f"{API}/periodos/{c2}/nomina").headers["content-disposition"]
    assert admin.put(f"{API}/periodos/{c2}/nombre", json={"nombre": "General"}).status_code == 409
    assert admin.put(f"{API}/periodos/{c2}/nombre", json={"nombre": "Cliente Norte"}).json()["data"]["nombre"] == "Cliente Norte"
    # Al eliminar la General, Cliente Norte se recalcula y pasa a aplicar el préstamo
    assert admin.delete(f"{API}/periodos/{pid}").status_code == 200
    n2 = _nomina(admin, c2, "1001")
    assert n2["prestamos"] == 40000 and not n2["otras_nominas"]


def test_pagado_por_mes_solo_suma_nominas_cerradas(admin):
    q1 = _periodo_con(admin, 2026, 11, 1, [[1, "1", *(["A"] * 15)], [2, "2", *(["A"] * 15)]], list(range(1, 16)))
    r = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Administrativa"})
    adm = r.json()["data"]["id"]
    _cargar(admin, adm, xlsx([["documento", "NOVIEMBRE", *range(1, 16)], [9, "9", *(["A"] * 15)]]))
    q2 = _periodo_con(admin, 2026, 11, 2, [[1, "1", *(["A"] * 15)]], list(range(16, 31)))
    for pid in (q1, q2):
        assert admin.post(f"{API}/periodos/{pid}/cerrar").status_code == 200
    r = admin.get(f"{API}/reportes/pagado?anio=2026&mes=11").json()
    extra = r["meta"]["extra"]
    assert [(n["nombre"], n["tipo"], n["incluida"]) for n in extra["nominas"]] == [
        ("General", "Q1", True), ("Administrativa", "Q1", False), ("General", "Q2", True)]
    assert extra["abiertas"] == 1
    netos = {n["id"]: n["neto"] for n in extra["nominas"]}
    assert extra["totales"]["neto"] == netos[q1] + netos[q2]
    personas = {x["documento"]: x for x in r["data"]}
    assert set(personas) == {"1", "2"} and personas["1"]["nominas"] == ["General (Q1)", "General (Q2)"]
    from openpyxl import load_workbook as _lw
    libro = _lw(BytesIO(admin.get(f"{API}/reportes/pagado/excel?anio=2026&mes=11").content))
    assert libro.sheetnames == ["Por nomina", "Por persona"] and libro["Por persona"].max_row == 3


def test_costeo_de_nomina_no_paga_ni_crea_empleados(admin):
    from tests.conftest import archivo

    pid = _periodo_con(admin, 2026, 8, 2, [[1001, "1001", *(["A"] * 16)]], list(range(16, 32)))
    eid = _resultados(admin, pid)["1001"]["empleado_id"]
    assert admin.put(f"{API}/empleados/{eid}/salario", json={"salario": 3000000}).status_code == 200
    assert admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "prestamo", "descripcion": "Celular",
                                                 "valor_mensual": 80000, "monto_total": 160000, "desde": "2026-09-01"}).status_code == 201
    contenido = xlsx([["documento", "SEPTIEMBRE", *range(16, 31)], [1001, "1001", *(["A"] * 15)], [7777, "Aspirante", *(["A"] * 15)]])
    r = admin.post(f"{API}/costeos", files=archivo(contenido, "costeo.xlsx"),
                   data={"nombre": "Cliente nuevo", "anio": "2026", "mes": "9", "quincena": "2"})
    assert r.status_code == 201, r.text
    c = r.json()["data"]
    por_doc = {x["documento"]: x for x in c["resultados"]}
    assert por_doc["7777"]["existe"] is False and por_doc["7777"]["nomina"]["salario"] == 1750905
    assert por_doc["1001"]["existe"] and por_doc["1001"]["nomina"]["salario"] == 3000000 and por_doc["1001"]["nomina"]["prestamos"] == 0
    x = por_doc["7777"]["costo"]
    n = por_doc["7777"]["nomina"]
    assert x["exonerado"] and x["aportes"]["salud"] == 0 and x["aportes"]["icbf"] == 0 and x["aportes"]["sena"] == 0
    assert x["aportes"]["pension"] == round(n["ibc"] * 0.12) and x["aportes"]["caja"] == round(n["ibc"] * 0.04)
    assert x["base_prestaciones"] == n["devengado"] and x["prestaciones"]["prima"] == round(n["devengado"] * 0.0833)
    assert x["costo_total"] == n["devengado"] + x["total_aportes"] + x["total_prestaciones"]
    assert c["totales"]["personas"] == 2 and c["totales"]["no_registradas"] == 1
    assert c["totales"]["costo_total"] == sum(y["costo"]["costo_total"] for y in c["resultados"])
    # No crea al aspirante ni descuenta el préstamo
    assert admin.get(f"{API}/empleados?q=7777").json()["data"] == []
    assert admin.get(f"{API}/descuentos?q=1001").json()["data"][0]["descontado"] == 0
    assert admin.get(f"{API}/costeos").json()["data"][0]["nombre"] == "Cliente nuevo"
    # Sin exoneración se suman salud, ICBF y SENA
    par = {**c["parametros"], "exonerado": False}
    c2 = admin.put(f"{API}/costeos/{c['id']}", json={"nombre": "Cliente nuevo", "parametros": par}).json()["data"]
    y = {z["documento"]: z for z in c2["resultados"]}["7777"]["costo"]
    assert y["aportes"]["salud"] == round(n["ibc"] * 0.085) and c2["totales"]["costo_total"] > c["totales"]["costo_total"]
    libro = load_workbook(BytesIO(admin.get(f"{API}/costeos/{c['id']}/excel").content))
    assert libro.sheetnames == ["Costo empresa", "Nomina", "Parametros"] and libro["Costo empresa"].max_row == 4
    assert admin.post(f"{API}/costeos", files=archivo(contenido, "costeo.xlsx"),
                      data={"nombre": "x", "anio": "2026", "mes": "10", "quincena": "1"}).status_code == 422  # fechas de otro periodo
    assert admin.delete(f"{API}/costeos/{c['id']}").status_code == 200



def test_arl_por_empleado_y_auxilio_del_costeo(admin):
    from tests.conftest import archivo

    # En la nómina la ARL se configura en Empleados: la columna arl del Excel se ignora
    pid = _periodo_con(admin, 2026, 8, 2, [[1001, "1001", 3, *(["A"] * 16)], [3003, "3003", "", *(["A"] * 16)]], ["arl", *range(16, 32)])
    emp = {e["documento"]: e for e in admin.get(f"{API}/empleados").json()["data"]}
    assert emp["1001"]["arl_clase"] is None and emp["1001"]["quincenas"] == 1
    base = {"dias_habiles": [0, 1, 2, 3, 4, 5]}
    assert admin.put(f"{API}/empleados/{emp['3003']['id']}/configuracion", json={**base, "arl_clase": 5}).status_code == 200
    assert admin.put(f"{API}/empleados/{emp['3003']['id']}/configuracion", json={**base, "arl_clase": 6}).status_code == 422
    assert {e["documento"]: e for e in admin.get(f"{API}/empleados").json()["data"]}["3003"]["arl_clase"] == 5
    assert admin.get(f"{API}/periodos/{pid}").status_code == 200

    # Costeo: la clase viene del Excel (1001 → 2 en romano); si no viene, la del empleado (3003 → 5); si tampoco,
    # la clase por defecto (7777 → 4). "x" no es válida. Auxilio del costeo de 300.000 mensual; 3003 tiene 1 AUS
    contenido = xlsx([["documento", "SEPTIEMBRE", "arl", *range(16, 31)], [1001, "1001", "II", *(["A"] * 15)],
                      [3003, "3003", "", "AUS", *(["A"] * 14)], [7777, "Nuevo", "x", *(["A"] * 15)]])
    c = admin.post(f"{API}/costeos", files=archivo(contenido, "c.xlsx"),
                   data={"nombre": "Con ARL", "anio": "2026", "mes": "9", "quincena": "2", "auxilio_mensual": "300000"}).json()["data"]
    r = {x["documento"]: x for x in c["resultados"]}
    assert (r["1001"]["arl_clase"], r["3003"]["arl_clase"], r["7777"]["arl_clase"]) == (2, 5, None)
    assert (r["1001"]["arl_origen"], r["3003"]["arl_origen"], r["7777"]["arl_origen"]) == ("excel", "empleado", None)
    assert (r["1001"]["costo"]["arl_clase"], r["3003"]["costo"]["arl_clase"], r["7777"]["costo"]["arl_clase"]) == (2, 5, 4)
    assert any("ARL 'x' no es válida" in a for a in c["advertencias"])
    assert r["3003"]["costo"]["aportes"]["arl"] == round(r["3003"]["costo"]["base_arl"] * 0.0696)
    assert c["totales"]["sin_arl"] == 1 and c["auxilio_mensual"] == 300000
    assert r["7777"]["nomina"]["auxilios_no_prestacionales"] == 150000
    assert r["3003"]["nomina"]["auxilios_no_prestacionales"] == 140000  # 150.000 − 1 ausencia × 10.000
    # Sin auxilio en el costeo no se suma ninguno (ni los registrados en el módulo de auxilios)
    assert admin.post(f"{API}/auxilios", json={"documento": "1001", "descripcion": "Rodamiento", "valor_mensual": 200000,
                                               "desde": "2026-09-01"}).status_code == 201
    c0 = admin.post(f"{API}/costeos", files=archivo(contenido, "c.xlsx"),
                    data={"nombre": "Sin auxilio", "anio": "2026", "mes": "9", "quincena": "2"}).json()["data"]
    assert all(x["nomina"]["auxilios"] == [] for x in c0["resultados"])
    # La clase por defecto se puede cambiar y recalcula solo a quien no tiene clase
    par = {**c["parametros"], "arl_clase": 1}
    c2 = admin.put(f"{API}/costeos/{c['id']}", json={"nombre": "Con ARL", "parametros": par}).json()["data"]
    r2 = {x["documento"]: x for x in c2["resultados"]}
    assert r2["7777"]["costo"]["arl_clase"] == 1 and r2["3003"]["costo"]["arl_clase"] == 5
    assert load_workbook(BytesIO(admin.get(f"{API}/costeos/{c['id']}/excel").content))["Costo empresa"]["D4"].value == "1 (por defecto)"
    hoja = load_workbook(BytesIO(admin.get(f"{API}/costeos/plantilla?anio=2026&mes=9&quincena=2").content)).active
    assert next(hoja.iter_rows(values_only=True))[:4] == ("documento", "SEPTIEMBRE 2026", "arl", 16)


def test_ficha_del_empleado(admin):
    pid = _quincena(admin)
    _cargar(admin, pid)
    eid = _resultados(admin, pid)["3003"]["empleado_id"]
    assert admin.post(f"{API}/descuentos", json={"documento": "3003", "tipo": "prestamo", "descripcion": "Celular",
                                                 "valor_mensual": 80000, "desde": "2026-09-01"}).status_code == 201
    assert admin.post(f"{API}/auxilios", json={"documento": "3003", "descripcion": "Rodamiento", "valor_mensual": 100000,
                                               "desde": "2026-09-01"}).status_code == 201
    assert admin.post(f"{API}/auxilios", json={"documento": "1001", "descripcion": "Otro", "valor_mensual": 100000,
                                               "desde": "2026-09-01"}).status_code == 201
    f = admin.get(f"{API}/empleados/{eid}").json()["data"]
    assert f["documento"] == "3003" and f["quincenas"] == 1 and f["nominas"][0]["nombre"] == "General" and "devengado" in f["nominas"][0]
    # Todo se guarda de una vez y queda auditado con lo que cambió
    datos = {"salario": 2500000, "arl_clase": 4, "dias_habiles": [0, 1, 2, 3, 4], "fecha_ingreso": "2025-01-15", "vacaciones_previas": 2}
    r = admin.put(f"{API}/empleados/{eid}/ficha", json=datos)
    assert r.status_code == 200 and set(r.json()["data"]["cambios"]) == {"salario", "arl_clase", "dias_habiles", "fecha_ingreso", "vacaciones_previas"}
    f = admin.get(f"{API}/empleados/{eid}").json()["data"]
    assert (f["salario"], f["arl_clase"], f["dias_habiles"], f["fecha_ingreso"]) == (2500000, 4, [0, 1, 2, 3, 4], "2025-01-15")
    assert f["saldo_vacaciones"]["previos"] == 2
    assert admin.put(f"{API}/empleados/{eid}/ficha", json={**datos, "salario": None}).json()["data"]["cambios"] == ["salario"]
    # Préstamos y auxilios de esa persona
    assert [x["descripcion"] for x in admin.get(f"{API}/descuentos?empleado_id={eid}").json()["data"]] == ["Celular"]
    assert [x["descripcion"] for x in admin.get(f"{API}/auxilios?empleado_id={eid}").json()["data"]] == ["Rodamiento"]
    assert admin.get(f"{API}/empleados/999999").status_code == 404


def test_cuota_exequial_y_cargo_por_defecto(admin):
    dias = list(range(1, 16))
    p1 = _periodo_con(admin, 2027, 1, 1, [[1001, "1001", *(["A"] * 15)]], dias)
    base = {"documento": "1001", "desde": "2027-01-01"}
    assert admin.post(f"{API}/descuentos", json={**base, "tipo": "cuota", "descripcion": "Plan exequial", "valor_mensual": 30000}).status_code == 201
    assert admin.post(f"{API}/descuentos", json={**base, "tipo": "prestamo", "descripcion": "Préstamo", "valor_mensual": 40000}).status_code == 201
    assert admin.post(f"{API}/descuentos", json={**base, "tipo": "otro", "descripcion": "x", "valor_mensual": 1}).status_code == 422
    admin.post(f"{API}/periodos/{p1}/recalcular")

    n = _nomina(admin, p1, "1001")
    assert [x["tipo"] for x in n["descuentos"]] == ["prestamo", "cuota"]  # la cuota va después del préstamo
    assert n["cuotas"] == 15000 and n["prestamos"] == 20000
    assert n["deducciones"] == n["salud"] + n["pension"] + 35000

    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{p1}/nomina").content))
    filas = list(libro.worksheets[0].iter_rows(values_only=True))
    assert dict(zip(filas[0], filas[1]))["CUOTAS"] == 15000
    assert any(f[0] == "1001" and f[2] == "Cuota" for f in libro.worksheets[1].iter_rows(values_only=True))

    # SERA no tiene vigilantes: sin columna de cargo en el Excel, la persona queda como operario de aseo
    e = next(x for x in admin.get(f"{API}/empleados?q=1001").json()["data"])
    assert e["cargo"] == "Operario de aseo"
