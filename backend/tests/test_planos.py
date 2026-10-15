"""Planos SIESA: reglas del plano de empleados y de contratos (casos de aceptación de los prompts maestros) y sus rutas."""

import base64
import json
from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest

from app.apps.planos.services import contratos, empleados
from app.apps.planos.services.comun import PlanoInvalido, col, normalizar_fecha
from app.apps.planos.services.dane import ResolverDane
from app.apps.planos.services.entidades import Codificador
from tests.conftest import archivo

DANE = ResolverDane()

# ---------------- Empleados ----------------


@pytest.mark.parametrize("entrada,esperado", [
    (("169", "Valle", "Cali"), ("76", "001")),
    (("169", "VALLE", "Buga"), ("76", "111")),  # alias: no Bugalagrande
    (("169", "Valle del Cauca", "Buenaventura"), ("76", "109")),
    (("169", "Cuaca", "Popayán"), ("19", "001")),
    (("169", "Bogotá", "Bogotá"), ("11", "001")),
    (("169", "Cundinamarca", "Bogota D.C."), ("11", "001")),
    (("169", "Cundinamarca", "Nariño"), ("25", "483")),
    (("169", "Antioquia", "Nariño"), ("05", "483")),
    (("169", "Nariño", "Nariño"), ("52", "480")),
    (("169", "Atlántico", "Barranquilla"), ("08", "001")),
    (("169", "Cundinamarca", "Sopó"), ("25", "758")),
    (("169", "Arauca", "Fortul"), ("81", "300")),
    (("850", "Distrito Capital", "Caracas"), ("01", "001")),
    (("239", "Pichincha", "Quito"), ("01", "001")),
    (("580", "Lima", "Lima"), ("01", "001")),
    (("169", "76", "1"), ("76", "001")),  # ya codificado: idempotente
    (("Colombia", "N. de Santander", "Cúcuta"), ("54", "001")),
])
def test_dane_exactos(entrada, esperado):
    r = DANE.resolver(*entrada)
    assert (r.depto, r.ciudad) == esperado and not r.problemas, r.problemas


@pytest.mark.parametrize("entrada,ciudad,aviso", [
    (("169", "Cordoba", "Tuchin"), "182", "aproximada"),  # Tuchín no está: no se acepta Chinú en silencio
    (("169", "Cauca", "El Bordo"), "256", "aproximada"),
    (("169", "", "Cali"), "001", "único municipio"),
    (("169", "Valle", "Ciudad Inventada"), None, "no encontrada"),
])
def test_dane_aproximadas_se_reportan(entrada, ciudad, aviso):
    r = DANE.resolver(*entrada)
    assert r.ciudad == ciudad and aviso in " ".join(m for _, m, _ in r.problemas)


def test_alias_del_usuario():
    r = ResolverDane(alias_ciudad={("169", "19", "EL BORDO"): "532"}).resolver("169", "Cauca", "El Bordo")
    assert r.ciudad == "532" and not r.problemas


def test_reglas_sueltas():
    assert normalizar_fecha("19980231") == ("19980231", False)
    assert normalizar_fecha("1998085") == ("1998085", False)
    assert normalizar_fecha("18991229")[1] is False
    assert normalizar_fecha(datetime(2012, 12, 18)) == ("20121218", True)
    assert normalizar_fecha("18/12/2012") == ("20121218", True)
    assert empleados.limpiar_nombre("  Muñoz, Peña ") == "MUNOZ PENA"
    assert empleados.limpiar_nombre("Muñoz", conservar_enie=True) == "MUÑOZ"
    assert empleados.convertir_nivel_educativo("tecnico")[0] == "TÉCNICO"
    assert empleados.convertir_nivel_educativo("Maestria")[0] == "MAESTRÍA"
    assert empleados.convertir_nivel_educativo("Tecnico laboral")[0] == "TECNICO LABORAL"
    assert empleados.convertir_nivel_educativo("Bachiller")[1] is False
    assert empleados.convertir_genero("Femenino") == ("1", True) and empleados.convertir_genero(None) == ("", True)
    assert empleados.convertir_estado_civil("Casado") == ("2", True)


def _libro(encabezados: list[str], filas: list[dict], extra_hojas=()) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(encabezados)
    for f in filas:
        fila = [None] * len(encabezados)
        for c, v in f.items():
            fila[col(c) - 1] = v
        ws.append(fila)
    for nombre in extra_hojas:
        wb.create_sheet(nombre).append(["Fila", "Cédula", "Columna", "Tipo", "Problema"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


ENC_EMPLEADOS = ["ID DEL TERCERO (Obligatorio)"] + [f"C{i}" for i in range(2, 64)]
EMPLEADO = {"A": "14835286", "B": "c", "C": "  Muñoz, ", "E": "Ana  María", "F": "003", "I": "20200101", "J": "19900101",
            "K": "169", "L": "Valle", "M": "Cali", "N": "20080101", "O": "Colombia", "P": "Valle", "Q": "Buga",
            "R": "Maestria", "S": "Femenino", "T": "Casado"}


def _leer(libro: bytes):
    wb = openpyxl.load_workbook(BytesIO(libro))
    return wb, wb.worksheets[0]


def test_empleados_completo_y_correccion():
    contenido = _libro(ENC_EMPLEADOS, [{}, EMPLEADO, {**EMPLEADO, "A": "1.007.897,943 ", "K": "Venezuela", "L": "Distrito capital",
                                                     "M": "Caracas", "I": "20240101", "J": "20240101"}], extra_hojas=["Errores"])
    libro, res, cruce = empleados.procesar(contenido, DANE, "completo", {("14835286", "N"): "20080805"})
    wb, ws = _leer(libro)
    assert wb.sheetnames == ["Hoja1", "Errores"]  # la hoja de errores anterior se reemplaza
    v = {c: ws.cell(2, col(c)).value for c in "A B C E F I L M N P Q R S T U AH AI".split()}
    assert v == {"A": "14835286", "B": "C", "C": "MUNOZ", "E": "ANA MARIA", "F": "001", "I": "20200101", "L": "76", "M": "001",
                 "N": "20080805", "P": "76", "Q": "111", "R": "MAESTRÍA", "S": "1", "T": "2", "U": "0", "AH": "1", "AI": "0"}
    assert ws.cell(3, 1).value == "1007897943" and ws.cell(3, col("U")).value == "1"
    assert ws.cell(2, col("L")).number_format == "@"
    assert res.resumen["filas"] == 2 and res.resumen["filas_eliminadas"] == 1 and res.resumen["correcciones"] == 1
    assert [h.columna for h in res.hallazgos if h.cedula == "14835286"] == []
    assert any(h.columna == "I/J" and h.cedula == "1007897943" for h in res.hallazgos)
    assert cruce == {"14835286": "20200101", "1007897943": "20240101"}
    errores = [[c.value for c in r] for r in wb["Errores"].iter_rows()]
    assert errores[0] == ["Fila", "Cédula", "Columna", "Tipo", "Problema"] and errores[1][0] == 4  # fila de la hoja de origen

    # Idempotente: procesar el resultado no cambia nada
    libro2, _, _ = empleados.procesar(libro, DANE, "completo", {("14835286", "N"): "20080805"})
    assert [[c.value for c in r] for r in _leer(libro2)[1].iter_rows()] == [[c.value for c in r] for r in ws.iter_rows()]


def test_empleados_modo_codigos_no_toca_nombres():
    libro, _, _ = empleados.procesar(_libro(ENC_EMPLEADOS, [EMPLEADO]), DANE, "codigos")
    ws = _leer(libro)[1]
    for c in "B C E F R S T".split():
        assert ws.cell(2, col(c)).value == EMPLEADO[c]
    assert ws.cell(2, col("U")).value is None and ws.cell(2, col("M")).value == "001"


def test_empleados_expedicion_17_anios_y_6_meses_y_vacios():
    filas = [{**EMPLEADO, "J": "20000115", "N": "20170714"}, {**EMPLEADO, "A": "2", "J": "20000115", "N": "20170715", "R": None, "S": None, "T": None}]
    _, res, _ = empleados.procesar(_libro(ENC_EMPLEADOS, filas), DANE, "completo")
    assert [h.cedula for h in res.hallazgos] == ["14835286"] and res.hallazgos[0].columna == "J/N"


def test_un_plano_una_hoja():
    wb = openpyxl.Workbook()
    wb.active.append(ENC_EMPLEADOS)
    wb.create_sheet("Otra").append(["datos"])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(PlanoInvalido, match="una sola hoja"):
        empleados.procesar(buf.getvalue(), DANE)


# ---------------- Contratos ----------------

COD = Codificador(
    {"BANCOS": {"07": "BANCOLOMBIA S.A.", "51": "BANCO DAVIVIENDA S.A.", "06": "ITAÚ CORPBANCA", "14": "ITAÚ (HELM BANK)"},
     "AFP": {"230301": "AFP PORVENIR", "230901": "SKANDIA FONDO DE PENSIONES OBLIGATORIAS", "230904": "SKANDIA FONDO ALTERNATIVO DE PENSIONES",
             "25-14": "AFP COLPENSIONES", "999999999": "SIN FONDO DE PENSIONES"},
     "EPS": {"EPS037": "NUEVA EPS", "EPS041": "NUEVA EPS S.A. MOVILIDAD", "EPS010": "EPS SURA"},
     "CESANTIAS": {"002": "CESANTIAS PORVENIR"},
     "CAJAS": {"CCF57": "CCF COMFANDI", "CCF56": "COMFENALCO VALLE", "CCF03": "COMFENALCO ANTIOQUIA CCF"},
     "CARGOS": {"032": "VIGILANTE", "033": "VIGILANTE MOTORIZADO", "031": "JEFE OPERACIONES", "041": "APRENDIZ GESTION HUMANA"}},
    {("AFP", "PORVENIR"): "230301", ("CESANTIAS", "PORVENIR"): "002", ("CAJAS", "COMANDI"): "CCF57", ("EPS", "SURA"): "EPS010"},
)
OBLIGATORIAS = {"A", "B", "C", "D", "F", "G", "I", "J", "K", "N", "O", "P", "Q", "S", "Z", "AA", "AD", "AE", "AF", "AG", "AH", "AJ",
                "AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AV", "AW", "AX", "AY", "AZ", "BA", "BB", "BC", "BD",
                "BE", "BF", "BG", "BH", "BI", "BJ", "BK", "BL", "BM", "BN", "BO", "BP", "BQ", "BR", "BS", "BT"}
ENC_CONTRATOS = ["ID CIA (Obligatorio)", "ID TERCERO (Obligatorio)"] + [
    f"{c} ({'Obligatorio' if c in OBLIGATORIAS else 'No obligatorio'})" for c in contratos.TODAS[2:]]
VIGILANTE = {"B": "1.144.195.447", "D": "002", "F": "Bancolombia", "N": "003", "P": "404", "S": "PORVENIR ", "T": "Sura",
             "U": "porvenir", "V": "COMANDI", "X": "PASENA ", "Z": "16/03/2026", "AJ": "1.500.000", "AW": "  123-456*789 ",
             "BN": "1", "BQ": "Vigilante", "G": None, "Q": None}
HOY = date(2026, 10, 15)


def test_contrato_vigilante():
    libro, res = contratos.procesar(_libro(ENC_CONTRATOS, [VIGILANTE]), COD, 1750905, hoy=HOY)
    wb, ws = _leer(libro)
    v = {c: ws.cell(2, col(c)).value for c in contratos.TODAS if ws.cell(2, col(c)).value is not None}
    assert wb.sheetnames == ["Hoja1", "Errores"]
    assert (v["B"], v["D"], v["F"], v["S"], v["T"], v["U"], v["V"], v["BQ"]) == ("1144195447", "001", "07", "230301", "EPS010", "002", "CCF57", "032")
    assert (v["I"], v["O"], v["P"], v["AJ"], v["AZ"], v["X"], v["BN"]) == ("999", "406999", "004", "1750905", "3", "PASENA", "0")
    assert (v["Z"], v["AA"], v["AI"], v["AC"]) == ("20260316", "20260316", "20260316", "20260615")  # 3 meses menos un día
    assert (v["AW"], v["AX"], v["G"], v["Q"], v["AY"], v["BM"], v["BO"]) == ("123456789", "2", "1", "Q", "1", "1", "1")
    assert "P: 404 → 004 (error de digitación)" in res.cambios and "D: 002 → 001" in res.cambios
    informados = {h.columna for h in res.hallazgos if h.tipo == "Informativo"}
    assert informados == {"S", "T", "U", "V"}  # alias confirmados: quedan en el reporte, sin pedir revisión
    assert not [h for h in res.hallazgos if h.tipo == "Error"], res.hallazgos
    libro2, _ = contratos.procesar(libro, COD, 1750905, hoy=HOY)  # idempotente
    assert [[c.value for c in r] for r in _leer(libro2)[1].iter_rows()] == [[c.value for c in r] for r in ws.iter_rows()]


def test_contrato_otros_cargos_y_alertas():
    base = {**VIGILANTE, "BQ": "Jefe operaciones", "I": "998", "O": "302001", "P": "004", "AJ": "4800000", "AC": "20260915"}
    filas = [
        base,
        {**base, "B": "222", "Z": "20060602", "AC": "20260901", "S": "SKANDIA", "F": "Itau", "N": "002", "AW": "123456789"},
        {**base, "B": "333", "AJ": "1500000", "AC": None, "Z": "20260131", "S": "999999999", "BQ": "Gerente de Marte", "I": None, "AW": None},
        {**base, "B": "444", "AJ": "1000000", "G": "19", "BQ": "41", "AC": "20260101"},
    ]
    libro, res = contratos.procesar(_libro(ENC_CONTRATOS, filas), COD, 1750905, hoy=HOY,
                                    empleados={"1144195447": "20260316", "333": "20260201"})
    ws = _leer(libro)[1]
    assert ws.cell(2, col("I")).value == "998" and ws.cell(2, col("O")).value == "302001" and ws.cell(2, col("AJ")).value == "4800000"
    assert ws.cell(2, col("AZ")).value == "0"  # más de 2 SMLV: sin auxilio
    assert ws.cell(4, col("AC")).value == "20260430" and ws.cell(4, col("AZ")).value == "3" and ws.cell(4, col("AX")).value == "0"
    por = lambda ced: {(h.columna, h.tipo) for h in res.hallazgos if h.cedula == ced}  # noqa: E731
    assert ("Z", "Error") in por("222") and ("AC", "Advertencia") in por("222") and ("N", "Advertencia") in por("222")
    assert ("S", "Error") in por("222") and ("F", "Error") in por("222")  # Skandia e Itaú: ambiguos, no se adivinan
    assert ("AW", "Error") in por("222")  # cuenta repetida en otra cédula
    assert {("AJ", "Error"), ("AY", "Advertencia"), ("BQ", "Error"), ("I", "Error"), ("Z", "Advertencia")} <= por("333")
    assert ("B", "Advertencia") in por("444") and ("AJ", "Error") not in por("444")  # aprendiz: puede ganar menos
    assert ("AC", "Error") in por("444")  # contrato hasta antes del ingreso
    assert ws.cell(5, col("BQ")).value == "041"


def test_fecha_hasta():
    assert [contratos.fecha_hasta(z) for z in ("20260316", "20260131", "20261130", "20260101")] == \
        ["20260615", "20260430", "20270228", "20260331"]


# ---------------- Rutas ----------------


def test_rutas_procesar_y_configurar(admin):
    r = admin.post("/api/planos/empleados", files=archivo(_libro(ENC_EMPLEADOS, [{**EMPLEADO, "Q": "El Bordo", "P": "Cauca"}]), "plano.xlsx"),
                   data={"modo": "completo"})
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["nombre"] == "plano_CARGUE.xlsx" and d["resumen"]["filas"] == 1 and d["cruce"] == {"14835286": "20200101"}
    assert any("aproximada" in h["problema"] for h in d["hallazgos"])
    assert openpyxl.load_workbook(BytesIO(base64.b64decode(d["archivo"]))).sheetnames == ["Hoja1", "Errores"]

    # Alias de ciudad creado por el usuario: la próxima vez es exacta
    assert admin.post("/api/planos/alias", json={"catalogo": "CIUDAD", "pais": "169", "departamento": "19", "texto": "el bordo",
                                                 "codigo": "999"}).status_code == 422
    a = admin.post("/api/planos/alias", json={"catalogo": "CIUDAD", "pais": "169", "departamento": "19", "texto": "el bordo", "codigo": "532"})
    assert a.status_code == 200 and a.json()["data"]["texto"] == "EL BORDO" and "Pat" in a.json()["data"]["destino"]
    d = admin.post("/api/planos/empleados", files=archivo(_libro(ENC_EMPLEADOS, [{**EMPLEADO, "Q": "El Bordo", "P": "Cauca"}])),
                   data={"modo": "completo"}).json()["data"]
    assert d["hallazgos"] == []

    # Corrección manual por cédula
    c = admin.post("/api/planos/correcciones", json={"plano": "empleados", "cedula": "14.835.286", "columna": "N", "valor": "20080805"})
    assert c.status_code == 200 and c.json()["data"]["cedula"] == "14835286"
    d = admin.post("/api/planos/empleados", files=archivo(_libro(ENC_EMPLEADOS, [EMPLEADO]))).json()["data"]
    assert d["resumen"]["correcciones"] == 1

    # Contratos: catálogos sembrados, alias de entidad y cruce con empleados
    assert len(admin.get("/api/planos/entidades?catalogo=CARGOS").json()["data"]) >= 48
    assert admin.post("/api/planos/alias", json={"catalogo": "EPS", "texto": "Coosalud", "codigo": "NOEXISTE"}).status_code == 422
    assert admin.put("/api/planos/entidades", json={"catalogo": "EPS", "codigo": "ESS024", "descripcion": "COOSALUD EPS"}).status_code == 200
    assert admin.post("/api/planos/alias", json={"catalogo": "EPS", "texto": "Coosalud", "codigo": "ESS024"}).status_code == 200
    plano = _libro(ENC_CONTRATOS, [{**VIGILANTE, "T": "coosalud"}])
    r = admin.post("/api/planos/contratos", files=archivo(plano), data={"empleados_cruce": json.dumps({"1144195447": "20260316"})})
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    hoja = openpyxl.load_workbook(BytesIO(base64.b64decode(d["archivo"]))).worksheets[0]
    assert hoja.cell(2, col("T")).value == "ESS024" and d["cruzado"] is True

    # Salario mínimo por año
    assert admin.put("/api/planos/salarios/2027", json={"valor": 1900000}).status_code == 200
    assert [s["anio"] for s in admin.get("/api/planos/salarios").json()["data"]][:2] == [2027, 2026]
    eventos = admin.get("/api/plataforma/auditoria?accion=plano_procesado").json()["data"]
    assert eventos and eventos[0]["detalle"]["plano"] == "contratos"


def test_rutas_rechazos(admin):
    r = admin.post("/api/planos/contratos", files=archivo(_libro(ENC_EMPLEADOS, [EMPLEADO])))
    assert r.status_code == 422 and "contratos" in r.json()["error"]["message"]
    assert admin.post("/api/planos/empleados", files=archivo(b"no es excel", "x.xlsx")).status_code == 415
