"""Validación de nómina: un escenario sintético (septiembre 2026) que ejercita cada regla y cada alerta."""

from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import ACCESOS, archivo, entrar, xlsx

API = "/api/nomina"
DIA_AUX = 249095 / 30  # auxilio de transporte por día


def _modalidades() -> bytes:
    return xlsx([
        ["Código", "Descripción", "Activo", "Método de liquidación", "Concepto adicional", "Descripción concepto adicional", "Valor concepto adicional"],
        ["4X2 1 GRP", "", "", "", "", "", ""],
        ["4X2 1 GRP", "4X2 1 GRP 433173*", "Si", "Valor por hora liquidada", "130", "HORAS EXTRAS Y RECARGOS", 20000],
        ["4X2 1 GRP", "4X2 1 GRP 433173*", "Si", "Valor por hora liquidada", "132", "APLICACION LEY 2101", 3000],
        ["6X1 ESP", "", "", "", "", "", ""],
        ["6X1 ESP", "6X1 ESPECIAL*", "Si", "Valor por hora liquidada", "130", "HORAS EXTRAS Y RECARGOS", 30000],
    ])


def _ubicaciones() -> bytes:
    enc = ["Código", "Descripción", "Modalidad de pago", "Descripción de Modalidad", "Codigo puesto", "Descripción puesto",
           "Código C.Costo", "Descripción C.Costo", "Modalidad pago"]
    return xlsx([
        enc,
        ["10", "CONJUNTO UNO", "4X2 1 GRP", "4X2 1 GRP 433173*", "10", "PORTERIA", "400010", "CR UNO", ""],
        # El puesto tiene su propia modalidad (viene la descripción): manda sobre la de la ubicación
        ["10", "CONJUNTO UNO", "4X2 1 GRP", "4X2 1 GRP 433173*", "10-1", "RONDA", "400010", "CR UNO", "6X1 ESPECIAL*"],
        ["20", "EDIFICIO DOS", "", "", "20", "RECEPCION", "400020", "ED DOS", ""],
    ])


CONTRATOS = [
    # cédula, nombre, tipo de nómina, grupo
    ("1001", "ANA COMPLETA", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1002", "BETO VACACIONES", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1003", "CARLA SIN AUXILIO", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1004", "DIEGO SIN PROGRAMACION", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1005", "EVA SIN NOMINA", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1006", "FABIO ADMIN", "QUINCENAL", "ADMINISTRATIVOS/COMERCIALES"),
    ("1007", "GINA DE LOS SIETE", "QUINCENAL", "OPERATIVOS FULL"),
    ("1008", "HUGO SIN MODALIDAD", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1009", "IVAN RONDA", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("2001", "JUAN MENSUAL", "MENSUAL", "OPERATIVOS CON VARIABLE"),
    ("1010", "KAREN SALE A VACACIONES", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1011", "LUIS REGRESA DE VACACIONES", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
    ("1012", "MARIA CAMBIA DE PUESTO", "QUINCENAL", "OPERATIVOS CON VARIABLE"),
]


def _contratos() -> bytes:
    filas = [["Código interno", "Empleado", "Nombre del empleado", "Estado", "Descripcion estado", "Descripcion del cargo",
              "Descripcion ccosto", "Descripcion tipo de nomina", "Descripcion grupo empleados", "Fecha ingreso"]]
    filas += [[i, ced, nombre, "1", "Activo", "VIGILANTE", "CC", tipo, grupo, "2024-01-01"] for i, (ced, nombre, tipo, grupo) in enumerate(CONTRATOS)]
    return xlsx(filas)


def _cuotas() -> bytes:
    enc = ["Empleado", "Nombre", "NDC", "Cpto", "Descripción Cpto", "Cuotas", "Valor máximo", "Valor acumulado", "Valor devengo",
           "Valor deducción", "Tasa", "Estado"]
    return xlsx([
        enc,
        ["1001", "ANA", "1", "600", "SEGUROS EXEQUIALES", "1", 0, 0, 0, 15900, 0, "En Proceso"],  # solo 1.ª quincena
        ["1001", "ANA", "1", "611", "EMBARGO 5A PARTE", "2", 5000000, 100000, 0, 0, 20, "En Proceso"],
        # Ajuste por mayor valor pagado: aunque no se descuente, no es alerta
        ["1001", "ANA", "1", "610", "DEDUCCION MAYOR VALOR PAGADO", "0", 50000, 0, 0, 50000, 0, "En Proceso"],
        ["1002", "BETO", "1", "605", "PRESTAMO EMPRESA", "1", 100000, 0, 0, 200000, 0, "En Proceso"],  # cuota > tope
        ["1003", "CARLA", "1", "605", "PRESTAMO EMPRESA", "5", 600000, 570000, 0, 100000, 0, "En Proceso"],  # saldo 30.000
        ["1006", "FABIO", "1", "605", "PRESTAMO EMPRESA", "0", 500000, 0, 0, 50000, 0, "Pendiente"],
        ["1009", "IVAN", "1", "605", "PRESTAMO EMPRESA", "3", 500000, 100000, 0, 50000, 0, "Inactivo"],
        # Rodamiento administrativo: en la nómina se paga con el concepto 129 RODAMIENTO
        ["1006", "FABIO", "1", "152", "RODAMIENTO_ADM", "4", 0, 0, 600000, 0, 0, "En Proceso"],
        ["1009", "IVAN", "1", "152", "RODAMIENTO_ADM", "4", 0, 0, 300000, 0, 0, "En Proceso"],
    ])


def _programacion() -> bytes:
    trabajo = ["06:00 - 18:00", "06:00 - 18:00", "18:00 - 06:00", "18:00 - 06:00", "Z", "L"] * 3  # 18 días, todos pagables
    dias: dict[str, dict[int, str]] = {
        "1001": {16 + i: trabajo[i] for i in range(15)},
        "1002": {**{16 + i: trabajo[i] for i in range(10)}, **{d: "[VAC]" for d in range(26, 31)}},
        "1003": {16 + i: trabajo[i] for i in range(15)},
        "1005": {16 + i: trabajo[i] for i in range(15)},
        "1008": {16 + i: trabajo[i] for i in range(15)},
        "1009": {**{16 + i: trabajo[i] for i in range(14)}, 30: "XYZ"},  # código nuevo sin corchetes: no descuenta
        "2001": {d: trabajo[(d - 1) % 6] for d in range(1, 31)},
        "1010": {**{16 + i: trabajo[i] for i in range(8)}, **{d: "[VAC]" for d in range(24, 31)}},
        "1011": {**{d: "[VAC]" for d in range(16, 21)}, **{21 + i: trabajo[i] for i in range(10)}},
    }
    puesto = {"1001": ("10", "10"), "1002": ("10", "10"), "1003": ("10", "10"), "1005": ("10", "10"), "1008": ("20", "20"),
              "1009": ("10", "10-1"), "2001": ("10", "10"), "1010": ("10", "10"), "1011": ("10", "10")}
    filas: list[list[object]] = [["Compañia", "EMPRESA"], ["Desde", "2026-09-01"], ["Hasta", "2026-09-30"], [],
                                 ["C.C. Empleado", "NOMBRE EMPLEADO", "UBICACION", "DESCRIPCION UBICACION", "PUESTO", "DESCRIPCION PUESTO",
                                  *[str(d) for d in range(1, 32)]]]
    for ced, d in dias.items():
        ubi, pto = puesto[ced]
        nombre = next(c[1] for c in CONTRATOS if c[0] == ced)
        filas.append([ced, nombre, ubi, f"UBI {ubi}", pto, f"PUESTO {pto}", *[d.get(i) for i in range(1, 32)]])
    # María: del 16 al 23 en el puesto 10 (4X2 1 GRP, 20.000/día) y del 24 al 30 en el 10-1 (6X1, 30.000/día)
    maria = next(c[1] for c in CONTRATOS if c[0] == "1012")
    filas.append(["1012", maria, "10", "UBI 10", "10", "PUESTO 10", *[trabajo[i - 16] if 16 <= i <= 23 else None for i in range(1, 32)]])
    filas.append(["1012", maria, "10", "UBI 10", "10-1", "PUESTO 10-1", *[trabajo[i - 24] if 24 <= i <= 30 else None for i in range(1, 32)]])
    return xlsx(filas)


COLS = ["HORAS_100", "DEVENGO_100", "DEVENGO_103", "DEVENGO_129", "DEVENGO_130", "DEVENGO_132", "DEDUCCION_600", "DEDUCCION_605", "DEDUCCION_611"]
TITULOS = {"HORAS_100": "100-SALARIO BASICO", "DEVENGO_103": "103-AUXILIO DE TRANSPORTE", "DEVENGO_129": "129-RODAMIENTO", "DEVENGO_130": "130-HORAS EXTRAS Y RECARGOS",
           "DEVENGO_132": "132-APLICACION LEY 2101", "DEDUCCION_600": "600-SEGUROS", "DEDUCCION_605": "605-PRESTAMO",
           "DEDUCCION_611": "611-EMBARGO 5A PARTE"}


def _nomina(personas: dict[str, dict[str, float]]) -> bytes:
    fijas = ["Tercero", "Descripción", "Periodo", "Centro Costos", "Centro Operación", "Unidad Negocio", "Salario básico", "Descripción del cargo"]
    filas: list[list[object]] = [[""] * len(fijas) + [TITULOS.get(c, "") for c in COLS] + [""], fijas + COLS + ["NETO"]]
    for ced, v in personas.items():
        dev = sum(v.get(c, 0) for c in COLS if c.startswith("DEVENGO"))
        ded = sum(v.get(c, 0) for c in COLS if c.startswith("DEDUCCION"))
        filas.append([ced, next(c[1] for c in CONTRATOS if c[0] == ced), 202609, "400010", "001", "01", 1750905, "VIGILANTE",
                      *[v.get(c, 0) for c in COLS], v.get("NETO", dev - ded)])
    return xlsx(filas)


def _completo(dias: int, aux: bool = True, v130: float = 20000, v132: float = 3000) -> dict[str, float]:
    return {"HORAS_100": dias * 7, "DEVENGO_100": round(1750905 / 30 * dias), "DEVENGO_103": round(DIA_AUX * dias) if aux else 0,
            "DEVENGO_130": dias * v130, "DEVENGO_132": dias * v132}


QUINCENAL = {
    # Ana: todo correcto; embargo = (875.453 + 300.000 − ½ SMLMV) × 20 % (sin 103 ni 132)
    "1001": {**_completo(15), "DEDUCCION_611": round((round(1750905 / 2) + 300000 - 1750905 / 2) * 0.2)},
    "1002": {**_completo(15), "DEDUCCION_605": 100000},  # 5 días de vacaciones y se le paga todo
    "1003": {**_completo(15, aux=False), "DEVENGO_130": 400000, "DEDUCCION_605": 100000},
    "1004": {"DEDUCCION_605": 50000, "NETO": -50000},  # sin programación, sin sueldo y neto negativo
    "1006": {**_completo(15), "DEVENGO_129": 600000},  # administrativo sin programación; rodamiento pagado como 129
    "1007": {"DEVENGO_130": 999999},  # de "los 7": no se revisa
    "1008": {**_completo(15, v130=0, v132=0)},  # puesto sin modalidad
    "1009": {**_completo(15, v130=30000, v132=0), "DEDUCCION_605": 50000},  # modalidad del puesto; cuota inactiva descontada
}
QUINCENAL["1010"] = {}  # nada en la nómina: lo trabajado antes va en la liquidación de vacaciones
QUINCENAL["1011"] = _completo(10)  # se le pagan solo los 10 días trabajados después del regreso
# María: 130 = 8 × 20.000 + 7 × 30.000; 132 solo en los 8 días del 4X2 (el 6X1 no lo tiene)
QUINCENAL["1012"] = {**_completo(15), "DEVENGO_130": 8 * 20000 + 7 * 30000, "DEVENGO_132": 8 * 3000}
MENSUAL = {"2001": _completo(30)}


def _cargar_todo(cliente, nomina: str = "quincenal") -> int:
    pid = cliente.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "nomina": nomina}).json()["data"]["id"]
    for tipo, contenido in [("modalidades", _modalidades()), ("ubicaciones", _ubicaciones()), ("contratos", _contratos()),
                            ("cuotas", _cuotas()), ("programacion", _programacion()),
                            (nomina, _nomina(QUINCENAL if nomina == "quincenal" else MENSUAL))]:
        r = cliente.post(f"{API}/periodos/{pid}/archivos/{tipo}", files=archivo(contenido, f"{tipo}.xlsx"))
        assert r.status_code == 200, (tipo, r.text)
    return pid


def _alertas(cliente, pid) -> dict[str, set[tuple[str, str]]]:
    datos = cliente.get(f"{API}/periodos/{pid}/alertas?tamano=150").json()["data"]
    res: dict[str, set] = {}
    for a in datos:
        res.setdefault(a["cedula"], set()).add((a["tipo"], a["referencia"]))
    return res


def test_escenario_completo(admin):
    pid = _cargar_todo(admin)
    p = admin.get(f"{API}/periodos/{pid}").json()["data"]
    assert p["calculado_en"] and p["faltan"] == []
    assert p["nombre"] == "Quincenal · 2.ª quincena de septiembre 2026"
    assert p["resumen"]["personas"] == {"quincenal": 10}  # Gina (los 7) queda fuera
    assert p["resumen"]["excluidas"] == {"quincenal": 1}
    assert p["historial"][0]["n"] == 1 and p["historial"][0]["alertas"] == p["resumen"]["alertas"]

    a = _alertas(admin, pid)
    assert "1001" not in a and "2001" not in a  # Juan es de nómina mensual: no se espera en la quincenal  # todo correcto (incluye la cuota 600, que solo va en la 1.ª quincena)
    # Beto trabajó del 16 al 25 y salió a vacaciones: esos días van en la liquidación de vacaciones, pero la nómina le pagó todo
    assert {("VACACIONES_CON_PAGO", "100"), ("VACACIONES_CON_PAGO", "130"), ("VACACIONES_CON_PAGO", "103"),
            ("CUOTA_MAESTRO", "605"), ("VACACIONES_VERIFICAR", "")} <= a["1002"]
    # Karen (trabajó y salió a vacaciones, nómina en cero) y Luis (regresó y se le pagan los días después): sin alertas reales
    assert a["1010"] == {("VACACIONES_VERIFICAR", "")} and a["1011"] == {("VACACIONES_VERIFICAR", "")}
    assert p["resumen"]["vacaciones"] == 3
    assert ("CUOTA_DE_MENOS", "605") not in a["1002"]  # el tope (100.000) es lo esperado y fue lo descontado
    assert {("AUX_CERO", "103"), ("MODALIDAD_MAYOR", "130"), ("CUOTA_DE_MAS", "605")} <= a["1003"]
    assert {("NOMINA_SIN_PROGRAMACION", ""), ("DESCUENTO_SIN_SUELDO", ""), ("NETO_NEGATIVO", "")} <= a["1004"]
    assert a["1005"] == {("PROGRAMADO_SIN_NOMINA", "")}
    assert a["1006"] == {("CUOTA_NO_DESCONTADA", "605")}  # administrativo: solo cuotas
    assert "1007" not in a and "1008" not in a
    # Iván: puesto con modalidad propia (6X1, 30.000/día) aunque la ubicación tenga otra; 15 días (XYZ no descuenta)
    # Préstamo inactivo que se sigue descontando y rodamiento (152) que no se pagó con el 129
    assert a["1009"] == {("CUOTA_DE_MAS", "605"), ("CUOTA_DEVENGO_DIFERENTE", "152")}


def test_detalle_de_persona_y_dias(admin):
    pid = _cargar_todo(admin)
    d = admin.get(f"{API}/periodos/{pid}/personas/quincenal/1002").json()["data"]
    assert d["conteo"] == {"pagables": 0, "novedad": 5, "vacaciones": 5, "vacios": 0, "sin_modalidad": 0, "antes_vacaciones": 10}
    assert d["vacaciones"] == {"desde": "2026-09-26", "hasta": "2026-09-30", "dias": 5, "antes": 10}
    assert d["dias_salario"] == 15 and d["auxilio"]["dias_esperados"] == 0
    c130 = next(x for x in d["conceptos"] if x["concepto"] == "130")
    assert c130 == {"concepto": "130", "descripcion": "HORAS EXTRAS Y RECARGOS", "esperado": 0, "pagado": 300000}
    assert d["dias"]["2026-09-26"]["clase"] == "vacaciones" and d["dias"]["2026-09-16"]["clase"] == "antes_vacaciones"
    luis = admin.get(f"{API}/periodos/{pid}/personas/quincenal/1011").json()["data"]
    assert luis["conteo"]["pagables"] == 10 and luis["conteo"]["antes_vacaciones"] == 0
    assert {x["tipo"] for x in d["alertas_lista"]} >= {"VACACIONES_CON_PAGO"}
    iv = admin.get(f"{API}/periodos/{pid}/personas/quincenal/1009").json()["data"]
    assert iv["modalidades"] == {"6X1 ESP": 15}


def test_puestos_sin_modalidad_y_aprobacion(admin):
    pid = _cargar_todo(admin)
    lista = admin.get(f"{API}/periodos/{pid}/sin-modalidad").json()["data"]
    p20 = next(x for x in lista if x["puesto"] == "20")
    assert p20["motivo"] == "sin_modalidad" and p20["dias"] == 15 and p20["estado"] == "pendiente"
    assert admin.put(f"{API}/puestos-sin-modalidad?periodo_id={pid}", json={"ubicacion": "20", "puesto": "20", "estado": "error"}).status_code == 422
    r = admin.put(f"{API}/puestos-sin-modalidad?periodo_id={pid}", json={"ubicacion": "20", "puesto": "20", "estado": "aprobado", "comentario": "Recepción sin modalidad"})
    assert r.status_code == 200
    p20 = next(x for x in admin.get(f"{API}/periodos/{pid}/sin-modalidad").json()["data"] if x["puesto"] == "20")
    assert p20["estado"] == "aprobado" and p20["aprobado"] is True


def test_decisiones_se_conservan_al_recalcular(admin):
    pid = _cargar_todo(admin)
    item = {"nomina": "quincenal", "cedula": "1003", "tipo": "AUX_CERO", "referencia": "103"}
    assert admin.put(f"{API}/periodos/{pid}/alertas/decision", json={"items": [item], "estado": "justificada"}).status_code == 422
    r = admin.put(f"{API}/periodos/{pid}/alertas/decision", json={"items": [item], "estado": "justificada", "comentario": "Vive cerca"})
    assert r.status_code == 200
    assert admin.post(f"{API}/periodos/{pid}/recalcular").status_code == 200
    lista = admin.get(f"{API}/periodos/{pid}/alertas?tipo=AUX_CERO").json()["data"]
    assert lista[0]["estado"] == "justificada" and lista[0]["comentario"] == "Vive cerca"
    pendientes = admin.get(f"{API}/periodos/{pid}/alertas?estado=pendiente&tamano=150").json()
    assert all(x["tipo"] != "AUX_CERO" for x in pendientes["data"])


def test_codigos_nuevos_y_check_descuenta(admin):
    pid = _cargar_todo(admin)
    r = admin.get(f"{API}/codigos?por_revisar=true").json()
    assert [c["codigo"] for c in r["data"]] == ["XYZ"] and r["meta"]["extra"]["por_revisar"] == 1
    # Si XYZ descuenta, Iván tiene 14 días y se le pagan 15: aparece la alerta
    admin.put(f"{API}/codigos/XYZ", json={"descripcion": "Prueba", "descuenta": True})
    admin.post(f"{API}/periodos/{pid}/recalcular")
    assert ("SALARIO_MAYOR_DIAS", "100") in _alertas(admin, pid)["1009"]
    vac = next(c for c in admin.get(f"{API}/codigos?q=VAC").json()["data"] if c["codigo"] == "VAC")
    assert vac["descuenta"] and vac["vacaciones"]


def test_grupos_y_parametros(admin):
    pid = _cargar_todo(admin)
    grupos = {g["nombre"]: g["tratamiento"] for g in admin.get(f"{API}/grupos").json()["data"]}
    assert grupos["OPERATIVOS FULL"] == "excluido" and grupos["ADMINISTRATIVOS/COMERCIALES"] == "administrativo"
    # Con tolerancia alta, la diferencia de modalidad de Carla (100.000) deja de ser alerta
    par = admin.get(f"{API}/parametros").json()["data"]
    assert par["tolerancia"] == 1000 and par["solo_primera_quincena"] == "600,630,631,632"
    admin.put(f"{API}/parametros", json={**par, "tolerancia": 150000})
    admin.post(f"{API}/periodos/{pid}/recalcular")
    assert ("MODALIDAD_MAYOR", "130") not in _alertas(admin, pid)["1003"]


def test_faltan_archivos_y_validacion_de_formato(admin):
    pid = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "nomina": "quincenal"}).json()["data"]["id"]
    r = admin.post(f"{API}/periodos/{pid}/archivos/modalidades", files=archivo(_modalidades()))
    assert r.status_code == 200 and r.json()["data"]["calculado"] is False
    faltan = r.json()["data"]["periodo"]["faltan"]
    assert "Contratos" in faltan and "Maestro de cuotas" in faltan and "Nómina quincenal" in faltan  # se exigen todos
    r = admin.post(f"{API}/periodos/{pid}/archivos/contratos", files=archivo(_modalidades()))
    assert r.status_code == 422 and "contratos" in r.json()["error"]["message"]
    r = admin.post(f"{API}/periodos/{pid}/archivos/mensual", files=archivo(_nomina(MENSUAL)))
    assert r.status_code == 422 and "quincenal" in r.json()["error"]["message"]
    assert admin.post(f"{API}/periodos/{pid}/recalcular").status_code == 422
    assert admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "nomina": "quincenal"}).status_code == 409
    assert admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 9, "nomina": "mensual"}).status_code == 201


def test_informe_excel(admin):
    pid = _cargar_todo(admin)
    r = admin.get(f"{API}/periodos/{pid}/informe")
    assert r.status_code == 200 and r.headers["content-disposition"] == 'attachment; filename="validacion_nomina_quincenal_2026-09.xlsx"'
    libro = load_workbook(BytesIO(r.content))
    assert libro.sheetnames == ["Resumen", "Revisiones", "Alertas", "Puestos sin modalidad"]
    assert libro["Alertas"].max_row > 10


def test_permisos(admin):
    admin.post("/api/roles", json={"nombre": "Solo ver nomina", "permisos": ["nomina.ver"]})
    rid = next(r["id"] for r in admin.get("/api/roles").json()["data"] if r["nombre"] == "Solo ver nomina")
    admin.post("/api/usuarios", json={"accesos": ACCESOS, "username": "vernom", "nombre": "Ver", "password": "clave-segura-1", "roles": [rid]})
    pid = _cargar_todo(admin)
    admin.post("/api/auth/logout")
    entrar(admin, "vernom", "clave-segura-1")
    assert admin.get(f"{API}/periodos/{pid}/alertas").status_code == 200
    assert admin.post(f"{API}/periodos/{pid}/recalcular").status_code == 403
    assert admin.put(f"{API}/periodos/{pid}/alertas/decision",
                     json={"items": [{"nomina": "quincenal", "cedula": "1003", "tipo": "AUX_CERO"}], "estado": "revisada"}).status_code == 403
    assert admin.put(f"{API}/codigos/Z", json={"descuenta": True}).status_code == 403


def test_periodo_mensual(admin):
    pid = _cargar_todo(admin, "mensual")
    p = admin.get(f"{API}/periodos/{pid}").json()["data"]
    assert p["resumen"]["personas"] == {"mensual": 1} and p["nombre"] == "Mensual · septiembre 2026"
    assert _alertas(admin, pid) == {}  # Juan: 30 días pagados y modalidad × 30; los quincenales no se esperan aquí


def test_recargar_la_nomina_muestra_lo_corregido(admin):
    pid = _cargar_todo(admin)
    antes = _alertas(admin, pid)
    assert ("AUX_CERO", "103") in antes["1003"] and ("MODALIDAD_MAYOR", "130") in antes["1003"]
    # Se corrige en SIESA el auxilio, la modalidad y el préstamo de Carla, pero aparece un error nuevo con Ana
    corregida = {**QUINCENAL, "1003": {**_completo(15), "DEDUCCION_605": 30000},
                 "1001": {**QUINCENAL["1001"], "DEVENGO_103": 0}}
    r = admin.post(f"{API}/periodos/{pid}/archivos/quincenal", files=archivo(_nomina(corregida), "nomina_v2.xlsx"))
    assert r.status_code == 200 and r.json()["data"]["calculado"] is True
    ultima = admin.get(f"{API}/periodos/{pid}").json()["data"]["historial"][0]
    assert ultima["n"] == 2 and ultima["motivo"] == "Nueva carga de la nómina" and ultima["archivo_nomina"] == "nomina_v2.xlsx"
    assert ultima["corregidas"] == 3 and ultima["nuevas"] == 1  # Carla: auxilio, modalidad y préstamo; Ana: auxilio en 0
    assert ultima["persisten"] == ultima["alertas"] - 1
    assert "1003" not in _alertas(admin, pid)
    nuevas = admin.get(f"{API}/periodos/{pid}/alertas?solo_nuevas=true").json()["data"]
    assert [(x["cedula"], x["tipo"], x["nueva"]) for x in nuevas] == [("1001", "AUX_CERO", True)]


def test_vacaciones_no_cuentan_como_pendientes(admin):
    pid = _cargar_todo(admin)
    p = admin.get(f"{API}/periodos/{pid}").json()["data"]
    todas = admin.get(f"{API}/periodos/{pid}/alertas?tamano=150").json()["meta"]["extra"]["total"]
    assert p["pendientes"] == todas - p["resumen"]["vacaciones"]  # la lista de vacaciones es para verificar, no son pendientes
    sin_vac = admin.get(f"{API}/periodos/{pid}/alertas?excluir_tipo=VACACIONES_VERIFICAR&tamano=150").json()["data"]
    assert all(x["tipo"] != "VACACIONES_VERIFICAR" for x in sin_vac)
    vac = admin.get(f"{API}/periodos/{pid}/alertas?tipo=VACACIONES_VERIFICAR").json()["data"]
    karen = next(x for x in vac if x["cedula"] == "1010")
    assert "8 días trabajados antes (del 16 al 23)" in karen["mensaje"] and karen["severidad"] == "info"


def test_cambio_de_puesto_paga_la_modalidad_de_cada_puesto(admin):
    pid = _cargar_todo(admin)
    assert "1012" not in _alertas(admin, pid)  # pagada con la modalidad de cada puesto por sus días
    d = admin.get(f"{API}/periodos/{pid}/personas/quincenal/1012").json()["data"]
    assert d["modalidades"] == {"4X2 1 GRP": 8, "6X1 ESP": 7}
    esperado = {x["concepto"]: x["esperado"] for x in d["conceptos"]}
    assert esperado == {"130": 370000, "132": 24000}
    puestos = {x["puesto"]: x for x in d["puestos"]}
    assert puestos["10"]["dias_pagables"] == 8 and puestos["10"]["origen"] == "ubicacion"
    assert puestos["10-1"]["dias_pagables"] == 7 and puestos["10-1"]["origen"] == "puesto"
    assert puestos["10-1"]["conceptos"] == {"130": 30000}
    assert d["dias"]["2026-09-23"]["modalidad"] == "4X2 1 GRP" and d["dias"]["2026-09-24"]["modalidad"] == "6X1 ESP"
