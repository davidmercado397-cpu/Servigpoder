"""Umbral semanal de horas del liquidador: se marca por periodo y ajusta las extras por semana (lunes a domingo).

Octubre 2026, quincena 1 (jueves 1 al jueves 15). Semanas: 28-sep a 4-oct (4 días en el periodo → umbral
42 × 4/7 = 24), 5 a 11-oct (7 días → 42). Turno D: 06-18, 8,4 ordinarias + 2,8 extras diurnas; N: 18-06, igual
pero de noche; INC: incapacidad (novedad justificada, baja 7 h el umbral)."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.conftest import xlsx
from tests.test_liquidador import API, _cargar, _quincena, _resultados


def _excel() -> bytes:
    dias = list(range(1, 16))
    turnos = {
        "5001": {5: "D", 6: "D", 7: "D", 8: "D"},  # 44,8 h en la semana (11,2 extras) → se pagan 2,8
        "5002": {5: "N", 6: "N", 7: "N", 8: "N"},  # igual, de noche: lo descontado conserva el recargo nocturno
        "5003": {1: "D", 2: "D"},  # semana partida: umbral 24, trabaja 22,4 → no se pagan extras
        "5004": {5: "D", 6: "D", 7: "D", 8: "D", 9: "INC"},  # incapacidad: umbral 35 → se pagan 9,8
        "5005": {5: "D", 6: "D", 7: "D", 8: "D", 9: "D", 10: "Z", 11: "L"},  # 56 h, Z y L no bajan el umbral → 11,2 + 2,8 = 14 pagadas
    }
    filas = [[doc, doc, *[t.get(d, "") for d in dias]] for doc, t in turnos.items()]
    return xlsx([["documento", "OCTUBRE", *dias], *filas])


def _horas(cliente, pid, doc) -> tuple[dict, list]:
    r = _resultados(cliente, pid)[doc]
    det = cliente.get(f"{API}/periodos/{pid}/empleados/{r['empleado_id']}").json()["data"]["empleado"]
    return r["horas"], det["umbral_semanas"]


def test_umbral_semanal(admin):
    pid = _quincena(admin, 2026, 10, 1)
    # Se marca antes de subir el Excel: queda guardado y aplica a la carga
    r = admin.put(f"{API}/periodos/{pid}/umbral", json={"activo": True, "horas": 42, "horas_novedad": 7})
    assert r.status_code == 200 and r.json()["data"]["umbral_activo"] is True
    assert _cargar(admin, pid, _excel()).status_code == 200

    h, semanas = _horas(admin, pid, "5001")
    s = next(x for x in semanas if x["desde"] == "2026-10-05")
    assert (s["umbral"], s["trabajadas"], s["extras"], s["pagadas"], s["descontadas"]) == (42, pytest.approx(44.8), pytest.approx(11.2),
                                                                                          pytest.approx(2.8), pytest.approx(8.4))
    assert h["overtime_day"] == pytest.approx(2.8) and h["ordinary_day"] == pytest.approx(33.6 + 8.4)

    # Extras nocturnas descontadas: pasan a ordinarias pero conservan el recargo nocturno
    h_on, _ = _horas(admin, pid, "5002")
    admin.put(f"{API}/periodos/{pid}/umbral", json={"activo": False, "horas": 42, "horas_novedad": 7})
    h_off, sem_off = _horas(admin, pid, "5002")
    assert sem_off == [] and h_off["overtime_night"] == pytest.approx(11.2)
    assert h_on["overtime_night"] == pytest.approx(2.8)
    assert h_on["ordinary_night"] == pytest.approx(h_off["ordinary_night"] + 8.4)
    assert h_on["ordinary_day"] == pytest.approx(h_off["ordinary_day"] + 8.4)
    admin.put(f"{API}/periodos/{pid}/umbral", json={"activo": True, "horas": 42, "horas_novedad": 7})

    # Semana partida (4 días del periodo): umbral 24
    h, semanas = _horas(admin, pid, "5003")
    assert semanas[0]["dias_periodo"] == 4 and semanas[0]["umbral"] == 24 and semanas[0]["pagadas"] == 0
    assert h.get("overtime_day", 0) == 0

    # Novedad justificada: 7 h menos de umbral
    _, semanas = _horas(admin, pid, "5004")
    s = semanas[0]
    assert (s["justificados"], s["umbral"], s["pagadas"]) == (1, 35, pytest.approx(9.8))

    # Descansos Z y L no bajan el umbral
    _, semanas = _horas(admin, pid, "5005")
    assert semanas[0]["justificados"] == 0 and semanas[0]["pagadas"] == pytest.approx(14)

    # La nómina paga con las horas ajustadas: 2,8 extras diurnas al 125 %
    n = admin.get(f"{API}/periodos/{pid}/empleados/{_resultados(admin, pid)['5001']['empleado_id']}").json()["data"]["empleado"]["nomina"]
    assert n["valores"]["overtime_day"] == round(2.8 * 1750905 / 210 * 1.25)
    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/nomina").content))
    assert "Umbral semanal" in libro.sheetnames and libro["Umbral semanal"].max_row >= 7
    assert load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/liquidacion").content)).sheetnames == ["2026-10-Q1"]  # no cambia
    eventos = admin.get("/api/plataforma/auditoria?accion=liq_umbral_semanal").json()["data"]
    assert len(eventos) == 3 and eventos[0]["detalle"]["recalculado"] is True


def test_umbral_no_se_cambia_con_el_periodo_cerrado(admin):
    pid = _quincena(admin, 2026, 10, 1)
    _cargar(admin, pid, _excel())
    assert admin.post(f"{API}/periodos/{pid}/cerrar").status_code == 200
    assert admin.put(f"{API}/periodos/{pid}/umbral", json={"activo": True, "horas": 42}).status_code == 409
