"""Pagado por mes: suma de las nóminas cerradas (pagadas) de un mes, por nómina y por persona. Las nóminas abiertas
se listan aparte y no suman."""

from collections import defaultdict

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.apps.liquidador.models import CERRADA, LiqPeriodo, LiqResultado
from app.apps.liquidador.services.exportar import XLSX, _FUENTE, _RELLENO, _CENTRO, _autoajustar, _guardar

CAMPOS = (
    ("basico", "Sueldo básico"), ("incapacidad", "Incapacidad"), ("vacaciones_valor", "Vacaciones"), ("recargos", "Recargos"),
    ("extras", "Horas extras"), ("auxilios_no_prestacionales", "Auxilios no prestacionales"), ("auxilio", "Auxilio de transporte"),
    ("novedades_devengo", "Novedades (devengo)"), ("devengado", "Total devengado"), ("salud", "Salud"), ("pension", "Pensión"), ("embargos", "Embargos"),
    ("prestamos", "Préstamos"), ("cuotas", "Cuotas"), ("novedades_deduccion", "Novedades (deducción)"), ("deducciones", "Total deducciones"), ("neto", "Neto pagado"),
)


def _suma(destino: dict, nomina: dict) -> None:
    for k, _ in CAMPOS:
        destino[k] = destino.get(k, 0) + float(nomina.get(k, 0) or 0)


def pagado(db: Session, anio: int, mes: int) -> dict:
    nominas = list(db.scalars(select(LiqPeriodo).where(LiqPeriodo.anio == anio, LiqPeriodo.mes == mes)
                              .order_by(LiqPeriodo.quincena, LiqPeriodo.id)))
    cerradas = {p.id for p in nominas if p.estado == CERRADA}
    por_nomina: dict[int, dict] = {p.id: {} for p in nominas}
    personas: dict[int, dict] = {}
    conteo: dict[int, int] = defaultdict(int)
    if nominas:
        for r in db.scalars(select(LiqResultado).where(LiqResultado.periodo_id.in_([p.id for p in nominas]))):
            conteo[r.periodo_id] += 1
            _suma(por_nomina[r.periodo_id], r.nomina or {})
            if r.periodo_id not in cerradas:
                continue
            x = personas.setdefault(r.empleado_id, {"empleado_id": r.empleado_id, "documento": r.empleado.documento,
                                                    "nombre": r.empleado.nombre, "nominas": []})
            _suma(x, r.nomina or {})
            x["nominas"].append(next(f"{p.nombre} ({p.tipo})" for p in nominas if p.id == r.periodo_id))
    totales: dict = {}
    for pid in cerradas:
        for k, _ in CAMPOS:
            totales[k] = totales.get(k, 0) + por_nomina[pid].get(k, 0)
    return {
        "nominas": [{"id": p.id, "nombre": p.nombre, "tipo": p.tipo, "etiqueta": p.etiqueta, "estado": p.estado,
                     "cerrado_en": p.cerrado_en.isoformat() if p.cerrado_en else None, "personas": conteo[p.id],
                     "incluida": p.id in cerradas, **{k: round(v) for k, v in por_nomina[p.id].items()}} for p in nominas],
        "totales": {k: round(v) for k, v in totales.items()},
        "personas": sorted(({**x, **{k: round(x.get(k, 0)) for k, _ in CAMPOS}} for x in personas.values()),
                           key=lambda x: x["nombre"]),
        "abiertas": sum(1 for p in nominas if p.id not in cerradas),
    }


def excel(db: Session, anio: int, mes: int) -> tuple[bytes, str, str]:
    datos = pagado(db, anio, mes)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Por nomina"
    hoja.append(["NOMINA", "PERIODO", "ESTADO", "PERSONAS", *(t.upper() for _, t in CAMPOS)])
    for n in datos["nominas"]:
        hoja.append([n["nombre"], n["etiqueta"], "Pagada (cerrada)" if n["incluida"] else "Abierta: no suma", n["personas"],
                     *(n.get(k, 0) for k, _ in CAMPOS)])
    hoja.append(["TOTAL PAGADO", f"{anio}-{mes:02d}", "", sum(n["personas"] for n in datos["nominas"] if n["incluida"]),
                 *(datos["totales"].get(k, 0) for k, _ in CAMPOS)])
    personas = libro.create_sheet("Por persona")
    personas.append(["DOCUMENTO", "EMPLEADO", "NOMINAS", *(t.upper() for _, t in CAMPOS)])
    for x in datos["personas"]:
        personas.append([x["documento"], x["nombre"], ", ".join(x["nominas"]), *(x.get(k, 0) for k, _ in CAMPOS)])
    for h, inicio in ((hoja, 5), (personas, 4)):
        for celda in h[1]:
            celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
        for col in range(inicio, inicio + len(CAMPOS)):
            for (celda,) in h.iter_rows(min_row=2, min_col=col, max_col=col):
                celda.number_format = "$ #,##0"
        _autoajustar(h, minimo=10, maximo=36)
        h.freeze_panes = "C2"
    return _guardar(libro), XLSX, f"pagado_{anio}_{mes:02d}.xlsx"
