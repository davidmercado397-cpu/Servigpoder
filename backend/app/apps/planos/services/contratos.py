"""Plano de contratos SIESA (94 columnas, A–CP).

Tipos de regla: FIJO (se escribe siempre), DEFECTO (solo si viene vacío), COPIA (de la fecha de ingreso Z),
CATÁLOGO (entidades y cargo), CÁLCULO (salario de vigilantes, auxilio de transporte, forma de pago, fecha hasta)
y LIMPIEZA (cédula, cuenta, fechas). Las columnas sin regla se dejan como vienen.
"""

import re
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl.utils import get_column_letter

from app.apps.planos.services.comun import (
    PlanoInvalido, Resultado, a_fecha, abrir, col, fila_vacia, formato_texto, guardar, hoja_de_datos, limpiar_id,
    normalizar, normalizar_fecha, ordenar, quitar_filas, texto, ultima_fila,
)
from app.apps.planos.services.entidades import Codificador

COLUMNAS = 94
VIGILANTES = {"032", "033", "052"}  # vigilante, vigilante motorizado y vigilante líder
APRENDICES = {"12", "19"}  # tipo de cotizante: lectivo y productivo (pueden ganar menos del mínimo)
GRUPOS_VALIDOS = {"001", "003"}
SIN_FONDO_PENSIONES = "999999999"
FIJOS = {**{c: "1" for c in ("A", "C", "AU", "BA", "BT")}, "D": "001", "K": "001", "J": "01", "W": "14-4", "X": "PASENA",
         "Y": "PAICBF", "AQ": "210", "AS": "100",
         **{c: "0" for c in "AK AL AM AN AO AP AR AT AV BB BC BD BE BF BG BH BI BJ BK BL BN BP BR BS".split()}}
FIJOS_VIGILANTE = {"I": "999", "O": "406999", "P": "004"}
DEFECTOS = {"G": "1", "Q": "Q", "AY": "1", "BM": "1", "BO": "1"}
PERMITIDOS = {"Q": {"Q", "M"}, "BM": {"1", "2", "3"}}
ENTIDADES = {"F": "BANCOS", "S": "AFP", "T": "EPS", "U": "CESANTIAS", "V": "CAJAS", "BQ": "CARGOS"}
COPIAS_INGRESO = ["AA", "AD", "AE", "AF", "AG", "AH", "AI"]
TODAS = [get_column_letter(i) for i in range(1, COLUMNAS + 1)]
MAX_DIAS_CONTRATO = 366


def fecha_hasta(ingreso: str) -> str:
    """Ingreso + 3 meses, menos un día (16/03 → 15/06). Si el mes de llegada no tiene ese día, su último día."""
    d = a_fecha(ingreso)
    mes, anio = d.month + 3, d.year
    if mes > 12:
        mes, anio = mes - 12, anio + 1
    dia = d.day
    while True:
        try:
            fin = datetime(anio, mes, dia)
            break
        except ValueError:
            dia -= 1
    if dia == d.day:
        fin -= timedelta(days=1)
    return fin.strftime("%Y%m%d")


def leer_salario(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = texto(v).replace("$", "").replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d{1,2})?", s):  # 1.750.905 o 1.750.905,50
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d{1,2})?", s):  # 1,750,905.50
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def es_plano_contratos(ws) -> bool:
    return texto(ws.cell(1, 1).value).upper().startswith("ID CIA") and texto(ws.cell(1, 2).value).upper().startswith("ID TERCERO")


def procesar(contenido: bytes, codificador: Codificador, smlv: int, correcciones: dict[tuple[str, str], str] | None = None,
             empleados: dict[str, str] | None = None, fila_ejemplo: int | None = None,
             hoy: date | None = None) -> tuple[bytes, Resultado]:
    """empleados: {cédula: fecha de ingreso} del plano de empleados procesado en la misma sesión (cruce opcional)."""
    wb = abrir(contenido)
    ws = hoja_de_datos(wb)
    if not es_plano_contratos(ws):
        raise PlanoInvalido("La hoja no parece un plano de contratos: las columnas A y B deben ser «ID CIA» e «ID TERCERO»")
    hoy = hoy or date.today()
    res = Resultado()
    correcciones = correcciones or {}
    obligatorias = [c for c in TODAS if "(OBLIGATORIO" in re.sub(r"\(\s+", "(", texto(ws.cell(1, col(c)).value).upper())]
    por_cedula: dict[str, list[int]] = defaultdict(list)
    por_cuenta: dict[str, list[tuple[int, str]]] = defaultdict(list)
    ultima = ultima_fila(ws)
    borrar: list[int] = []

    for f in range(2, ultima + 1):
        if f == fila_ejemplo or fila_vacia(ws, f):
            borrar.append(f)
            continue
        res.sumar("filas")
        cel = lambda c: ws.cell(f, col(c))  # noqa: E731

        def poner(c: str, v: str, motivo: str = "") -> None:
            antes = cel(c).value
            if antes != v:
                cel(c).value = v
                if texto(antes) and texto(antes) != v:  # completar un vacío o quitar espacios no es un cambio
                    res.cambio(f"{c}: {texto(antes)} → {v}" + (f" ({motivo})" if motivo else ""))

        ced, problema = limpiar_id(cel("B").value)
        cel("B").value = ced or None
        if problema:
            res.agregar(f, ced, "B", problema)
        else:
            por_cedula[ced].append(f)
            if not 5 <= len(ced) <= 10:
                res.agregar(f, ced, "B", f"Cédula con {len(ced)} dígitos: revisar", "Advertencia")
        for (c_ced, c_col), valor in correcciones.items():
            if c_ced == ced:
                cel(c_col).value = valor
                res.sumar("correcciones")

        for c, v in FIJOS.items():
            poner(c, v)
        for c, v in DEFECTOS.items():
            if not texto(cel(c).value):
                cel(c).value = v
        for c, validos in PERMITIDOS.items():
            v = texto(cel(c).value).upper()
            cel(c).value = v
            if v not in validos:
                res.agregar(f, ced, c, f"Valor '{v}' no permitido (válidos: {', '.join(sorted(validos))})")

        # Entidades y cargo
        for c, catalogo in ENTIDADES.items():
            original = texto(cel(c).value)
            r = codificador.codificar(catalogo, cel(c).value)
            sugerencia = {"catalogo": catalogo, "pais": "", "departamento": "", "texto": normalizar(original), "codigo": r.codigo or ""}
            if r.codigo is None:
                if r.como == "vacio":
                    if c in obligatorias:
                        res.agregar(f, ced, c, f"{catalogo}: vacío y es obligatorio")
                elif r.como == "ambiguo":
                    res.agregar(f, ced, c, f"{catalogo}: '{original}' es ambiguo ({r.detalle}); escriba el código o cree un alias",
                                alias=sugerencia)
                else:
                    res.agregar(f, ced, c, f"{catalogo}: '{original}' no está en el catálogo" + (f" ({r.detalle})" if r.detalle else ""),
                                alias=sugerencia)
                continue
            cel(c).value = r.codigo
            if r.revisar:
                forma = {"ceros": "se completaron ceros", "alias": "alias", "contenido": "nombre parcial",
                         "similitud": f"parecido {r.detalle}"}[r.como]
                res.agregar(f, ced, c, f"{catalogo}: '{original}' → {r.codigo} {codificador.descripcion(catalogo, r.codigo)} ({forma})",
                            "Revisar" if r.como in ("contenido", "similitud") else "Informativo",
                            sugerencia if r.como in ("contenido", "similitud") else None)
        vigilante = texto(cel("BQ").value) in VIGILANTES

        # Reglas de vigilantes; en los demás cargos se respeta lo que viene
        if vigilante:
            if texto(cel("P").value) == "404":
                poner("P", "004", "error de digitación")
            for c, v in FIJOS_VIGILANTE.items():
                poner(c, v, "vigilante")
            poner("AJ", str(smlv), "vigilante: salario mínimo")
        else:
            for c in FIJOS_VIGILANTE:
                if not texto(cel(c).value):
                    res.agregar(f, ced, c, "Vacío: en cargos que no son vigilante no hay valor por defecto")
        grupo = texto(cel("N").value)
        if grupo and grupo not in GRUPOS_VALIDOS:
            res.agregar(f, ced, "N", f"Grupo de empleados '{grupo}' distinto de 001 y 003", "Advertencia")

        # Fechas
        ingreso, ok = normalizar_fecha(cel("Z").value)
        cel("Z").value = ingreso or None
        if not ok:
            res.agregar(f, ced, "Z", f"Fecha de ingreso inválida o vacía: '{ingreso}'" if ingreso else "Fecha de ingreso vacía")
            ingreso = ""
        else:
            if ingreso[:4] != str(hoy.year):
                res.agregar(f, ced, "Z", f"Fecha de ingreso {ingreso} fuera del año en curso: solo se crean contratos de {hoy.year}")
            for c in COPIAS_INGRESO:
                poner(c, ingreso)
        for c in ("AC", "CD"):
            v, ok = normalizar_fecha(cel(c).value)
            if not v:
                if c == "AC" and ingreso:
                    cel(c).value = fecha_hasta(ingreso)
                    res.sumar("fechas_hasta_calculadas")
                continue
            cel(c).value = v
            if not ok:
                res.agregar(f, ced, c, f"Fecha inválida: '{v}'")
            elif c == "AC" and ingreso:
                if v < ingreso:
                    res.agregar(f, ced, "AC", f"Contrato hasta ({v}) anterior al ingreso ({ingreso})")
                elif (a_fecha(v) - a_fecha(ingreso)).days > MAX_DIAS_CONTRATO:
                    res.agregar(f, ced, "AC", f"Contrato de más de un año (ingreso {ingreso}, hasta {v}): revisar", "Advertencia")

        # Salario y auxilio de transporte (3 = hasta 2 SMLV; 0 = no tiene derecho)
        salario = leer_salario(cel("AJ").value)
        if salario is None:
            res.agregar(f, ced, "AJ", f"Salario inválido o vacío: '{texto(cel('AJ').value)}'")
        else:
            cel("AJ").value = str(int(salario)) if salario.is_integer() else f"{salario:.2f}"
            if salario < smlv and texto(cel("G").value) not in APRENDICES:
                res.agregar(f, ced, "AJ", f"Salario {salario:,.0f} menor al salario mínimo ({smlv:,})".replace(",", "."))
            hasta_dos = salario <= 2 * smlv
            poner("AZ", "3" if hasta_dos else "0", "salario hasta 2 SMLV" if hasta_dos else "salario mayor a 2 SMLV")

        # Cuenta bancaria y forma de pago
        cuenta = re.sub(r"[\s*%$/.,\-]", "", texto(cel("AW").value))
        cel("AW").value = cuenta or None
        if cuenta:
            por_cuenta[cuenta].append((f, ced))
        poner("AX", "2" if cuenta else "0")

        if texto(cel("S").value) == SIN_FONDO_PENSIONES:
            res.agregar(f, ced, "AY", "AFP «sin fondo de pensiones»: posible pensionado, revisar el régimen laboral", "Advertencia")
        if empleados is not None and ced and not problema:
            if ced not in empleados:
                res.agregar(f, ced, "B", "La cédula no está en el plano de empleados procesado", "Advertencia")
            elif ingreso and empleados[ced] and empleados[ced] != ingreso:
                res.agregar(f, ced, "Z", f"Fecha de ingreso ({ingreso}) distinta a la del plano de empleados ({empleados[ced]})",
                            "Advertencia")
        for c in obligatorias:
            if c not in ENTIDADES and c not in FIJOS_VIGILANTE and c not in ("B", "Z", "AJ") and not texto(cel(c).value):
                res.agregar(f, ced, c, "Columna obligatoria vacía")

    for ced, filas in por_cedula.items():
        if len(filas) > 1:
            for f in filas:
                res.agregar(f, ced, "B", f"Cédula repetida en el plano (filas {', '.join(map(str, filas))})")
    for cuenta, filas in por_cuenta.items():
        if len({c for _, c in filas}) > 1:
            for f, ced in filas:
                res.agregar(f, ced, "AW", f"Cuenta bancaria repetida en otra cédula (filas {', '.join(str(x) for x, _ in filas)})")

    formato_texto(ws, TODAS, ultima)
    quitar_filas(ws, ultima, borrar)
    res.resumen["filas_eliminadas"] = len(borrar)
    res.hallazgos = ordenar(res.hallazgos)
    return guardar(wb, res.hallazgos), res
