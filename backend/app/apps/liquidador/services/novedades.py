"""Novedades manuales de nómina: plantilla de Excel e importación (el archivo se procesa y se descarta).

Plantilla: una fila por novedad con la cédula, el tipo, la descripción, el valor o las horas por concepto (horas extras y
recargos pendientes), la fecha en que se causaron y si es salarial (solo «otro devengo»).
Importar reemplaza las novedades cargadas antes por Excel en esa nómina; las registradas a mano se conservan.
"""

from openpyxl import Workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.apps.liquidador.models import (
    AUXILIO_VARIABLE, HORAS_EXTRAS, MAYOR_VALOR, MENOR_VALOR, NOVEDADES_HORAS, OTRA_DEDUCCION, OTRO_DEVENGO, RECARGOS_PENDIENTES, LiqNovedad,
    LiqPeriodo, LiqResultado,
)
from app.apps.liquidador.services.exportar import (
    ENCABEZADOS, ORDEN_CONCEPTOS, XLSX, _CENTRO, _FUENTE, _RELLENO, _autoajustar, _guardar, _titulo, sufijo,
)
from app.apps.liquidador.services.personal import ErrorArchivo, Indice, _abrir, _fecha, _norm, _texto, documento
from app.apps.liquidador.dominio.tipos import HourConcept

TIPOS = {AUXILIO_VARIABLE: "Auxilio no prestacional variable", MAYOR_VALOR: "Mayor valor pagado", MENOR_VALOR: "Menor valor pagado", HORAS_EXTRAS: "Horas extras pendientes",
         RECARGOS_PENDIENTES: "Recargos pendientes", OTRO_DEVENGO: "Otro devengo", OTRA_DEDUCCION: "Otra deducción"}
# Concepto de horas → encabezado de la plantilla (los mismos nombres del archivo de liquidación)
COLUMNAS_HORAS = tuple((c.value, ENCABEZADOS[2 + i]) for i, c in enumerate(ORDEN_CONCEPTOS) if c != HourConcept.ORDINARY_DAY)
ENCABEZADO = ("CEDULA", "NOMBRE", "TIPO", "DESCRIPCION", "VALOR", "FECHA CAUSACION", "SALARIAL",
              *(t for _, t in COLUMNAS_HORAS))


def tipo_de(texto: object) -> str | None:
    t = _norm(texto)
    for clave, tipo in (("AUXILIO", AUXILIO_VARIABLE), ("RODAMIENTO", AUXILIO_VARIABLE), ("INTERMUNICIPAL", AUXILIO_VARIABLE), ("MAYOR", MAYOR_VALOR), ("MENOR", MENOR_VALOR), ("EXTRA", HORAS_EXTRAS), ("RECARGO", RECARGOS_PENDIENTES),
                        ("DEDUCCION", OTRA_DEDUCCION), ("DESCUENTO", OTRA_DEDUCCION), ("DEVENGO", OTRO_DEVENGO), ("PAGO", OTRO_DEVENGO)):
        if clave in t:
            return tipo
    return t.lower() if t.lower() in TIPOS else None


def plantilla(periodo: LiqPeriodo) -> tuple[bytes, str, str]:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Novedades"
    hoja.append(list(ENCABEZADO))
    for celda in hoja[1]:
        celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
    hoja.append([""] * len(ENCABEZADO))
    _autoajustar(hoja, minimo=12, maximo=26)
    hoja.freeze_panes = "C2"
    hoja.row_dimensions[1].height = 42
    ayuda = libro.create_sheet("Instrucciones")
    for fila in (
        ["TIPO", "EFECTO", "QUÉ LLENAR"],
        ["AUXILIO NO PRESTACIONAL VARIABLE", "Se paga como auxilio no prestacional (p. ej. rodamiento intermunicipal); no hace base salvo el exceso del 40 %", "VALOR"],
        ["MAYOR VALOR PAGADO", "Se descuenta del neto (no cambia salud y pensión)", "VALOR"],
        ["MENOR VALOR PAGADO", "Se paga; hace base de salud y pensión", "VALOR"],
        ["HORAS EXTRAS PENDIENTES", "Se paga; hace base de salud y pensión", "Horas en las columnas de cada concepto y FECHA CAUSACION (la tarifa de esa fecha). VALOR opcional: reemplaza el calculado"],
        ["RECARGOS PENDIENTES", "Se paga; hace base de salud y pensión", "Igual que horas extras"],
        ["OTRO DEVENGO", "Se paga; SALARIAL = SI hace base de salud y pensión", "VALOR y SALARIAL (SI/NO)"],
        ["OTRA DEDUCCION", "Se descuenta del neto", "VALOR"],
        [],
        ["La cédula debe estar en esta nómina. Importar reemplaza las novedades cargadas antes por Excel; las registradas a mano se conservan."],
    ):
        ayuda.append(fila)
    for celda in ayuda[1]:
        celda.fill, celda.font, celda.alignment = _RELLENO, _FUENTE, _CENTRO
    _autoajustar(ayuda, minimo=14, maximo=70)
    return _guardar(libro), XLSX, f"novedades_{_titulo(periodo)}{sufijo(periodo)}.xlsx"


def importar(db: Session, periodo: LiqPeriodo, contenido: bytes, usuario_id: int | None) -> dict:
    libro = _abrir(contenido)
    hoja = libro.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        raise ErrorArchivo("El archivo está vacío")
    cols = {_norm(c): i for i, c in enumerate(filas[0]) if c is not None}
    if "CEDULA" not in cols or "TIPO" not in cols:
        raise ErrorArchivo("No se encontraron las columnas CEDULA y TIPO: use la plantilla de novedades de esta nómina")

    def v(fila, nombre):
        i = cols.get(nombre)
        return fila[i] if i is not None and i < len(fila) else None

    en_nomina = {r.empleado.documento: r.empleado for r in db.scalars(select(LiqResultado).where(LiqResultado.periodo_id == periodo.id))}
    indice = Indice(en_nomina.values())
    hallazgos, nuevas = [], []
    for num, fila in enumerate(filas[1:], start=2):
        doc, _ = documento(v(fila, "CEDULA"))
        if doc is None:
            continue
        e = indice.buscar(doc)
        if e is None:
            hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": "No está en esta nómina: no se cargó"})
            continue
        tipo = tipo_de(v(fila, "TIPO"))
        if tipo is None:
            hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": f"Tipo «{_texto(v(fila, 'TIPO'))}» no reconocido: no se cargó"})
            continue
        horas = {}
        for concepto, titulo in COLUMNAS_HORAS:
            h = v(fila, titulo)
            if isinstance(h, (int, float)) and h:
                horas[concepto] = float(h)
        valor = v(fila, "VALOR")
        valor = valor if isinstance(valor, (int, float)) and valor > 0 else None
        if tipo in NOVEDADES_HORAS and not horas and valor is None:
            hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": "Sin horas ni valor: no se cargó"})
            continue
        if tipo not in NOVEDADES_HORAS and valor is None:
            hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": "Sin valor: no se cargó"})
            continue
        fecha, valida = _fecha(v(fila, "FECHA CAUSACION"))
        if not valida:
            hallazgos.append({"hoja": hoja.title, "fila": num, "documento": doc, "detalle": "Fecha de causación inválida: se usa la del inicio de la nómina"})
        salarial = _norm(v(fila, "SALARIAL")) not in ("NO", "N", "0", "FALSO")
        nuevas.append(LiqNovedad(periodo_id=periodo.id, empleado_id=e.id, tipo=tipo, descripcion=_texto(v(fila, "DESCRIPCION"), 200) or TIPOS[tipo],
                                 horas=horas if tipo in NOVEDADES_HORAS else {}, fecha_causacion=fecha if tipo in NOVEDADES_HORAS else None,
                                 valor=valor, salarial=(salarial if tipo == OTRO_DEVENGO else tipo != AUXILIO_VARIABLE),
                                 nomina_origen=_texto(v(fila, "NOMINA ORIGEN"), 80), origen_registro="excel", creado_por=usuario_id))
    reemplazadas = db.execute(delete(LiqNovedad).where(LiqNovedad.periodo_id == periodo.id, LiqNovedad.origen_registro == "excel")).rowcount
    db.add_all(nuevas)
    db.flush()
    return {"cargadas": len(nuevas), "reemplazadas": reemplazadas or 0, "hallazgos": hallazgos}

