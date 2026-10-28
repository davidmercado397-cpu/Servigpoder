"""Costeo de nómina: subir un Excel de turnos para saber cuánto costaría, sin que quede como pagado."""

import json
from io import BytesIO
from types import SimpleNamespace

from fastapi import APIRouter, Depends, File, Form, Query, Request, UploadFile
from openpyxl import Workbook
from pydantic import BaseModel, Field, ValidationError, field_validator
from sqlalchemy import select

from app.api.deps import DbSession, require
from app.core.archivos import leer_xlsx
from app.core.auditoria import auditar
from app.core.paginacion import Paginacion, paginar_consulta
from app.core.rate_limit import limitar
from app.core.respuestas import ApiError, ApiResponse, ok
from app.models import Usuario
from app.apps.liquidador.models import LiqCosteo
from app.apps.liquidador.services import calculo, costeo_nomina, exportar, importador

router = APIRouter(tags=["liquidador: costeo de nómina"])
VER = "liquidador.costeo.ver"
GESTIONAR = "liquidador.costeo.gestionar"


class ParametrosIn(BaseModel):
    salud: float = Field(ge=0, le=100)
    pension: float = Field(ge=0, le=100)
    arl_tarifas: dict[str, float]  # % por clase de riesgo "1".."5"
    arl_clase: int = Field(4, ge=1, le=5)  # para quien no la trae en el Excel ni la tiene configurada en Empleados
    caja: float = Field(ge=0, le=100)
    icbf: float = Field(ge=0, le=100)
    sena: float = Field(ge=0, le=100)
    prima: float = Field(ge=0, le=100)
    cesantias: float = Field(ge=0, le=100)
    intereses_cesantias: float = Field(ge=0, le=100)
    vacaciones: float = Field(ge=0, le=100)
    exonerado: bool = True

    @field_validator("arl_tarifas")
    @classmethod
    def _tarifas(cls, v: dict[str, float]) -> dict[str, float]:
        if set(v) != {"1", "2", "3", "4", "5"} or any(not 0 <= x <= 100 for x in v.values()):
            raise ValueError("Indique la tarifa de ARL (0 a 100 %) de las clases 1 a 5")
        return v


class CosteoEditarIn(BaseModel):
    nombre: str = Field(min_length=1, max_length=120)
    parametros: ParametrosIn


def _resumen(c: LiqCosteo) -> dict:
    t = c.totales or {}
    return {"id": c.id, "nombre": c.nombre, "anio": c.anio, "mes": c.mes, "quincena": c.quincena, "etiqueta": c.etiqueta,
            "desde": c.desde.isoformat(), "hasta": c.hasta.isoformat(), "archivo": c.archivo, "personas": t.get("personas", 0),
            "no_registradas": t.get("no_registradas", 0), "devengado": t.get("devengado", 0), "costo_total": t.get("costo_total", 0),
            "creado_en": c.creado_en.isoformat() if c.creado_en else None,
            "actualizado_en": c.actualizado_en.isoformat() if c.actualizado_en else None}


def _detalle(c: LiqCosteo) -> dict:
    return {**_resumen(c), "umbral_activo": c.umbral_activo, "umbral_horas": float(c.umbral_horas),
            "umbral_horas_novedad": float(c.umbral_horas_novedad), "auxilio_mensual": float(c.auxilio_mensual or 0),
            "parametros": costeo_nomina._completos(c.parametros), "advertencias": c.advertencias or [],
            "totales": c.totales or {}, "resultados": c.resultados or []}


def _costeo(db: DbSession, costeo_id: int) -> LiqCosteo:
    c = db.get(LiqCosteo, costeo_id)
    if c is None:
        raise ApiError(404, "Costeo no encontrado")
    return c


@router.get("/costeos/parametros", response_model=ApiResponse[dict])
def parametros(_=Depends(require(VER))):
    """% por defecto de aportes y prestaciones, con sus nombres."""
    return ok({"valores": costeo_nomina.parametros_defecto(),
               "nombres": {**{k: n for k, (_, n) in costeo_nomina.PARAMETROS.items()}, **costeo_nomina.NOMBRES_EXTRA},
               "aportes": list(costeo_nomina.APORTES), "prestaciones": list(costeo_nomina.PRESTACIONES)})


@router.get("/costeos/plantilla")
def plantilla(anio: int = Query(ge=2020, le=2100), mes: int = Query(ge=1, le=12), quincena: int = Query(ge=0, le=2),
              _=Depends(require(VER))):
    """Plantilla del costeo: la de la nómina con la columna "arl" (clase de riesgo 1 a 5) de cada persona."""
    from app.apps.liquidador.routes.periodos import _archivo

    desde, hasta = calculo.rango_quincena(anio, mes, quincena)
    contenido, tipo, _nombre = exportar.plantilla(SimpleNamespace(anio=anio, mes=mes, quincena=quincena, desde=desde, hasta=hasta,
                                                                   tipo="Mensual" if quincena == 0 else f"Q{quincena}"), con_arl=True)
    return _archivo(contenido, tipo, f"plantilla_costeo_{anio}_{mes:02d}_{'Mensual' if quincena == 0 else f'Q{quincena}'}.xlsx")


@router.get("/costeos", response_model=ApiResponse[list[dict]])
def listar(db: DbSession, p: Paginacion, q: str = Query("", max_length=100), _=Depends(require(VER))):
    consulta = select(LiqCosteo).order_by(LiqCosteo.id.desc())
    if q.strip():
        consulta = consulta.where(LiqCosteo.nombre.ilike(f"%{q.strip()}%"))
    lista, meta = paginar_consulta(db, consulta, p)
    return ok([_resumen(c) for c in lista], **meta)


@router.post("/costeos", response_model=ApiResponse[dict], status_code=201, dependencies=[Depends(limitar("liq_costeo", 10, 60))])
async def crear(request: Request, db: DbSession, archivo: UploadFile = File(...), nombre: str = Form(..., min_length=1, max_length=120),
                anio: int = Form(..., ge=2020, le=2100), mes: int = Form(..., ge=1, le=12), quincena: int = Form(..., ge=0, le=2),
                umbral_activo: bool = Form(False), umbral_horas: float = Form(42, gt=0, le=168),
                umbral_horas_novedad: float = Form(7, ge=0, le=24), auxilio_mensual: float = Form(0, ge=0, le=100_000_000),
                parametros: str = Form("{}", max_length=2000),
                actual: Usuario = Depends(require(GESTIONAR))):
    """Costea el Excel de turnos (mismo formato que una nómina). El archivo se procesa en memoria y se descarta."""
    try:
        par = ParametrosIn(**{**costeo_nomina.parametros_defecto(), **json.loads(parametros or "{}")}).model_dump()
    except (ValueError, ValidationError) as e:
        raise ApiError(422, "Los porcentajes de aportes y prestaciones no son válidos", "DATOS_INVALIDOS") from e
    desde, hasta = calculo.rango_quincena(anio, mes, quincena)
    c = LiqCosteo(nombre=" ".join(nombre.split()), anio=anio, mes=mes, quincena=quincena, desde=desde, hasta=hasta,
                  umbral_activo=umbral_activo, umbral_horas=umbral_horas, umbral_horas_novedad=umbral_horas_novedad,
                  auxilio_mensual=auxilio_mensual, parametros=par, archivo=(archivo.filename or "")[:255], creado_por=actual.id)
    contenido = await leer_xlsx(archivo)
    try:
        res = costeo_nomina.costear(db, c, contenido)
    except importador.ErrorImportacion as e:
        db.rollback()
        raise ApiError(422, str(e), "EXCEL_INVALIDO") from e
    db.add(c)
    db.flush()
    auditar(db, request, "liq_costeo_creado", actual.id, costeo=c.nombre, periodo=c.etiqueta, personas=res.empleados,
            costo_total=c.totales.get("costo_total"))
    db.commit()
    return ok(_detalle(c))


@router.get("/costeos/{costeo_id}", response_model=ApiResponse[dict])
def ver(costeo_id: int, db: DbSession, _=Depends(require(VER))):
    return ok(_detalle(_costeo(db, costeo_id)))


@router.put("/costeos/{costeo_id}", response_model=ApiResponse[dict])
def editar(costeo_id: int, data: CosteoEditarIn, request: Request, db: DbSession, actual: Usuario = Depends(require(GESTIONAR))):
    """Cambia el nombre o los % y recalcula el costo empresa (la nómina de cada persona no cambia)."""
    c = _costeo(db, costeo_id)
    c.nombre, c.parametros = " ".join(data.nombre.split()), data.parametros.model_dump()
    costeo_nomina.recostear(db, c)
    auditar(db, request, "liq_costeo_editado", actual.id, costeo=c.nombre, parametros=c.parametros)
    db.commit()
    return ok(_detalle(c))


@router.delete("/costeos/{costeo_id}", response_model=ApiResponse[dict])
def eliminar(costeo_id: int, request: Request, db: DbSession, actual: Usuario = Depends(require(GESTIONAR))):
    c = _costeo(db, costeo_id)
    auditar(db, request, "liq_costeo_eliminado", actual.id, costeo=c.nombre)
    db.delete(c)
    db.commit()
    return ok({"eliminado": costeo_id})


@router.get("/costeos/{costeo_id}/excel", dependencies=[Depends(limitar("exportar", 10, 60))])
def excel(costeo_id: int, request: Request, db: DbSession, actual: Usuario = Depends(require(VER))):
    from app.apps.liquidador.routes.periodos import _archivo

    c = _costeo(db, costeo_id)
    aportes, prest = costeo_nomina.APORTES, costeo_nomina.PRESTACIONES
    nombres = {**{k: n for k, (_, n) in costeo_nomina.PARAMETROS.items()}, **costeo_nomina.NOMBRES_EXTRA}
    parametros = costeo_nomina._completos(c.parametros)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Costo empresa"
    hoja.append(["DOCUMENTO", "EMPLEADO", "REGISTRADO", "CLASE ARL", "AUXILIO NO PRESTACIONAL MENSUAL", "SALARIO", "TOTAL DEVENGADO", *(nombres[k].upper() for k in aportes),
                 "TOTAL APORTES", *(nombres[k].upper() for k in prest), "TOTAL PRESTACIONES", "COSTO TOTAL"])
    for r in c.resultados or []:
        x = r["costo"]
        hoja.append([r["documento"], r["nombre"], "Sí" if r["existe"] else "No: salario mínimo",
                     x.get("arl_clase") if r.get("arl_clase") else f"{x.get('arl_clase')} (por defecto)", r.get("auxilio_mensual", 0),
                     r["nomina"].get("salario", 0),
                     x["devengado"], *(x["aportes"][k] for k in aportes), x["total_aportes"],
                     *(x["prestaciones"][k] for k in prest), x["total_prestaciones"], x["costo_total"]])
    t = c.totales or {}
    hoja.append(["TOTAL", f"{t.get('personas', 0)} personas", "", None, None, None, t.get("devengado", 0),
                 *(t.get("aportes", {}).get(k, 0) for k in aportes), t.get("total_aportes", 0),
                 *(t.get("prestaciones", {}).get(k, 0) for k in prest), t.get("total_prestaciones", 0), t.get("costo_total", 0)])
    nom = libro.create_sheet("Nomina")
    nom.append(list(exportar.ENCABEZADOS_NOMINA))
    for r in c.resultados or []:
        nom.append(exportar.fila_nomina(SimpleNamespace(empleado=SimpleNamespace(documento=r["documento"], nombre=r["nombre"]),
                                                        nomina=r["nomina"])))
    par = libro.create_sheet("Parametros")
    par.append(["CONCEPTO", "PORCENTAJE"])
    for k in (*aportes, *prest):
        if k == "arl":
            for clase, tarifa in parametros["arl_tarifas"].items():
                par.append([f"ARL clase {clase}", tarifa / 100])
        else:
            par.append([nombres[k], parametros.get(k, 0) / 100])
    filas_pct = par.max_row
    par.append(["Clase de ARL para quien no la tiene", parametros["arl_clase"]])
    par.append(["Auxilio no prestacional mensual por persona", float(c.auxilio_mensual or 0)])
    par.append(["Exoneración art. 114-1 E.T. (sin salud, ICBF ni SENA bajo 10 mínimos)", "Sí" if parametros.get("exonerado") else "No"])
    for h in (hoja, nom, par):
        for celda in h[1]:
            celda.fill, celda.font, celda.alignment = exportar._RELLENO, exportar._FUENTE, exportar._CENTRO
    for col in (5, *range(6, hoja.max_column + 1)):
        for (celda,) in hoja.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = "$ #,##0"
    for (celda,) in par.iter_rows(min_row=2, max_row=filas_pct, min_col=2, max_col=2):
        celda.number_format = "0.00%"
    for h in (hoja, nom, par):
        exportar._autoajustar(h, minimo=10, maximo=40)
    buf = BytesIO()
    libro.save(buf)
    auditar(db, request, "liq_costeo_descargado", actual.id, costeo=c.nombre)
    db.commit()
    return _archivo(buf.getvalue(), exportar.XLSX, f"costeo_{c.anio}_{c.mes:02d}_{c.tipo}{exportar.sufijo(c)}.xlsx")
