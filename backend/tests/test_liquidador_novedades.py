"""Novedades manuales de nómina, Excel de todo el periodo y comparación con el reporte de netos de Siigo."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from tests.conftest import archivo, xlsx
from tests.test_liquidador import API, _cargar
from tests.test_liquidador_nomina import HORA, _nomina, _periodo_con

DIAS = list(range(1, 16))


def _nov(cliente, pid, **datos):
    return cliente.post(f"{API}/periodos/{pid}/novedades", json={"documento": "1001", "descripcion": "Ajuste", **datos})


def test_novedades_manuales(admin):
    pid = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)], [2002, "2002", *(["A"] * 15)]], DIAS)
    n0 = _nomina(admin, pid, "1001")

    r = _nov(admin, pid, tipo="mayor_valor", valor=50000, nomina_origen="Octubre Q2")
    assert r.status_code == 201, r.text
    n = _nomina(admin, pid, "1001")
    assert n["novedades_deduccion"] == 50000 and n["salud"] == n0["salud"] and n["neto"] == n0["neto"] - 50000

    assert _nov(admin, pid, tipo="menor_valor", valor=100000).status_code == 201
    n = _nomina(admin, pid, "1001")
    assert n["ibc"] == n0["ibc"] + 100000 and n["salud"] == n0["salud"] + 4000 and n["devengado"] == n0["devengado"] + 100000

    # Horas causadas en junio (dominical al 80 %) y agosto (al 90 %): cada una con la tarifa de su fecha
    assert _nov(admin, pid, tipo="recargos", horas={"sunday_surcharge": 8}, fecha_causacion="2026-06-14").status_code == 201
    assert _nov(admin, pid, tipo="horas_extras", horas={"overtime_day": 2, "sunday_surcharge": 0}, fecha_causacion="2026-08-10").status_code == 201
    n = _nomina(admin, pid, "1001")
    por_tipo = {x["tipo"]: x for x in n["novedades"]}
    assert por_tipo["recargos"]["valor"] == round(8 * HORA * 0.80) and "80 %" in por_tipo["recargos"]["calculo"]
    assert por_tipo["horas_extras"]["valor"] == round(2 * HORA * 1.25) and "tarifa del 01/07/2026" in por_tipo["horas_extras"]["calculo"]
    devengo = 100000 + round(8 * HORA * 0.80) + round(2 * HORA * 1.25)
    assert n["novedades_devengo"] == devengo and n["ibc"] == n0["ibc"] + devengo

    # Otro devengo no salarial: suma al pago, no a la base
    assert _nov(admin, pid, tipo="otro_devengo", valor=60000, salarial=False).status_code == 201
    n2 = _nomina(admin, pid, "1001")
    assert n2["ibc"] == n["ibc"] and n2["devengado"] == n["devengado"] + 60000

    # Validaciones
    assert _nov(admin, pid, tipo="menor_valor").status_code == 422
    assert _nov(admin, pid, tipo="recargos", horas={"inventado": 3}).status_code == 422
    assert _nov(admin, pid, tipo="mayor_valor", valor=1, documento="9999").status_code == 409
    assert _nov(admin, pid, tipo="recargos", horas={"ordinary_night": 2}, fecha_causacion="2026-12-01").status_code == 422

    # Lo pagado de más que no alcanza: se descuenta hasta el neto y avisa
    _nov(admin, pid, documento="2002", tipo="mayor_valor", valor=5_000_000)
    n3 = _nomina(admin, pid, "2002")
    assert n3["neto"] == 0 and any("quedan" in a for a in n3["alertas"])

    lista = admin.get(f"{API}/periodos/{pid}/novedades").json()["data"]
    assert len(lista) == 6 and all(x["valor"] is not None for x in lista)
    mayor = next(x for x in lista if x["documento"] == "1001" and x["tipo"] == "mayor_valor")
    assert admin.put(f"{API}/novedades/{mayor['id']}", json={"documento": "1001", "tipo": "mayor_valor", "descripcion": "Ajuste",
                                                             "valor": 20000}).status_code == 200
    assert _nomina(admin, pid, "1001")["novedades_deduccion"] == 20000
    assert admin.delete(f"{API}/novedades/{mayor['id']}").status_code == 200

    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/nomina").content))
    filas = list(libro.worksheets[0].iter_rows(values_only=True))
    fila = dict(zip(filas[0], next(f for f in filas if f[0] == "1001")))
    assert fila["NOVEDADES DEVENGO"] == devengo + 60000 and fila["NOVEDADES DEDUCCION"] == 0
    assert "Novedades" in libro.sheetnames

    admin.post(f"{API}/periodos/{pid}/cerrar")
    assert _nov(admin, pid, tipo="menor_valor", valor=1000).status_code == 409


def test_importar_novedades(admin):
    pid = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)]], DIAS)
    _nov(admin, pid, tipo="menor_valor", valor=1000)  # manual: se conserva al importar
    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/novedades/plantilla").content))
    hoja = libro["Novedades"]
    enc = [c.value for c in hoja[1]]
    hoja.delete_rows(2)

    def fila(**v):
        hoja.append([v.get(c) for c in enc])

    fila(CEDULA=1001, TIPO="MAYOR VALOR PAGADO", DESCRIPCION="Pago doble", VALOR=30000, **{"NOMINA ORIGEN": "Oct Q2"})
    fila(CEDULA=1001, TIPO="RECARGOS PENDIENTES", **{"FECHA CAUSACION": "15/10/2026", "RECARGO NOCTURNO": 10})
    fila(CEDULA=1001, TIPO="OTRO DEVENGO", VALOR=5000, SALARIAL="NO")
    fila(CEDULA=7777, TIPO="MENOR VALOR PAGADO", VALOR=1)
    fila(CEDULA=1001, TIPO="BONO", VALOR=1)
    fila(CEDULA=1001, TIPO="HORAS EXTRAS")
    buf = BytesIO()
    libro.save(buf)
    r = admin.post(f"{API}/periodos/{pid}/novedades/importar", files=archivo(buf.getvalue(), "novedades.xlsx"))
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["cargadas"] == 3 and len(d["hallazgos"]) == 3
    n = _nomina(admin, pid, "1001")
    assert n["novedades_deduccion"] == 30000 and n["novedades_devengo"] == 1000 + round(10 * HORA * 0.35) + 5000
    # Importar de nuevo reemplaza las de Excel y conserva la manual
    r = admin.post(f"{API}/periodos/{pid}/novedades/importar", files=archivo(buf.getvalue(), "novedades.xlsx"))
    assert r.json()["data"]["reemplazadas"] == 3 and len(admin.get(f"{API}/periodos/{pid}/novedades").json()["data"]) == 4


def _siigo(hojas: dict[str, list[list]]) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        h = libro.create_sheet(nombre)
        h.append(["NOMBRE", None, "CEDULA", "CONTRATO", "DETALLE", "INGRESOS", "DEDUCCIONES", "NETO A PAGAR"])
        for f in filas:
            h.append(f)
        h.append([None] * 7 + [sum(f[7] for f in filas)])
    buf = BytesIO()
    libro.save(buf)
    return buf.getvalue()


def test_excel_del_periodo_y_comparacion_con_siigo(admin, en_empresa):
    general = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)], [2002, "2002", *(["A"] * 15)]], DIAS)
    otra = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Cliente 2"}).json()["data"]["id"]
    _cargar(admin, otra, xlsx([["documento", "NOVIEMBRE", 1, 2], [3003, "3003", "A", "A"]]))
    a, b, c = (_nomina(admin, pid, doc) for pid, doc in ((general, "1001"), (general, "2002"), (otra, "3003")))

    r = admin.get(f"{API}/nominas-del-periodo/excel?anio=2026&mes=11&quincena=1")
    assert r.status_code == 200 and "nomina_periodo_2026_11_Q1" in r.headers["content-disposition"]
    libro = load_workbook(BytesIO(r.content))
    assert libro.sheetnames == ["Nomina", "Por persona", "Descuentos", "Novedades"]
    filas = list(libro["Nomina"].iter_rows(values_only=True))
    assert filas[0][0] == "NOMINA" and {f[1] for f in filas[1:-1]} == {"1001", "2002", "3003"} and filas[-1][0] == "TOTAL"
    assert admin.get(f"{API}/nominas-del-periodo/excel?anio=2030&mes=1&quincena=1").status_code == 404

    # Siigo: 1001 igual (con centavos), 2002 sin el auxilio de transporte, 3003 no está y 4004 solo en Siigo
    siigo = _siigo({
        "UNIDADES": [["A", None, 1001, "1001", "(4)", a["devengado"] + 0.3, a["deducciones"], a["neto"] + 0.3],
                     ["B", None, 2002, "2002-1", "(4)", b["devengado"] - b["auxilio"], b["deducciones"], b["neto"] - b["auxilio"]]],
        "CVC": [["D", None, 4004, "4004", "(4)", 500000, 40000, 460000]],
        "ACUAVALLE": [],
    })
    datos = {"anio": "2026", "mes": "11", "quincena": "1", "tolerancia": "1"}
    r = admin.post(f"{API}/comparar-siigo", files=archivo(siigo, "netos.xlsx"), data=datos)
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["resumen"]["cuadran"] == 1 and d["resumen"]["con_diferencia"] == 1
    assert d["resumen"]["solo_siigo"] == 1 and d["resumen"]["solo_sistema"] == 1
    filas = {f["documento"]: f for f in d["filas"]}
    assert filas["2002"]["estado"] == "diferencia" and filas["2002"]["diferencia"]["neto"] == round(b["auxilio"], 2)
    assert any("auxilio de transporte" in p for p in filas["2002"]["pistas"])
    assert filas["2002"]["siigo"]["hojas"] == [{"hoja": "UNIDADES", "ingresos": b["devengado"] - b["auxilio"], "deducciones": b["deducciones"],
                                                "neto": b["neto"] - b["auxilio"]}]
    assert filas["4004"]["estado"] == "solo_siigo" and filas["3003"]["estado"] == "solo_sistema"
    assert d["hojas_siigo"]["CVC"] == 460000 and set(d["nominas_sistema"]) == {"General", "Cliente 2"}
    assert round(d["resumen"]["neto_sistema"], 2) == round(a["neto"] + b["neto"] + c["neto"], 2)

    assert d["carga"]["archivo"] == "netos.xlsx"

    # Los datos quedan 24 horas: se vuelve a comparar sin subir el archivo, contra la nómina actual
    consulta = "anio=2026&mes=11&quincena=1"
    r = admin.get(f"{API}/comparar-siigo/excel?{consulta}")
    assert r.status_code == 200 and load_workbook(BytesIO(r.content)).sheetnames == ["Comparacion", "Resumen"]
    admin.post(f"{API}/descuentos", json={"documento": "1001", "tipo": "cuota", "descripcion": "Exequial", "valor_mensual": 20000,
                                          "desde": "2026-11-01"})
    admin.post(f"{API}/periodos/{general}/recalcular")
    otra_vez = admin.get(f"{API}/comparar-siigo?{consulta}").json()["data"]
    assert {f["documento"]: f for f in otra_vez["filas"]}["1001"]["estado"] == "diferencia"
    assert admin.get(f"{API}/comparar-siigo?anio=2026&mes=11&quincena=2").status_code == 404

    # Vencidos (más de 24 horas) se borran solos
    from datetime import datetime, timedelta, timezone

    from app.apps.liquidador.models import LiqSiigoCarga

    s = en_empresa("sera")
    s.query(LiqSiigoCarga).update({"vence_en": datetime.now(timezone.utc) - timedelta(minutes=1)})
    s.commit()
    assert admin.get(f"{API}/comparar-siigo?{consulta}").status_code == 404
    assert s.query(LiqSiigoCarga).count() == 0

    vacio = admin.post(f"{API}/comparar-siigo", files=archivo(xlsx([["hola"]]), "x.xlsx"), data=datos)
    assert vacio.status_code == 422


def test_reporte_de_siigo_solo_con_netos(admin):
    pid = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)], [2002, "2002", *(["A"] * 15)]], DIAS)
    a, b = _nomina(admin, pid, "1001"), _nomina(admin, pid, "2002")
    libro = Workbook()
    h = libro.active
    h.title = "ALCALDIA 1"
    h.append(["NOMBRE", "CEDULA", "NETO A PAGAR"])
    h.append(["A", 1001, a["neto"] + 0.37])
    h.append(["B", 2002, b["neto"] - 9262])
    buf = BytesIO()
    libro.save(buf)
    r = admin.post(f"{API}/comparar-siigo", files=archivo(buf.getvalue(), "netos.xlsx"), data={"anio": "2026", "mes": "11", "quincena": "1"})
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert d["resumen"]["solo_neto"] and d["resumen"]["tolerancia"] == 1000
    filas = {f["documento"]: f for f in d["filas"]}
    # Sin columnas de ingresos ni deducciones no se inventan diferencias: solo el neto
    assert filas["1001"]["estado"] == "cuadra" and filas["1001"]["siigo"]["ingresos"] is None
    assert filas["1001"]["diferencia"]["ingresos"] is None and filas["1001"]["diferencia"]["deducciones"] is None
    assert filas["2002"]["estado"] == "diferencia" and filas["2002"]["diferencia"]["neto"] == 9262
    hoja = load_workbook(BytesIO(admin.get(f"{API}/comparar-siigo/excel?anio=2026&mes=11&quincena=1").content))["Comparacion"]
    enc = [c.value for c in hoja[1]]
    assert "INGRESOS SIIGO" not in enc and "DIF INGRESOS" not in enc and "NETO SIIGO" in enc and "INGRESOS SISTEMA" in enc


def test_codigo_180_y_auxilio_variable(admin):
    # 1001: 10 días trabajados y 5 con código 180; 2002: toda la quincena con código 180
    pid = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 10 + ["180"] * 5)], [2002, "2002", *(["180"] * 15)]], DIAS)
    base_dia = 1750905 / 30
    n = _nomina(admin, pid, "2002")
    ss = n["seguridad_social_180"]
    assert n["neto"] == 0 and n["devengado"] == 0 and n["salud"] == 0
    assert ss["dias"] == 15 and ss["base"] == int(base_dia * 15 + 0.5) and ss["exonerado"] and ss["salud_empleador"] == 0
    assert ss["salud_empleado"] == round(ss["base"] * 0.04) and ss["pension_empleador"] == round(ss["base"] * 0.12)
    assert n["provision_ss_180"] == ss["total"] and any("código 180" in a for a in n["alertas"])
    n1 = _nomina(admin, pid, "1001")
    assert n1["seguridad_social_180"]["dias"] == 5 and n1["dias_salario"] == 10 and n1["neto"] > 0

    # Auxilio variable (rodamiento intermunicipal): suma a los auxilios no prestacionales, sin descuento por TNL
    r = admin.post(f"{API}/periodos/{pid}/novedades", json={"documento": "1001", "tipo": "auxilio_variable",
                                                           "descripcion": "Rodamiento intermunicipal", "valor": 45000})
    assert r.status_code == 201, r.text
    assert r.json()["data"]["valor"] == 45000 and r.json()["data"]["deduccion"] is False
    n2 = _nomina(admin, pid, "1001")
    assert n2["auxilios_no_prestacionales"] == 45000 and n2["novedades_devengo"] == 0
    assert n2["ibc"] == n1["ibc"] and n2["devengado"] == n1["devengado"] + 45000
    assert any(x["descripcion"] == "Rodamiento intermunicipal" for x in n2["auxilios"])

    libro = load_workbook(BytesIO(admin.get(f"{API}/periodos/{pid}/nomina").content))
    filas = list(libro.worksheets[0].iter_rows(values_only=True))
    fila = dict(zip(filas[0], next(f for f in filas if f[0] == "2002")))
    assert fila["DIAS COD 180"] == 15 and fila["PROVISION SEG SOCIAL COD 180"] == ss["total"]


def test_nomina_por_horas(admin):
    turno = {"hora_inicio": "08:00", "hora_fin": "12:00", "horas_extras": 0, "remunerado": True, "incapacidad": False}
    assert admin.post(f"{API}/turnos", json={**turno, "codigo": "H4", "nombre": "4 horas", "horas_ordinarias": 4}).status_code == 201
    assert admin.post(f"{API}/turnos", json={**turno, "codigo": "H2", "nombre": "2 horas", "horas_ordinarias": 2}).status_code == 201

    # Normal (no cambia nada) y dos por horas del mismo periodo
    normal = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)]], DIAS)
    n_normal_antes = _nomina(admin, normal, "1001")
    r = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Puesto A", "tipo_nomina": "horas",
                                            "descuenta_salud": True, "descuenta_pension": False})
    assert r.status_code == 201 and r.json()["data"]["tipo_nomina"] == "horas" and not r.json()["data"]["descuenta_pension"]
    a = r.json()["data"]["id"]
    # 5005: 2 h el 3 y el 4 (y un descanso Z el 5) en el Puesto A; el 4 también 4 h en el Puesto B
    _cargar(admin, a, xlsx([["documento", "NOVIEMBRE", 3, 4, 5], [5005, "5005", "H2", "H2", "Z"]]))
    b = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Puesto B", "tipo_nomina": "horas"}).json()["data"]["id"]
    _cargar(admin, b, xlsx([["documento", "NOVIEMBRE", 4], [5005, "5005", "H4"]]))

    hora = 1750905 / 210
    aux_dia = 249095 / 30
    na = _nomina(admin, a, "5005")
    assert na["por_horas"] and na["horas_pagadas"] == 4 and na["basico"] == int(4 * hora + 0.5)
    assert na["dias_auxilio"] == 2 and na["auxilio"] == int(2 * aux_dia + 0.5)  # el Z no paga auxilio
    assert na["pension"] == 0 and na["salud"] == int(na["basico"] * 0.04 + 0.5)
    nb = _nomina(admin, b, "5005")
    # El 4 ya tiene auxilio en el Puesto A (creado antes): en el B solo las horas
    assert nb["horas_pagadas"] == 4 and nb["dias_auxilio"] == 0 and nb["auxilio"] == 0 and nb["dias_auxilio_otra_nomina"] == 1
    assert any("«Puesto A»" in x and "auxilio de transporte" in x for x in nb["alertas"])
    assert nb["salud"] > 0 and nb["pension"] > 0

    # La nómina normal no cambió
    assert _nomina(admin, normal, "1001") == n_normal_antes

    # Cambiar a normal: se recalcula como nómina normal
    r = admin.put(f"{API}/periodos/{a}/tipo", json={"tipo_nomina": "normal"})
    assert r.status_code == 200 and r.json()["data"]["descuenta_pension"]
    n = _nomina(admin, a, "5005")
    assert not n["por_horas"] and n["dias_salario"] == 3 and n["pension"] > 0


def test_varias_lineas_el_mismo_dia_solo_en_nomina_por_horas(admin):
    turno = {"hora_inicio": "08:00", "hora_fin": "12:00", "horas_extras": 0, "remunerado": True, "incapacidad": False}
    admin.post(f"{API}/turnos", json={**turno, "codigo": "H4", "nombre": "4 horas", "horas_ordinarias": 4})
    excel = xlsx([["documento", "NOVIEMBRE", 3, 4], [5005, "5005", "H4", "H4"], [5005, "5005", "H4", ""], [6006, "6006", "A", "A"]])

    horas = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Por horas", "tipo_nomina": "horas"}).json()["data"]["id"]
    r = _cargar(admin, horas, excel)
    assert r.status_code == 200, r.text
    assert any("más de una línea" in a for a in r.json()["data"]["advertencias"])
    n = _nomina(admin, horas, "5005")
    assert n["horas_pagadas"] == 12 and n["dias_auxilio"] == 2  # 8 h el 3 (dos puestos) + 4 h el 4; auxilio una vez por día
    dias = admin.get(f"{API}/periodos/{horas}/empleados/{_resultado_id(admin, horas, '5005')}").json()["data"]["dias"]
    assert [d["codigo"] for d in dias] == ["H4+H4", "H4"]

    # En la nómina normal la segunda línea se omite como siempre
    normal = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Normal"}).json()["data"]["id"]
    r = _cargar(admin, normal, excel)
    assert any("solo se permiten en la nómina por horas" in a for a in r.json()["data"]["advertencias"])
    assert _nomina(admin, normal, "5005")["dias_salario"] == 2


def _resultado_id(cliente, pid, doc):
    from tests.test_liquidador import _resultados

    return _resultados(cliente, pid)[doc]["empleado_id"]


def test_pensionado_no_descuenta_pension(en_empresa, admin):
    from app.apps.liquidador.models import LiqEmpleado

    normal = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 15)]], DIAS)
    horas = admin.post(f"{API}/periodos", json={"anio": 2026, "mes": 11, "quincena": 1, "nombre": "Por horas", "tipo_nomina": "horas",
                                                "descuenta_salud": True, "descuenta_pension": True}).json()["data"]["id"]
    _cargar(admin, horas, xlsx([["documento", "NOVIEMBRE", 3], [1001, "1001", "A"]]))
    assert _nomina(admin, normal, "1001")["pension"] > 0 and _nomina(admin, horas, "1001")["pension"] > 0

    s = en_empresa("sera")
    s.query(LiqEmpleado).filter(LiqEmpleado.documento == "1001").update({"pensionado": True})
    s.commit()
    for pid in (normal, horas):
        admin.post(f"{API}/periodos/{pid}/recalcular")
        n = _nomina(admin, pid, "1001")
        assert n["pensionado"] and n["pension"] == 0 and n["pension_pct"] == 0 and n["salud"] > 0


def test_auxilio_intermunicipal(admin):
    pid = _periodo_con(admin, 2026, 11, 1, [[1001, "1001", *(["A"] * 13 + ["AUS"] * 2)]], DIAS)
    n0 = _nomina(admin, pid, "1001")
    r = admin.post(f"{API}/periodos/{pid}/novedades", json={"documento": "1001", "tipo": "auxilio_intermunicipal",
                                                           "descripcion": "Auxilio intermunicipal", "valor": 60000})
    assert r.status_code == 201, r.text
    n = _nomina(admin, pid, "1001")
    # No prestacional y completo aunque tenga ausencias: no cambia la base de salud y pensión
    assert n["auxilios_no_prestacionales"] == 60000 and n["ibc"] == n0["ibc"] and n["devengado"] == n0["devengado"] + 60000
    assert any("intermunicipal" in x["calculo"] for x in n["auxilios"])
    assert admin.get(f"{API}/periodos/{pid}/novedades").json()["data"][0]["valor"] == 60000


def test_importar_cuotas_desde_la_lista_de_la_entidad(admin):
    pid = _periodo_con(admin, 2026, 11, 2, [[1001, "1001", *(["A"] * 15)], [2002, "2002", *(["A"] * 15)]], list(range(16, 31)))
    mutual = xlsx([[None] * 8 + [36800], ["No.", None, "TIPO", "IDENTIFICACION", "NOMBRES Y APELLIDOS", "F. NAC", "TEL", "CIUDAD", "VALOR A PAGAR"],
                   [1, 1, "Titular", 1001, "UNO PRUEBA", None, None, "CALI", 15900], [2, 2, "Titular", 2002, "DOS PRUEBA", None, None, "CALI", 15900],
                   [3, 3, "Titular", 9999, "NO ESTA", None, None, "CALI", 5000]])
    datos = {"descripcion": "Cuota mutual", "columna_valor": "I", "quincena": "segunda", "desde": "2026-11-01"}
    r = admin.post(f"{API}/descuentos/importar-cuotas", files=archivo(mutual, "mutual.xlsx"), data=datos)
    assert r.status_code == 200, r.text
    d = r.json()["data"]
    assert (d["creadas"], d["personas"]) == (2, 3) and any("no está en la base" in h["detalle"] for h in d["hallazgos"])
    admin.post(f"{API}/periodos/{pid}/recalcular")
    n = _nomina(admin, pid, "1001")
    assert n["cuotas"] == 15900 and n["descuentos"][0]["descripcion"] == "Cuota mutual"
    # Valores como texto («20900.00») y filas de total al final; solo en la 2.ª quincena
    recordar = xlsx([["Fecha", "Empresa"], [], ["Identificacion", "Afiliado", "Inicio", "Plan", "GBF", "Adicional", "Desc", "Valor Total"],
                     [1001, "UNO PRUEBA", None, "PLAN", "20900.00", 0, None, "20900.00"], [None] * 6 + ["Valor total", "20900.00"]])
    r = admin.post(f"{API}/descuentos/importar-cuotas", files=archivo(recordar, "recordar.xlsx"),
                   data={**datos, "descripcion": "Cuota recordar", "columna_valor": "H"})
    assert r.json()["data"]["creadas"] == 1 and r.json()["data"]["total"] == 20900
    q1 = _periodo_con(admin, 2026, 12, 1, [[1001, "1001", *(["A"] * 15)]], DIAS)
    assert _nomina(admin, q1, "1001")["cuotas"] == 0  # solo 2.ª quincena
    # La siguiente lista ya no trae a 2002: se le desactiva la cuota
    mutual2 = xlsx([["IDENTIFICACION", "NOMBRE", "", "", "", "", "", "", "VALOR"], [1001, "UNO PRUEBA", None, None, None, None, None, None, 16500]])
    d = admin.post(f"{API}/descuentos/importar-cuotas", files=archivo(mutual2, "mutual.xlsx"), data=datos).json()["data"]
    assert (d["actualizadas"], d["desactivadas"]) == (1, 1)
